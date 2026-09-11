import json
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from lqe_provenance import AUDIT_HEADER_BASES
from lqe_report_contract import (
    CONTEXT_AUDIT_HEADERS,
    SHEET_NAME,
    attach_report_contract,
    context_audit_values,
    validate_report_contract,
    verify_native_report,
)
from lqe_delivery import (
    begin_report_attempt,
    complete_report_attempt,
    validate_report_receipt,
)
from lqe_io import _build_xlsx


def workbook_with_results(*, context_headers: bool) -> openpyxl.Workbook:
    workbook = openpyxl.Workbook()
    results = workbook.active
    results.title = "LQE Results"
    headers = [
        *AUDIT_HEADER_BASES.values(),
        "错误详情",
    ]
    if context_headers:
        headers.extend(CONTEXT_AUDIT_HEADERS)
    headers.append("LQE_Iter")
    results.append(headers)
    workbook.create_sheet("LQA Scorecard")
    return workbook


class ReportContextContractTests(unittest.TestCase):
    def test_native_output_provenance_is_required_and_bound_to_job(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            state = {
                "artifact_contract_version": 1,
                "job_runtime_contract_version": 2,
                "context_contract_version": 1,
                "job_id": "job-a",
                "created_at": "2026-09-11T00:00:00Z",
                "scoring_policy": {"scorecard_profile": "lqe_2026"},
                "segments": [],
            }
            results = []
            output = root / "job-a_lqe.xlsx"
            workbook = workbook_with_results(context_headers=True)
            attach_report_contract(workbook, state, results)
            workbook.save(output)
            workbook.close()

            verified = verify_native_report(output, state, results)
            self.assertEqual(verified["job_id"], "job-a")
            self.assertEqual(verified["writer"], "lqe_io._build_xlsx")

            foreign_state = dict(state, job_id="job-b")
            with self.assertRaisesRegex(ValueError, "stale|provenance"):
                verify_native_report(output, foreign_state, results)

    def test_report_without_contract_is_not_native_delivery(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "copied.xlsx"
            workbook = openpyxl.Workbook()
            workbook.save(output)
            workbook.close()
            state = {
                "job_runtime_contract_version": 2,
                "job_id": "job-a",
                "segments": [],
            }
            with self.assertRaisesRegex(ValueError, "missing _LQE_CONTRACT"):
                verify_native_report(output, state, [])

    def test_delivery_receipt_invalidates_previous_report_attempt(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            state = {
                "artifact_contract_version": 1,
                "job_runtime_contract_version": 2,
                "context_contract_version": 1,
                "job_id": "job-a",
                "created_at": "2026-09-11T00:00:00Z",
                "segments": [],
            }
            results = []
            output = root / "job-a_lqe.xlsx"
            workbook = workbook_with_results(context_headers=True)
            attach_report_contract(workbook, state, results)
            workbook.save(output)
            workbook.close()

            generation = begin_report_attempt(root, state)
            complete_report_attempt(root, state, results, output, generation)
            validate_report_receipt(root, state, results, output)

            begin_report_attempt(root, state)
            with self.assertRaisesRegex(ValueError, "not complete"):
                validate_report_receipt(root, state, results, output)

    def test_current_runtime_requires_context_audit_columns_and_writes_v7(self):
        state = {
            "artifact_contract_version": 1,
            "job_runtime_contract_version": 2,
            "context_contract_version": 1,
            "segments": [],
        }
        workbook = workbook_with_results(context_headers=False)
        with self.assertRaisesRegex(ValueError, "context audit columns"):
            attach_report_contract(workbook, state, [])

        workbook = workbook_with_results(context_headers=True)
        attach_report_contract(workbook, state, [])
        contract = json.loads(workbook[SHEET_NAME]["A1"].value)
        self.assertEqual(contract["version"], 7)
        self.assertEqual(contract["job_runtime_contract_version"], 2)
        self.assertEqual(
            contract["context_audit_headers"], list(CONTEXT_AUDIT_HEADERS)
        )
        self.assertIn("Content Type", contract["context_audit_headers"])
        self.assertIn(
            "Capability Resolution Digest",
            contract["context_audit_headers"],
        )

    def test_current_runtime_validates_each_context_audit_value(self):
        segment = {
            "id": 0,
            "segment_key": "dialogue:001",
            "input_status": "ready",
            "content_type": "dialogue",
            "context": {
                "core": {"content_type": "dialogue"},
                "status": "ready",
            },
            "context_provenance": {"context.core.content_type": {"source": "input"}},
            "segment_revision_digest": "a" * 64,
            "resolved_constraints": [{"id": "length.max", "value": 42}],
        }
        state = {
            "artifact_contract_version": 1,
            "job_runtime_contract_version": 2,
            "context_contract_version": 1,
            "capability_resolution_digest": "b" * 64,
            "project_asset_snapshot_digest": "c" * 64,
            "segments": [segment],
        }
        results = [{"id": 0, "errors": [], "corrected": None}]
        workbook = workbook_with_results(context_headers=True)
        workbook["LQE Results"].append([
            0,
            None,
            "不适用",
            "不适用",
            "不适用",
            "",
            *context_audit_values(state, segment),
            0,
        ])
        attach_report_contract(workbook, state, results)

        headers = [cell.value for cell in workbook["LQE Results"][1]]
        context_column = headers.index("Context JSON") + 1
        workbook["LQE Results"].cell(2, context_column).value = "{}"
        with self.assertRaisesRegex(ValueError, "rows do not match"):
            attach_report_contract(workbook, state, results)

    def test_generated_v7_report_exposes_context_coverage_and_input_warning(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "report.xlsx"
            segments = [
                {
                    "id": 0,
                    "source": "Start",
                    "target": "开始",
                    "segment_key": "ui:start",
                    "input_status": "ready",
                    "content_type": "ui_button",
                    "context": {
                        "core": {"content_type": "ui_button"},
                        "status": "ready",
                    },
                    "context_provenance": {},
                    "segment_revision_digest": "1" * 64,
                    "resolved_constraints": [],
                    "input_warnings": [
                        {"code": "UNVERIFIED_TARGET_PROVENANCE"}
                    ],
                },
                {
                    "id": 1,
                    "source": "Continue",
                    "target": "继续",
                    "segment_key": "ui:continue",
                    "input_status": "blocked",
                    "content_type": "ui_button",
                    "context": {
                        "core": {"content_type": "ui_button"},
                        "status": "ready",
                    },
                    "context_provenance": {},
                    "segment_revision_digest": "2" * 64,
                    "resolved_constraints": [],
                },
            ]
            state = {
                "artifact_contract_version": 1,
                "job_runtime_contract_version": 2,
                "context_contract_version": 1,
                "input_path": str(Path(temp_dir) / "source.xlsx"),
                "headers": ["Source", "Target"],
                "rows_raw": [["Start", "开始"], ["Continue", "继续"]],
                "source_col": 0,
                "target_col": 1,
                "wordcount": 2,
                "capability_resolution_digest": "3" * 64,
                "project_asset_snapshot_digest": "4" * 64,
                "segments": segments,
            }
            results = [
                {"id": segment["id"], "errors": [], "corrected": None}
                for segment in segments
            ]
            history = [{"iteration": 0, "errors": results}]

            _build_xlsx(
                state,
                history,
                100,
                98,
                output,
                announce=False,
                report_contract_results=results,
            )

            workbook = openpyxl.load_workbook(output)
            try:
                report = workbook["LQE Results"]
                headers = [cell.value for cell in report[1]]
                content_type_column = headers.index("Content Type") + 1
                capability_column = (
                    headers.index("Capability Resolution Digest") + 1
                )
                self.assertEqual(report.cell(2, content_type_column).value, "ui_button")
                self.assertEqual(report.cell(2, capability_column).value, "3" * 64)
                self.assertEqual(report.cell(3, capability_column).value, "3" * 64)

                scorecard = workbook["LQA Scorecard"]
                self.assertEqual(scorecard["I5"].value, "Segment coverage")
                self.assertEqual(
                    scorecard["J5"].value,
                    "Total: 2 | Reviewable: 1 | Blocked: 1",
                )
                self.assertEqual(scorecard["I6"].value, "Input warning")
                self.assertEqual(
                    scorecard["J6"].value,
                    "UNVERIFIED_TARGET_PROVENANCE: 1 segment(s)",
                )
                self.assertGreaterEqual(scorecard.row_dimensions[6].height, 30)
                validate_report_contract(workbook, state, results)
                report.cell(3, capability_column).value = "forged"
                with self.assertRaisesRegex(ValueError, "rows do not match"):
                    validate_report_contract(workbook, state, results)
            finally:
                workbook.close()

    def test_historical_runtime_keeps_v5_without_context_columns(self):
        state = {"artifact_contract_version": 1, "segments": []}
        workbook = workbook_with_results(context_headers=False)
        attach_report_contract(workbook, state, [])
        contract = json.loads(workbook[SHEET_NAME]["A1"].value)
        self.assertEqual(contract["version"], 5)
        self.assertNotIn("context_audit_headers", contract)


if __name__ == "__main__":
    unittest.main()
