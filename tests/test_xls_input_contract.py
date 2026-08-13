from datetime import datetime
from pathlib import Path
import sys
import tempfile
import types
import unittest
from unittest import mock

import json


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from lqe_inputs.xls import (
    XLSImportError,
    read_xls,
    workbook_for_corrected_export,
)


class FakeCell:
    def __init__(self, value, ctype=1):
        self.value = value
        self.ctype = ctype


class FakeSheet:
    def __init__(self, rows):
        self.rows = rows
        self.nrows = len(rows)
        self.ncols = max((len(row) for row in rows), default=0)

    def cell(self, row_index, column_index):
        if column_index >= len(self.rows[row_index]):
            return FakeCell(None, 0)
        return self.rows[row_index][column_index]


class FakeBook:
    datemode = 0

    def __init__(self, sheets):
        self.sheets = sheets
        self.released = False

    def sheet_names(self):
        return list(self.sheets)

    def sheet_by_name(self, name):
        return self.sheets[name]

    def release_resources(self):
        self.released = True


def fake_xlrd(book, *, version="2.0.1"):
    module = types.ModuleType("xlrd")
    module.__version__ = version
    module.XL_CELL_EMPTY = 0
    module.XL_CELL_TEXT = 1
    module.XL_CELL_NUMBER = 2
    module.XL_CELL_DATE = 3
    module.XL_CELL_BOOLEAN = 4
    module.XL_CELL_ERROR = 5
    module.XL_CELL_BLANK = 6
    module.open_workbook = mock.Mock(return_value=book)
    module.xldate = types.SimpleNamespace(
        xldate_as_datetime=lambda value, datemode: datetime(2026, 8, int(value), 12, 30)
    )
    return module


class LegacyXLSInputContractTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.root = Path(self.tempdir.name)
        self.path = self.root / "neutral.xls"
        self.path.write_bytes(b"neutral fixture; parsing is mocked")

    def tearDown(self):
        self.tempdir.cleanup()

    def book(self):
        return FakeBook(
            {
                "Strings": FakeSheet(
                    [
                        [FakeCell("Key"), FakeCell("Source"), FakeCell("Target")],
                        [FakeCell(1.0, 2), FakeCell("Open"), FakeCell("Ouvrir")],
                        [FakeCell(2.5, 2), FakeCell(13, 3), FakeCell(1, 4)],
                        [FakeCell(None, 0), FakeCell(7, 5), FakeCell(None, 6)],
                    ]
                ),
                "Notes": FakeSheet([[FakeCell("Note")], [FakeCell("Keep values")]]),
            }
        )

    def test_read_preserves_sheet_values_coordinates_and_manifest_limitations(self):
        book = self.book()
        module = fake_xlrd(book)
        with mock.patch.dict(sys.modules, {"xlrd": module}):
            result = read_xls(self.path, sheet_name="Strings")
        self.assertTrue(book.released)
        self.assertEqual(result.sheet_name, "Strings")
        self.assertEqual(result.sheet_names, ["Strings", "Notes"])
        self.assertEqual(result.headers, ["Key", "Source", "Target"])
        self.assertEqual(result.data_rows[0], [1, "Open", "Ouvrir"])
        self.assertEqual(result.data_rows[1][0], 2.5)
        self.assertEqual(result.data_rows[1][1], "2026-08-13 12:30:00")
        self.assertIs(result.data_rows[1][2], True)
        self.assertEqual(result.data_rows[2], [None, "#XLERR:7", None])
        self.assertEqual(result.manifest["adapter"], "xls.xlrd@1")
        self.assertEqual(result.manifest["coordinate_system"], "zero_based_data_row_and_column")
        self.assertIn("corrected_export_is_xlsx", result.manifest["limitations"])
        self.assertIn(
            "formulas_styles_comments_drawings_not_round_tripped",
            result.manifest["limitations"],
        )

    def test_no_header_keeps_first_row_as_data_and_uses_stable_default_headers(self):
        book = self.book()
        with mock.patch.dict(sys.modules, {"xlrd": fake_xlrd(book)}):
            result = read_xls(self.path, no_header=True)
        self.assertEqual(result.headers, ["Key", "Source", "Target"])
        self.assertEqual(result.data_rows[0], ["Key", "Source", "Target"])
        self.assertEqual(result.manifest["rows"], 4)
        self.assertEqual(result.manifest["columns"], 3)

    def test_missing_sheet_and_wrong_suffix_fail_without_publishing(self):
        book = self.book()
        with mock.patch.dict(sys.modules, {"xlrd": fake_xlrd(book)}):
            with self.assertRaisesRegex(XLSImportError, "not found"):
                read_xls(self.path, sheet_name="Missing")
        self.assertTrue(book.released)

        wrong = self.root / "neutral.xlsx"
        wrong.write_bytes(b"not used")
        with self.assertRaisesRegex(XLSImportError, "only accepts .xls"):
            read_xls(wrong)

    def test_missing_or_legacy_xlrd_fails_before_opening_workbook(self):
        book = self.book()
        legacy = fake_xlrd(book, version="1.2.0")
        with mock.patch.dict(sys.modules, {"xlrd": legacy}):
            with self.assertRaisesRegex(XLSImportError, "requires xlrd>=2.0"):
                read_xls(self.path)
        legacy.open_workbook.assert_not_called()

        with mock.patch.dict(sys.modules, {"xlrd": None}):
            with self.assertRaisesRegex(XLSImportError, "requires xlrd>=2.0"):
                read_xls(self.path)

    def test_corrected_export_bridge_copies_all_sheet_value_grids_to_xlsx(self):
        book = self.book()
        with mock.patch.dict(sys.modules, {"xlrd": fake_xlrd(book)}):
            workbook = workbook_for_corrected_export(self.path)
        try:
            self.assertTrue(book.released)
            self.assertEqual(workbook.sheetnames, ["Strings", "Notes"])
            self.assertEqual(workbook["Strings"]["A2"].value, 1)
            self.assertEqual(workbook["Strings"]["B3"].value, "2026-08-13 12:30:00")
            self.assertEqual(workbook["Notes"]["A2"].value, "Keep values")
        finally:
            workbook.close()


class LegacyXLSCLIIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.root = Path(self.tempdir.name)
        self.source = self.root / "neutral.xls"
        self.source.write_bytes(b"neutral fixture; parsing is mocked")
        self.state_path = self.root / "job" / "state.json"

    def tearDown(self):
        self.tempdir.cleanup()

    def book(self):
        return FakeBook(
            {
                "Strings": FakeSheet(
                    [
                        [FakeCell("Key"), FakeCell("Source"), FakeCell("Target")],
                        [FakeCell("row-1"), FakeCell("Open"), FakeCell("Ouvrir")],
                    ]
                ),
                "Notes": FakeSheet([[FakeCell("Note")], [FakeCell("Keep me")]]),
            }
        )

    def run_main(self, arguments, book):
        import lqe_io

        with mock.patch.dict(sys.modules, {"xlrd": fake_xlrd(book)}), mock.patch.object(
            sys, "argv", ["lqe_io.py", *map(str, arguments)]
        ):
            return lqe_io.main()

    def test_cli_read_publishes_xls_manifest_and_export_uses_xlsx_bridge(self):
        self.run_main(
            [
                "read",
                "--input",
                self.source,
                "--sheet",
                "Strings",
                "--source-col",
                "Source",
                "--target-col",
                "Target",
                "--key-col",
                "Key",
                "--source-lang",
                "en",
                "--target-lang",
                "en",
                "--no-terminology",
                "--out",
                self.state_path,
            ],
            self.book(),
        )
        state = json.loads(self.state_path.read_text(encoding="utf-8"))
        manifest = json.loads(
            Path(state["tabular_source_manifest_path"]).read_text(encoding="utf-8")
        )
        self.assertEqual(manifest["adapter"], "xls.xlrd@1")
        self.assertEqual(manifest["sheet_names"], ["Strings", "Notes"])
        self.assertIn("corrected_export_is_xlsx", manifest["limitations"])
        state["segments"][0]["corrected"] = "Open corrected"
        self.state_path.write_text(json.dumps(state), encoding="utf-8")

        self.run_main(["export", "--state", self.state_path], self.book())
        output = self.state_path.parent / "job_corrected.xlsx"
        self.assertTrue(output.is_file())
        import openpyxl

        workbook = openpyxl.load_workbook(output, data_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ["Strings", "Notes"])
            self.assertEqual(workbook["Strings"]["C2"].value, "Open corrected")
            self.assertEqual(workbook["Notes"]["A2"].value, "Keep me")
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
