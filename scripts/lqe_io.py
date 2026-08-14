"""
LQE I/O utilities.

Subcommands:
  read          Excel/CSV/TSV + project profile → state.json
  pre-check     确定性错误自动检测（标点/Markup/术语/长度等）
  protect-segments 把已确认的 TM/100% 匹配段标记为已保护
  apply-fixes   把程序生成的建议译文写回 state.json
  write         state.json + errors.json → *_lqe.xlsx
  ingest-corpus 建议译文回传 AIPE 语料库（接口尚未确定，暂不执行）
"""
import argparse
import csv
import io
import json
import math
import os
import re
import shutil
import sys
import tempfile
import unicodedata
from copy import copy, deepcopy
from datetime import date, datetime, timezone
from pathlib import Path

import openpyxl
import regex
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from lqe_engine import (
    read_json, RE_CJK as _RE_CJK, _source_lang, _target_lang, _load_lang, _LANG_DIR, _SKILL_ROOT,
    CATEGORY_ORDER as _ALL_CATS, CATEGORY_PARENT as _PARENT,
    VALID_CATEGORIES as _VALID_CATEGORIES, VALID_SEVERITIES as _VALID_SEVERITIES,
    apply_severity, build_check_scope, build_review_policy,
    get_check_scope, get_review_policy,
    current_target,
    job_runtime_contract_version,
    load_terms as _load_terms, group_terms as _group_terms,
    require_current_job_runtime,
    requires_bound_artifacts,
    raw_points, weighted_points,
    load_scorecard_profile, normalize_category_for_profile, scorecard_category_order,
    scorecard_category_parent, scorecard_category_weight,
    scorecard_severity_points,
    language_tags_match, normalize_language_tag,
    resolve_language_assets,
    validate_scope_entries,
)
from lqe_corrections import (
    CheckFormatError,
    build_results,
    normalize_check_entries,
    validate_error_history_term_contract,
    verify_results,
)
from lqe_inputs import (
    SDLXLIFFImportError,
    XLSImportError,
    detect_input_format,
    read_sdlxliff,
    read_xls,
)
from lqe_inputs.xls import workbook_for_corrected_export
from lqe_inputs.sdlxliff import (
    is_exact_tm,
    validate_options as validate_sdlxliff_options,
)
from lqe_excel_diff import build_review_rich_texts
from lqe_paths import (
    file_sha256,
    paths_alias as _paths_alias,
    publish_replacement_transaction,
    state_reference_paths,
    validate_artifact_paths,
    write_json_atomic,
)
from lqe_terms import (
    canonicalize_terms,
    load_canonical_terminology,
    terminology_issue_fields,
)
from lqe_scoring import (
    resolve_scoring_policy,
    score_errors,
    scoring_policy_overrides,
)
from lqe_report_contract import attach_report_contract, context_audit_values
from lqe_provenance import (
    AUDIT_HEADER_BASES,
    issue_detail,
    issue_review_columns as _issue_review_columns,
)
from lqe_result_contract import (
    build_result_contract,
    result_contract_path,
    validate_result_contract,
)
from lqe_suggestions import ARTIFACT_NAME, load_reference_suggestions
from lqe_capabilities import (
    JOB_RUNTIME_CONTRACT_VERSION,
    normalize_profile,
    resolve_capabilities,
)
from lqe_project_assets import (
    asset_statuses,
    copy_project_assets,
    inspect_project_assets,
)
from lqe_input_guard import (
    apply_pivot_comparison,
    apply_target_source_digest_guard,
    build_segment_identity,
    canonical_digest as input_guard_digest,
    ensure_unique_business_keys,
    input_guard_summary,
    source_digest,
)
from lqe_context import (
    CORE_CAPABILITY_ID,
    ContextContractError,
    canonical_field_ref,
    descriptor_registry,
    extract_segment_context,
    module_review_equivalence_key,
    parse_context_columns,
    resolve_context_columns,
)
from lqe_language_policies import (
    evaluate_language_policy,
    trusted_provider_registry,
)
from lqe_profile_ingest import (
    apply_segment_context_overrides,
    load_project_source_manifest,
    validate_context_rules,
)
from lqe_shadow import build_shadow_context_artifact
from lqe_context_overrides import (
    build_context_gap_report,
    load_and_apply_job_context_overrides,
)


def _validate_scope_or_exit(
    state: dict,
    entries: list[dict],
    *,
    issues_key: str,
    label: str,
    command: str,
) -> None:
    try:
        get_check_scope(state)
        get_review_policy(state)
        validate_scope_entries(
            state, entries, issues_key=issues_key, label=label
        )
    except ValueError as exc:
        raise SystemExit(f"[{command}] {exc}") from exc


def _processing_label(entry: dict) -> str:
    errors = entry.get("errors") or []
    if any(error.get("protected") for error in errors):
        return "已保护，不修改"
    if any(error.get("needs_confirmation") for error in errors):
        return "需要人工确认"
    if entry.get("corrected") is not None:
        return "建议修改"
    if errors:
        return "仅提醒"
    return "无需修改"


def _issue_processing_label(
    issue: dict | None,
    entry: dict,
    *,
    protected: bool,
) -> str:
    if protected or (issue is not None and issue.get("protected")):
        return "已保护，不修改"
    if issue is None:
        return _processing_label(entry)
    if not isinstance(issue.get("review_provenance"), dict):
        return _processing_label(entry)
    if issue.get("needs_confirmation"):
        return "需要人工确认"
    if issue.get("edit") is not None and entry.get("corrected") is not None:
        return "建议修改"
    return "仅提醒"


# ── read ──────────────────────────────────────────────────────────────────────

def _clean_terms(items: list) -> list:
    return canonicalize_terms(items)
_MAXLEN_KEYS = {"maxlen", "max_len", "max length", "maxlength", "max_length",
                "char_limit", "charlimit", "limit", "width", "ui_max",
                "限长", "字符上限", "长度上限", "字数上限"}


def _parse_maxlen(val) -> int | None:
    if val is None or str(val).strip() == "":
        return None
    try:
        n = int(float(str(val).strip()))
        return n if n > 0 else None
    except ValueError:
        return None


def _job_label(state_path) -> str:
    """输出文件名前缀：标注产物来自哪个任务。取 jobs/ 下的子路径用 _ 连接
    （如 jobs/LQE测试用/剧情/ → 'LQE测试用_剧情'），否则退回 job 目录名。
    避免多文件/多 sheet 拆分时所有 job 都叫 src_*，看不出来源。"""
    d = Path(state_path).resolve().parent
    parts = d.parts
    if "jobs" in parts:
        sub = parts[parts.index("jobs") + 1:]
        if sub:
            return "_".join(sub)
    return d.name


def _load_terminology(
    path: str,
    *,
    term_status_map: object = None,
    protected_statuses: object = None,
) -> list:
    return load_canonical_terminology(
        Path(path),
        term_status_map=term_status_map,
        protected_statuses=protected_statuses,
    )


def _load_project(name_or_path: str) -> dict:
    # 项目名 = <game>/<track>（如 nrc/zh-th、wwm/zh-en），在 skill 根 projects/ 下解析（CWD 无关）；
    # 带后缀或绝对路径按字面处理（支持任意位置的 profile.json）。
    p = Path(name_or_path)
    if not p.suffix and not p.is_absolute():
        p = _SKILL_ROOT / "projects" / p
    if p.is_dir():
        p = p / "profile.json"
    if not p.exists():
        print(f"[ERROR] project profile not found: {p}", file=sys.stderr)
        sys.exit(1)
    prof = read_json(p)
    prof["_dir"] = str(p.parent.resolve())
    prof["_path"] = str(p.resolve())
    return prof


def _deep_merge_profile(base: object, overlay: object, location: str = "profile"):
    if isinstance(base, dict) and isinstance(overlay, dict):
        merged = deepcopy(base)
        for key, value in overlay.items():
            if key in merged:
                merged[key] = _deep_merge_profile(
                    merged[key], value, f"{location}.{key}"
                )
            else:
                merged[key] = deepcopy(value)
        return merged
    if isinstance(base, dict) != isinstance(overlay, dict):
        raise ValueError(f"profile overlay type conflict at {location}")
    return deepcopy(overlay)


def _load_profile_overlay(path: str, base: dict) -> dict:
    overlay_path = Path(path)
    if not overlay_path.is_file():
        raise FileNotFoundError(f"profile overlay not found: {overlay_path}")
    overlay = read_json(overlay_path)
    if not isinstance(overlay, dict):
        raise ValueError("profile overlay must be an object")
    forbidden = {
        "language_pair",
        "source_lang",
        "target_lang",
        "name",
        "profile_contract_version",
    }
    for field in forbidden:
        if field in overlay and overlay[field] != base.get(field):
            raise ValueError(f"profile overlay cannot change {field}")
    merged = _deep_merge_profile(base, overlay)
    merged["_dir"] = base["_dir"]
    merged["_path"] = base["_path"]
    merged["_overlay_path"] = str(overlay_path.resolve())
    return merged


def _validate_project_profile(prof: dict):
    required = ("language_pair", "source_lang", "target_lang")
    missing = [k for k in required if not str(prof.get(k, "")).strip()]
    if missing:
        print("[ERROR] project profile must define language_pair, source_lang, and target_lang; "
              f"missing: {', '.join(missing)}", file=sys.stderr)
        sys.exit(1)
    protected_statuses = prof.get("protected_term_statuses")
    if "protected_term_statuses" in prof and (
        not isinstance(protected_statuses, list)
        or any(
            not isinstance(value, str) or not value.strip()
            for value in protected_statuses
        )
    ):
        print(
            "[ERROR] project profile protected_term_statuses must be an "
            "array of non-empty strings",
            file=sys.stderr,
        )
        sys.exit(1)
    if "scoring_policy" in prof and not isinstance(
        prof["scoring_policy"], dict
    ):
        print(
            "[ERROR] project profile scoring_policy must be an object",
            file=sys.stderr,
        )
        sys.exit(1)
    try:
        normalized = normalize_profile(prof)
        inspection = inspect_project_assets(
            normalized,
            profile_dir=Path(prof["_dir"]),
            allow_outside_root=normalized.get("legacy_adapter", False),
            strict_required=True,
        )
        resolution = resolve_capabilities(
            normalized,
            asset_statuses=asset_statuses(inspection["snapshot"]),
            provider_registry=trusted_provider_registry(),
        )
    except ValueError as exc:
        print(f"[ERROR] project profile contract: {exc}", file=sys.stderr)
        sys.exit(1)
    prof["_normalized_profile"] = normalized
    prof["_asset_inspection"] = inspection
    prof["_capability_resolution"] = resolution


def _project_path(prof: dict, val: str) -> str:
    if not val:
        return ""
    q = Path(val)
    return str(q if q.is_absolute() else Path(prof["_dir"]) / q)


def _profile_reference_paths(prof: dict | None) -> dict[str, Path]:
    if not prof:
        return {}
    references = {"--project": Path(prof["_path"])}
    if prof.get("_overlay_path"):
        references["--profile-overlay"] = Path(prof["_overlay_path"])
    inspection = prof.get("_asset_inspection")
    if isinstance(inspection, dict):
        for asset_id, path in inspection.get("resolved_paths", {}).items():
            references[f"project.asset.{asset_id}"] = Path(path)
    for field in ("style_guide", "terminology", "checks", "confirmed_rules"):
        value = prof.get(field)
        if isinstance(value, str) and value.strip():
            references[f"project.{field}"] = Path(_project_path(prof, value))
    confirmed = references.get("project.confirmed_rules")
    if confirmed is not None:
        references["project.common_confirmed_rules"] = (
            confirmed.parent.parent / "common" / "confirmed_rules_common.md"
        )
    return references


def _profile_asset_by_kind(prof: dict | None, kind: str) -> tuple[str, dict] | None:
    if not prof:
        return None
    normalized = prof.get("_normalized_profile")
    assets = normalized.get("assets", {}) if isinstance(normalized, dict) else {}
    matches = [
        (asset_id, declaration)
        for asset_id, declaration in assets.items()
        if declaration.get("kind") == kind
    ]
    if len(matches) > 1:
        raise ValueError(f"profile declares multiple {kind!r} assets")
    return matches[0] if matches else None


def _active_profile_asset_path(prof: dict | None, kind: str) -> str:
    match = _profile_asset_by_kind(prof, kind)
    if match is None:
        return ""
    asset_id, declaration = match
    snapshot = (prof.get("_asset_inspection") or {}).get("snapshot", {})
    status = snapshot.get("assets", {}).get(asset_id, {}).get("status")
    if status != "present":
        return ""
    return _project_path(prof, declaration["path"])


def _docx_list_value(value: int, fmt: str) -> str:
    if fmt in {"lowerLetter", "upperLetter"}:
        chars = []
        while value > 0:
            value -= 1
            chars.append(chr(ord("a") + value % 26))
            value //= 26
        text = "".join(reversed(chars)) or "0"
        return text.upper() if fmt == "upperLetter" else text
    if fmt in {"lowerRoman", "upperRoman"}:
        parts = []
        remainder = value
        for amount, token in (
            (1000, "M"), (900, "CM"), (500, "D"), (400, "CD"),
            (100, "C"), (90, "XC"), (50, "L"), (40, "XL"),
            (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I"),
        ):
            while remainder >= amount:
                parts.append(token)
                remainder -= amount
        text = "".join(parts) or "0"
        return text.lower() if fmt == "lowerRoman" else text
    return str(value)


def _docx_numbering_levels(doc) -> dict[int, dict[int, dict]]:
    from docx.oxml.ns import qn

    root = doc.part.numbering_part.element
    abstract_by_id = {
        int(node.get(qn("w:abstractNumId"))): node
        for node in root.findall(qn("w:abstractNum"))
    }
    result = {}
    for num in root.findall(qn("w:num")):
        num_id = int(num.get(qn("w:numId")))
        abstract_ref = num.find(qn("w:abstractNumId"))
        if abstract_ref is None:
            continue
        abstract = abstract_by_id.get(int(abstract_ref.get(qn("w:val"))))
        if abstract is None:
            continue
        levels = {}
        for level in abstract.findall(qn("w:lvl")):
            level_id = int(level.get(qn("w:ilvl")))
            start = level.find(qn("w:start"))
            num_fmt = level.find(qn("w:numFmt"))
            level_text = level.find(qn("w:lvlText"))
            levels[level_id] = {
                "start": int(start.get(qn("w:val"))) if start is not None else 1,
                "format": num_fmt.get(qn("w:val")) if num_fmt is not None else "decimal",
                "text": level_text.get(qn("w:val")) if level_text is not None else f"%{level_id + 1}.",
            }
        for override in num.findall(qn("w:lvlOverride")):
            level_id = int(override.get(qn("w:ilvl")))
            start_override = override.find(qn("w:startOverride"))
            if start_override is not None and level_id in levels:
                levels[level_id]["start"] = int(start_override.get(qn("w:val")))
        result[num_id] = levels
    return result


def _docx_list_prefix(para, levels_by_num: dict, counters: dict) -> str:
    ppr = para._p.pPr
    numpr = ppr.numPr if ppr is not None else None
    if numpr is None or numpr.numId is None:
        return ""
    num_id = int(numpr.numId.val)
    level_id = int(numpr.ilvl.val) if numpr.ilvl is not None else 0
    levels = levels_by_num.get(num_id)
    if not levels or level_id not in levels:
        return ""
    level_counters = counters.setdefault(num_id, {})
    start = levels[level_id]["start"]
    level_counters[level_id] = level_counters.get(level_id, start - 1) + 1
    for deeper in [key for key in level_counters if key > level_id]:
        del level_counters[deeper]
    if levels[level_id]["format"] == "bullet":
        return "-"
    prefix = levels[level_id]["text"]
    for referenced_level in range(9):
        marker = f"%{referenced_level + 1}"
        if marker not in prefix:
            continue
        referenced = levels.get(referenced_level, levels[level_id])
        number = level_counters.get(referenced_level, referenced["start"])
        prefix = prefix.replace(marker, _docx_list_value(number, referenced["format"]))
    return prefix


def _docx_table_lines(table) -> list[str]:
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            parts = [part.strip() for part in cell.text.replace("\u00a0", " ").splitlines() if part.strip()]
            cells.append(" / ".join(parts).replace("|", "\\|"))
        rows.append(cells)
    if not rows:
        return []
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    lines = ["| " + " | ".join(rows[0]) + " |"]
    lines.append("| " + " | ".join(["---"] * width) + " |")
    lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return lines


def _docx_style_guide_text(doc) -> str:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    lines = []
    levels_by_num = _docx_numbering_levels(doc)
    counters = {}
    for block in doc.iter_inner_content():
        if isinstance(block, Table):
            table_lines = _docx_table_lines(block)
            if table_lines:
                lines.extend(["", *table_lines, ""])
            continue
        if not isinstance(block, Paragraph):
            continue
        text = block.text.strip()
        if not text:
            continue
        style = block.style.name
        if style.startswith("Heading 1"):
            lines.append(f"\n# {text}")
        elif style.startswith("Heading 2"):
            lines.append(f"\n## {text}")
        elif style.startswith("Heading 3"):
            lines.append(f"\n### {text}")
        else:
            prefix = _docx_list_prefix(block, levels_by_num, counters)
            lines.append(f"{prefix} {text}" if prefix else text)
    return "\n".join(lines)


def _load_style_guide(path: str) -> str:
    p = Path(path)
    if not p.exists():
        print(f"[warn] style-guide file not found: {path}", file=sys.stderr)
        return ""
    suffix = p.suffix.lower()
    if suffix == ".docx":
        import docx
        doc = docx.Document(str(p))
        return _docx_style_guide_text(doc)
    if suffix in (".xlsx", ".xlsm"):
        wb = openpyxl.load_workbook(str(p), data_only=True)
        out = []
        for ws in wb.worksheets:
            rows = []
            for r in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c).strip() for c in r]
                if any(cells):
                    rows.append(cells)
            if len(rows) < 2:
                continue
            header, data = rows[0], rows[1:]
            out.append(f"\n# {ws.title}")
            category = ""
            for r in data:
                title, body = "", []
                for i, v in enumerate(r):
                    h = header[i] if i < len(header) else ""
                    if i == 0 and not h:
                        if v:
                            category = v
                        continue
                    if not v:
                        continue
                    if not title:
                        title = v
                    else:
                        body.append(f"[{h}] {v}" if h else v)
                if not title and not body:
                    continue
                out.append(f"## {category} — {title}" if category else f"## {title}")
                out.extend(body)
        return "\n".join(out)
    return p.read_text(encoding="utf-8")


def _cell(row, idx):
    return row[idx] if idx is not None and idx < len(row) else None


def _text(v):
    return str(v).strip() if v is not None else ""


def _tabular_default_headers(width: int) -> list[str]:
    defaults = [
        "Key", "Source", "Target", "Status", "Comment", "Scope", "File",
        "Reviewer Note",
    ]
    return [
        defaults[index] if index < len(defaults) else f"col{index}"
        for index in range(width)
    ]


def _read_tabular_source(
    path: Path,
    *,
    sheet_name: str | None,
    no_header: bool,
) -> tuple[list, list[list], str, dict]:
    suffix = path.suffix.casefold()
    file_digest = file_sha256(path)
    if suffix in {".csv", ".tsv"}:
        if sheet_name:
            raise ValueError("--sheet is not valid for CSV/TSV input")
        delimiter = "\t" if suffix == ".tsv" else ","
        raw_rows = list(
            csv.reader(
                io.StringIO(path.read_bytes().decode("utf-8-sig")),
                delimiter=delimiter,
            )
        )
        width = max((len(row) for row in raw_rows), default=0)
        if no_header:
            headers = _tabular_default_headers(width)
            data_rows = raw_rows
        else:
            headers = [
                str(value).strip() if value is not None else ""
                for value in (raw_rows[0] if raw_rows else [])
            ]
            data_rows = raw_rows[1:] if raw_rows else []
        container = path.name
        manifest = {
            "schema": "lqe.tabular-source-manifest",
            "version": 1,
            "adapter": "csv@1" if suffix == ".csv" else "tsv@1",
            "format": suffix.lstrip("."),
            "sheet": None,
            "sheet_names": [],
            "rows": len(raw_rows),
            "columns": width,
            "coordinate_system": "zero_based_data_row_and_column",
            "limitations": [],
        }
    elif suffix == ".xls":
        result = read_xls(path, sheet_name=sheet_name, no_header=no_header)
        headers = list(result.headers)
        data_rows = [list(row) for row in result.data_rows]
        container = result.sheet_name
        manifest = deepcopy(result.manifest)
    else:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            if sheet_name is not None and sheet_name not in workbook.sheetnames:
                raise ValueError(
                    f"sheet {sheet_name!r} not found; available sheets: "
                    f"{workbook.sheetnames}"
                )
            selected = sheet_name or workbook.active.title
            worksheet = workbook[selected]
            raw_rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            width = max((len(row) for row in raw_rows), default=0)
            if no_header:
                headers = _tabular_default_headers(width)
                data_rows = raw_rows
            else:
                headers = raw_rows[0] if raw_rows else []
                data_rows = raw_rows[1:] if raw_rows else []
            container = selected
            manifest = {
                "schema": "lqe.tabular-source-manifest",
                "version": 1,
                "adapter": "xlsx.openpyxl@1",
                "format": suffix.lstrip("."),
                "sheet": selected,
                "sheet_names": list(workbook.sheetnames),
                "rows": len(raw_rows),
                "columns": width,
                "coordinate_system": "zero_based_data_row_and_column",
                "limitations": [],
            }
        finally:
            workbook.close()
    manifest.update(
        {
            "input_path": str(path.resolve()),
            "input_sha256": file_digest,
            "no_header": bool(no_header),
        }
    )
    return list(headers), data_rows, container, manifest


def _resolve_input_column(
    headers: list,
    value: object,
    *,
    no_header: bool,
    label: str,
) -> int:
    if no_header:
        try:
            index = int(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(
                f"--no-header mode requires an integer for {label}"
            ) from exc
        if index < 0 or index >= len(headers):
            raise ValueError(f"{label} column index is out of range: {index}")
        return index
    matches = [
        index for index, header in enumerate(headers)
        if str(header or "").strip() == str(value or "").strip()
    ]
    if not matches:
        raise ValueError(f"column {value!r} not found; available: {headers}")
    if len(matches) > 1:
        raise ValueError(f"column {value!r} is duplicated at indexes {matches}")
    return matches[0]


def _context_cli_specs(args) -> list[str]:
    specs = list(getattr(args, "context_cols", None) or [])
    aliases = {
        "content_type_col": "content_type",
        "speaker_col": "speaker_id",
        "addressee_col": "addressee_ids",
        "relationship_stage_col": "relationship_stage",
        "scene_id_col": "scene_id",
        "scene_tone_col": "scene_tone",
        "context_note_col": "context_note",
    }
    for argument, field in aliases.items():
        value = getattr(args, argument, None)
        if value is not None:
            specs.append(f"{field}={value}")
    return specs


def _context_cli_specs_for_registry(
    specs: list[str],
    *,
    source_registry: dict,
    target_registry: dict,
) -> dict[str, object]:
    target_fields = {
        f"{capability_id}.{field_name}"
        for capability_id, descriptor in target_registry.items()
        for field_name in descriptor.get("fields", {})
    }
    selected = {}
    for raw_ref, column in parse_context_columns(specs).items():
        canonical = canonical_field_ref(raw_ref, source_registry)
        if canonical in selected:
            raise ContextContractError(
                f"duplicate CLI mapping for context field: {canonical}"
            )
        if canonical in target_fields:
            selected[canonical] = column
    return selected


_PIVOT_BUILTIN_FIELDS = {"source", "target", "key"}


def _pivot_digest_value(value: object) -> object:
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    isoformat = getattr(value, "isoformat", None)
    if callable(isoformat):
        return isoformat()
    return str(value)


def _tabular_grid_digest(headers: list, rows: list[list]) -> str:
    return input_guard_digest(
        {
            "headers": [_pivot_digest_value(value) for value in headers],
            "rows": [
                [_pivot_digest_value(value) for value in row]
                for row in rows
            ],
        }
    )


def _column_binding(
    headers: list,
    index: int,
    *,
    no_header: bool,
    method: str,
) -> dict:
    return {
        "column_index": index,
        "column": index if no_header else _text(headers[index]),
        "method": method,
    }


def _parse_pivot_compare_specs(specs: list[str], registry: dict) -> list[dict]:
    comparisons = []
    seen = set()
    for spec in specs:
        if not isinstance(spec, str) or spec.count("=") != 1:
            raise ValueError(
                "--pivot-compare must use FIELD=COLUMN exactly once"
            )
        raw_field, raw_column = (part.strip() for part in spec.split("=", 1))
        if not raw_field or not raw_column:
            raise ValueError(
                "--pivot-compare field and column must be non-empty"
            )
        builtin = raw_field.casefold()
        if builtin in _PIVOT_BUILTIN_FIELDS:
            field = builtin
            kind = "builtin"
        else:
            field = canonical_field_ref(raw_field, registry)
            kind = "context"
        if field in seen:
            raise ValueError(f"duplicate --pivot-compare field: {field}")
        seen.add(field)
        comparisons.append(
            {
                "field": field,
                "kind": kind,
                "pivot_column": raw_column,
            }
        )
    return comparisons


def _context_comparison_normalizer(field_ref: str, registry: dict) -> tuple[str, str]:
    capability_id, field_name = field_ref.rsplit(".", 1)
    descriptor_normalizer = registry[capability_id]["fields"][field_name][
        "normalizer"
    ]
    comparison_normalizer = {
        "identity": "text",
        "trim": "trim",
        "lowercase": "text",
        "positive_integer": "integer",
        "string_list": "text",
    }.get(descriptor_normalizer)
    if comparison_normalizer is None:
        raise ValueError(
            f"unsupported pivot normalizer for {field_ref}: "
            f"{descriptor_normalizer!r}"
        )
    return descriptor_normalizer, comparison_normalizer


def _prepare_pivot_guard(
    args,
    *,
    input_path: Path,
    main_container: str,
    main_headers: list,
    source_index: int,
    target_index: int,
    key_index: int | None,
    resolved_context_columns: dict,
    registry: dict,
    no_header: bool,
) -> dict | None:
    pivot_sheet = getattr(args, "pivot_sheet", None)
    pivot_key_col = getattr(args, "pivot_key_col", None)
    pivot_compare = list(getattr(args, "pivot_compare", None) or [])
    pivot_authority = getattr(args, "pivot_authority", None)
    configured = any(
        value not in (None, [], "")
        for value in (
            pivot_sheet,
            pivot_key_col,
            pivot_compare,
            pivot_authority,
        )
    )
    if not configured:
        return None
    missing = [
        flag
        for flag, value in (
            ("--pivot-sheet", pivot_sheet),
            ("--pivot-key-col", pivot_key_col),
            ("--pivot-compare", pivot_compare),
            ("--pivot-authority", pivot_authority),
        )
        if value in (None, [], "")
    ]
    if missing:
        raise ValueError(
            "pivot configuration requires all of --pivot-sheet, "
            "--pivot-key-col, --pivot-compare, and --pivot-authority; missing: "
            + ", ".join(missing)
        )
    if key_index is None:
        raise ValueError("pivot join requires an explicit --key-col")
    if input_path.suffix.casefold() not in {".xlsx", ".xls"}:
        raise ValueError(
            "--pivot-sheet is only valid for .xlsx or .xls workbook input"
        )

    pivot_headers, pivot_rows, pivot_container, pivot_manifest = (
        _read_tabular_source(
            input_path,
            sheet_name=str(pivot_sheet),
            no_header=no_header,
        )
    )
    if pivot_container == main_container:
        raise ValueError("pivot sheet must differ from the main sheet")
    if not pivot_rows:
        raise ValueError("pivot sheet has no data rows")
    pivot_widths = {
        len(row) for row in pivot_rows if any(_text(cell) for cell in row)
    }
    if len(pivot_widths) > 1:
        raise ValueError(
            f"pivot sheet has inconsistent row widths: {sorted(pivot_widths)}"
        )

    pivot_key_index = _resolve_input_column(
        pivot_headers,
        pivot_key_col,
        no_header=no_header,
        label="--pivot-key-col",
    )
    parsed = _parse_pivot_compare_specs(pivot_compare, registry)
    pivot_rows_by_key = {}
    pivot_row_indexes = {}
    for row_index, row in enumerate(pivot_rows):
        if not any(_text(cell) for cell in row):
            continue
        key = _text(_cell(row, pivot_key_index))
        if not key:
            raise ValueError(
                f"pivot business key is empty at data row {row_index}"
            )
        if key in pivot_rows_by_key:
            raise ValueError(
                f"duplicate pivot business key {key!r}: data rows "
                f"{pivot_row_indexes[key]} and {row_index}"
            )
        pivot_rows_by_key[key] = list(row)
        pivot_row_indexes[key] = row_index
    if not pivot_rows_by_key:
        raise ValueError("pivot sheet has no keyed data rows")

    comparisons = []
    for item in parsed:
        field = item["field"]
        pivot_index = _resolve_input_column(
            pivot_headers,
            item["pivot_column"],
            no_header=no_header,
            label=f"--pivot-compare {field}",
        )
        if item["kind"] == "builtin":
            primary_index = {
                "source": source_index,
                "target": target_index,
                "key": key_index,
            }[field]
            primary_mapping = _column_binding(
                main_headers,
                primary_index,
                no_header=no_header,
                method=f"primary_{field}_column",
            )
            descriptor_normalizer = "trim"
            comparison_normalizer = "trim"
        else:
            if field not in resolved_context_columns:
                raise ValueError(
                    f"pivot comparison field {field} has no primary context "
                    "column mapping"
                )
            primary_mapping = deepcopy(resolved_context_columns[field])
            descriptor_normalizer, comparison_normalizer = (
                _context_comparison_normalizer(field, registry)
            )
        pivot_mapping = _column_binding(
            pivot_headers,
            pivot_index,
            no_header=no_header,
            method="pivot_cli",
        )
        if item["kind"] == "context":
            pivot_mapping.update(
                {
                    key: primary_mapping[key]
                    for key in (
                        "capability_id",
                        "descriptor_id",
                        "extension",
                        "field",
                    )
                }
            )
        comparisons.append(
            {
                "field": field,
                "kind": item["kind"],
                "normalizer": comparison_normalizer,
                "descriptor_normalizer": descriptor_normalizer,
                "primary": primary_mapping,
                "pivot": pivot_mapping,
            }
        )

    public_comparisons = [
        {
            "field": item["field"],
            "kind": item["kind"],
            "normalizer": item["normalizer"],
            "descriptor_normalizer": item["descriptor_normalizer"],
            "primary": deepcopy(item["primary"]),
            "pivot": deepcopy(item["pivot"]),
        }
        for item in comparisons
    ]
    mapping_payload = {
        "main_sheet": main_container,
        "pivot_sheet": pivot_container,
        "key": {
            "primary": _column_binding(
                main_headers,
                key_index,
                no_header=no_header,
                method="primary_key_column",
            ),
            "pivot": _column_binding(
                pivot_headers,
                pivot_key_index,
                no_header=no_header,
                method="pivot_key_column",
            ),
        },
        "comparisons": public_comparisons,
        "authority": pivot_authority,
    }
    return {
        "rows_by_key": pivot_rows_by_key,
        "row_indexes": pivot_row_indexes,
        "comparisons": comparisons,
        "comparison_registry": {
            capability_id: {
                **deepcopy(descriptor),
                "applies_when": {},
            }
            for capability_id, descriptor in registry.items()
        },
        "public": {
            "schema": "lqe.pivot-guard",
            "version": 1,
            "workbook_sha256": pivot_manifest["input_sha256"],
            "main_sheet": main_container,
            "pivot_sheet": pivot_container,
            "pivot_sheet_digest": _tabular_grid_digest(
                pivot_headers, pivot_rows
            ),
            "authority": pivot_authority,
            "key": deepcopy(mapping_payload["key"]),
            "comparisons": public_comparisons,
            "mapping_digest": input_guard_digest(mapping_payload),
        },
    }


def _context_values_for_pivot(
    row: list,
    comparisons: list[dict],
    registry: dict,
    *,
    side: str,
) -> dict:
    mappings = {
        item["field"]: item[side]
        for item in comparisons
        if item["kind"] == "context"
    }
    if not mappings:
        return {}
    context = extract_segment_context(row, mappings, registry)
    values = {}
    for field_ref, mapping in mappings.items():
        if mapping["capability_id"] == CORE_CAPABILITY_ID:
            value = context.get("core", {}).get(mapping["field"])
        else:
            value = (
                context.get("extensions", {})
                .get(mapping["extension"], {})
                .get(mapping["field"])
            )
        values[field_ref] = value
    return values


def _apply_pivot_guard(
    segments: list[dict],
    rows: list[list],
    pivot: dict,
) -> dict:
    if len(segments) != len(rows):
        raise ValueError("pivot join cannot align segments with primary rows")
    primary_keys = []
    for segment in segments:
        if segment.get("key_origin") != "input":
            raise ValueError(
                "pivot join requires a non-empty explicit business key for "
                f"segment {segment.get('id')}"
            )
        primary_keys.append(segment["segment_key"])
    primary_key_set = set(primary_keys)
    pivot_key_set = set(pivot["rows_by_key"])
    missing = sorted(primary_key_set - pivot_key_set)
    extra = sorted(pivot_key_set - primary_key_set)
    if missing or extra:
        details = []
        if missing:
            details.append(f"missing pivot keys: {missing}")
        if extra:
            details.append(f"extra pivot keys: {extra}")
        raise ValueError("pivot key coverage mismatch; " + "; ".join(details))

    guard_rules = [
        {"field": item["field"], "normalizer": item["normalizer"]}
        for item in pivot["comparisons"]
    ]
    for segment, primary_row in zip(segments, rows):
        key = segment["segment_key"]
        pivot_row = pivot["rows_by_key"][key]
        primary_context = _context_values_for_pivot(
            primary_row,
            pivot["comparisons"],
            pivot["comparison_registry"],
            side="primary",
        )
        pivot_context = _context_values_for_pivot(
            pivot_row,
            pivot["comparisons"],
            pivot["comparison_registry"],
            side="pivot",
        )
        primary_values = dict(primary_context)
        pivot_values = dict(pivot_context)
        for item in pivot["comparisons"]:
            field = item["field"]
            if item["kind"] != "builtin":
                continue
            primary_values[field] = {
                "source": segment.get("source"),
                "target": segment.get("target"),
                "key": key,
            }[field]
            pivot_values[field] = _cell(
                pivot_row, item["pivot"]["column_index"]
            )
        apply_pivot_comparison(
            segment,
            primary_values=primary_values,
            pivot_values=pivot_values,
            comparisons=guard_rules,
            authority=pivot["public"]["authority"],
        )

    public = deepcopy(pivot["public"])
    public.update(
        {
            "primary_keys": len(primary_key_set),
            "pivot_keys": len(pivot_key_set),
            "join_key_digest": input_guard_digest(sorted(primary_key_set)),
        }
    )
    public["digest"] = input_guard_digest(public)
    for segment in segments:
        key = segment["segment_key"]
        segment["pivot_guard_digest"] = public["digest"]
        segment["pivot_provenance"] = {
            "source_file_digest": public["workbook_sha256"],
            "sheet": public["pivot_sheet"],
            "row_index": pivot["row_indexes"][key],
            "authority": public["authority"],
        }
    return public


def _tabular_source_provenance(
    manifest: dict,
    *,
    row_index: int,
) -> dict:
    return {
        "adapter": manifest["adapter"],
        "source_file_digest": manifest["input_sha256"],
        "container": manifest.get("sheet") or Path(manifest["input_path"]).name,
        "row_index": row_index,
    }


def _segment_revision(segment: dict, state_context: dict) -> str:
    payload = {
        "segment_key": segment.get("segment_key"),
        "key_origin": segment.get("key_origin"),
        "source_digest": segment.get("source_digest"),
        "target": current_target(segment),
        "input_status": segment.get("input_status"),
        "input_block_reasons": segment.get("input_block_reasons", []),
        "input_warnings": segment.get("input_warnings", []),
        "context": segment.get("context"),
        "resolved_constraints": segment.get("resolved_constraints", []),
        "pivot_guard_digest": segment.get("pivot_guard_digest"),
        "capability_resolution_digest": state_context.get(
            "capability_resolution_digest"
        ),
        "project_asset_snapshot_digest": state_context.get(
            "project_asset_snapshot_digest"
        ),
        "context_overrides_fingerprint": state_context.get(
            "context_overrides_fingerprint"
        ),
    }
    return source_digest(
        json.dumps(
            payload,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        )
    )


def _foundation_context_registry() -> dict:
    return descriptor_registry(
        capability_resolution={
            "enabled": {
                CORE_CAPABILITY_ID: {},
                "source_provenance@1": {},
            }
        }
    )


def _registry_for_profile(
    prof: dict | None,
    *,
    include_shadow: bool = False,
) -> dict:
    normalized = prof.get("_normalized_profile") if prof else None
    resolution = prof.get("_capability_resolution") if prof else None
    if normalized and resolution:
        effective_resolution = deepcopy(resolution)
        if not include_shadow:
            effective_resolution["enabled"] = {
                capability_id: item
                for capability_id, item in resolution.get("enabled", {}).items()
                if item.get("effect") in {"foundation", "enforce"}
            }
        return descriptor_registry(
            normalized,
            capability_resolution=effective_resolution,
        )
    return _foundation_context_registry()


def _module_context_views_for_profile(prof: dict | None) -> tuple[dict, dict | None]:
    normalized = prof.get("_normalized_profile") if prof else None
    if not isinstance(normalized, dict):
        return {}, None
    views = deepcopy(normalized.get("module_context_views", {}))
    mode = normalized.get("context_pipeline", {}).get("mode", "off")
    if mode == "enforce":
        return views, None
    if mode == "shadow":
        return {}, views
    return {}, None


def _legacy_segment_context(
    segment: dict,
    registry: dict,
    *,
    source_provenance: dict,
) -> dict:
    headers = ["content_type", "context_note"]
    resolved = resolve_context_columns(
        headers,
        registry,
        cli_columns={
            "content_type": "content_type",
            "context_note": "context_note",
        },
    )
    return extract_segment_context(
        [segment.get("content_type"), segment.get("context_note")],
        resolved,
        registry,
        source_provenance=source_provenance,
    )


def _read_json_asset(path: str | Path, *, label: str) -> dict:
    value = read_json(Path(path))
    if not isinstance(value, dict):
        raise ValueError(f"{label} must be a JSON object")
    return value


def _refresh_context_extension_statuses(
    segment: dict,
    registry: dict,
) -> None:
    """Re-evaluate required extension fields after verified sidecar patches."""

    context = segment.get("context")
    if not isinstance(context, dict):
        return
    extensions = context.get("extensions")
    if not isinstance(extensions, dict):
        return
    missing = []
    for capability_id, descriptor in registry.items():
        if capability_id == CORE_CAPABILITY_ID:
            continue
        extension = capability_id.removeprefix("context.").split("@", 1)[0]
        value = extensions.get(extension)
        if not isinstance(value, dict) or value.get("status") in {
            "not_applicable",
            "disabled",
            "conflict",
        }:
            continue
        capability_missing = []
        for field_name, field in descriptor.get("fields", {}).items():
            if field.get("required_when_applicable") is not True:
                continue
            field_value = value.get(field_name)
            if field_value in (None, "", [], {}):
                capability_missing.append(
                    f"context.extensions.{extension}.{field_name}"
                )
        value["status"] = "incomplete" if capability_missing else "ready"
        missing.extend(capability_missing)
    context["missing_required"] = sorted(set(missing))
    context["status"] = "context_incomplete" if missing else "ready"
    segment["context_provenance"] = deepcopy(context.get("provenance", {}))
    segment["context_status"] = context["status"]
    segment["context_missing_required"] = list(context["missing_required"])


def _apply_runtime_project_context(
    segments: list[dict],
    common: dict,
    prof: dict | None,
    registry: dict,
    *,
    runtime_asset_paths: dict | None = None,
    phase: str = "all",
) -> list[dict]:
    """Apply only declared, copied canonical assets to the formal context."""

    if phase not in {"all", "project_overrides", "rules"}:
        raise ValueError(f"unknown project context phase: {phase}")
    if not prof:
        return segments
    mode = (common.get("context_pipeline") or {}).get("mode", "off")
    if mode != "enforce":
        return segments
    snapshot = common.get("project_asset_snapshot") or {}
    paths = runtime_asset_paths or common.get("project_asset_paths") or {}
    enabled = (common.get("capability_resolution") or {}).get("enabled", {})
    enabled_asset_ids = {
        item.get("asset")
        for item in enabled.values()
        if isinstance(item, dict) and isinstance(item.get("asset"), str)
    }
    source_provenance_enabled = "source_provenance@1" in enabled
    by_kind: dict[str, list[tuple[str, Path]]] = {}
    for asset_id, entry in (snapshot.get("assets") or {}).items():
        path = paths.get(asset_id)
        kind = str(entry.get("kind"))
        allowed = asset_id in enabled_asset_ids or (
            kind == "project_source_manifest" and source_provenance_enabled
        )
        if allowed and entry.get("status") == "present" and path:
            by_kind.setdefault(str(entry.get("kind")), []).append(
                (asset_id, Path(path))
            )
    for kind in by_kind:
        by_kind[kind].sort()

    source_ids = None
    manifests = by_kind.get("project_source_manifest", [])
    if manifests:
        if len(manifests) != 1:
            raise ValueError("only one active project_source_manifest is allowed")
        manifest = load_project_source_manifest(manifests[0][1])
        source_ids = [item["id"] for item in manifest["sources"]]
        common["project_source_manifest_digest"] = manifest["manifest_digest"]
        common["project_source_manifest_path"] = common[
            "project_asset_paths"
        ][manifests[0][0]]

    overrides = by_kind.get("segment_context_overrides", [])
    if phase in {"all", "project_overrides"} and overrides:
        if source_ids is None:
            raise ValueError(
                "segment_context_overrides requires an enabled project_source_manifest"
            )
        if len(overrides) != 1:
            raise ValueError("only one active segment_context_overrides asset is allowed")
        declared_extensions = {
            descriptor_id: descriptor
            for descriptor_id, descriptor in registry.items()
            if descriptor_id != CORE_CAPABILITY_ID
        }
        updated = apply_segment_context_overrides(
            _read_json_asset(overrides[0][1], label="segment_context_overrides"),
            segments=segments,
            declared_extensions=declared_extensions,
            source_ids=source_ids,
        )
        segments = updated
        for segment in segments:
            context = segment.get("context")
            if not isinstance(context, dict):
                continue
            sidecar_provenance = segment.pop("provenance", {})
            if isinstance(sidecar_provenance, dict):
                context.setdefault("provenance", {}).update(sidecar_provenance)
            _refresh_context_extension_statuses(segment, registry)
        common["segment_context_overrides_digest"] = input_guard_digest(
            _read_json_asset(overrides[0][1], label="segment_context_overrides")
        )

    rule_assets = by_kind.get("context_rules", [])
    if phase in {"all", "rules"} and rule_assets:
        if len(rule_assets) != 1:
            raise ValueError("only one active context_rules asset is allowed")
        target_lang = common.get("target_lang")
        rules = validate_context_rules(
            _read_json_asset(rule_assets[0][1], label="context_rules"),
            target_lang=target_lang,
        )
        enabled = (common.get("capability_resolution") or {}).get("enabled", {})
        policy_capabilities = [
            (capability_id, resolution)
            for capability_id, resolution in enabled.items()
            if capability_id.startswith("language_policy.")
            and resolution.get("effect") == "enforce"
        ]
        for capability_id, resolution in policy_capabilities:
            provider = resolution.get("provider")
            if not isinstance(provider, dict):
                raise ValueError(
                    f"enabled language policy {capability_id} lacks provider binding"
                )
            provider_ref = {
                "id": provider["id"],
                "api_version": provider["api_version"],
            }
            for segment in segments:
                evaluation = evaluate_language_policy(
                    rules,
                    segment.get("context", {}),
                    current_target(segment),
                    provider=provider_ref,
                    target_lang=target_lang,
                    as_of=common.get("created_at"),
                )
                constraint = deepcopy(evaluation["constraint"])
                constraint["runtime_evaluation"] = {
                    "status": evaluation["status"],
                    "reason_codes": deepcopy(evaluation.get("reason_codes", [])),
                    "observation": deepcopy(evaluation.get("observation")),
                    "evaluation": deepcopy(evaluation.get("evaluation")),
                    "evaluation_digest": evaluation["evaluation_digest"],
                }
                segment.setdefault("resolved_constraints", []).append(constraint)
                segment.setdefault("constraint_evaluations", []).append(
                    evaluation
                )
        common["context_rules_digest"] = input_guard_digest(rules)
    return segments


def _apply_explicit_context_overrides(
    args,
    segments: list[dict],
    common: dict,
    registry: dict,
    *,
    shadow_registry: dict | None,
    staging_dir: Path,
    job_dir: Path,
    staged_project_asset_paths: dict,
) -> tuple[list[dict], dict | None, dict | None, tuple[Path, str] | None]:
    source_path = getattr(args, "context_overrides", None)
    if not source_path:
        return segments, None, None, None
    result = load_and_apply_job_context_overrides(
        source_path,
        segments=segments,
        registry=registry,
        capability_resolution_digest=common.get("capability_resolution_digest"),
    )
    segments = result["segments"]
    published_sidecar = job_dir / "context_overrides.json"
    staged_sidecar = staging_dir / published_sidecar.name
    staged_sidecar.write_text(
        json.dumps(result["document"], ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    audit = deepcopy(result["audit"])
    audit["path"] = str(published_sidecar)
    report_state = {
        **common,
        "job_runtime_contract_version": JOB_RUNTIME_CONTRACT_VERSION,
        "resolved_context_descriptors": registry,
        "shadow_context_descriptors": shadow_registry,
        "project_asset_paths": staged_project_asset_paths,
        "segments": segments,
    }
    gap_report = build_context_gap_report(report_state)
    published_gap_report = job_dir / "context_gap_report.json"
    staged_gap_report = staging_dir / published_gap_report.name
    staged_gap_report.write_text(
        json.dumps(gap_report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    common.update(
        {
            "context_overrides": audit,
            "context_overrides_path": str(published_sidecar),
            "context_overrides_digest": audit["document_digest"],
            "context_overrides_fingerprint": audit["fingerprint"],
            "context_runtime_fingerprint": audit["fingerprint"],
            "context_gap_report_path": str(published_gap_report),
            "context_gap_report_digest": gap_report["report_digest"],
        }
    )
    guard = (Path(result["source_path"]), audit["source_file_sha256"])
    return segments, audit, gap_report, guard


def _extract_text_type_marker(
    src,
    tgt=None,
    content_type=None,
    *,
    rules: list[dict] | None = None,
):
    """Apply only profile-declared marker rows; ordinary text is never inferred."""

    source = _text(src)
    for rule in rules or []:
        if source != rule["source_equals"]:
            continue
        if "text_type_context" in rule:
            return rule["text_type_context"]
        values = {
            "source": source,
            "target": _text(tgt),
            "content_type": _text(content_type),
        }
        marker = values[rule["text_type_from"]]
        if not marker:
            raise ValueError(
                f"text type marker rule {rule['id']!r} resolved to an empty value"
            )
        return marker
    return None


def _prepare_read_assets(
    args,
    prof: dict | None,
    check_scope: dict,
    review_policy: dict,
    segments: list[dict],
    source_lang: str,
    target_lang: str,
    *,
    asset_dir: Path | None = None,
    publish_dir: Path | None = None,
) -> dict:
    output_dir = Path(args.out).parent
    write_dir = Path(asset_dir) if asset_dir is not None else output_dir
    final_dir = Path(publish_dir) if publish_dir is not None else write_dir
    write_dir.mkdir(parents=True, exist_ok=True)

    sg_path = ""
    if args.style_guide:
        sg_text = _load_style_guide(args.style_guide)
        if sg_text:
            staged_file = write_dir / "sg.txt"
            published_file = final_dir / "sg.txt"
            staged_file.write_text(sg_text, encoding="utf-8")
            sg_path = str(published_file)
            print(f"[lqe_io] style_guide: {len(sg_text)} chars → {published_file}")

    terms_path = ""
    if args.terminology:
        protected_statuses = (
            prof.get("protected_term_statuses", []) if prof else []
        )
        terms = _load_terminology(
            args.terminology,
            term_status_map=prof.get("term_status_map") if prof else None,
            protected_statuses=protected_statuses,
        )
        if terms:
            n_protected = sum(
                1
                for term in terms
                for sense in term.get("senses", [term])
                if sense.get("protected") is True
            )
            if n_protected:
                print(f"[lqe_io] protected term senses: {n_protected}")
            staged_file = write_dir / "terms.json"
            published_file = final_dir / "terms.json"
            staged_file.write_text(
                json.dumps(terms, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            terms_path = str(published_file)
            print(
                f"[lqe_io] terminology: {len(terms)} entries → {published_file}"
            )

    requested_asset_lang = (
        _target_lang({"target_lang": getattr(args, "target_lang", None)})
        or _target_lang(prof if prof else {})
        or str(target_lang or "").lower()
    )
    asset_lang, lang_cfg = resolve_language_assets(requested_asset_lang)
    if lang_cfg:
        print(
            f"[lqe_io] target language attributes: "
            f"target_languages/{asset_lang}/attributes.json"
        )

    lang_notes_path = ""
    if asset_lang:
        notes_path = _LANG_DIR / asset_lang / "eval_notes.md"
        if notes_path.exists():
            staged_file = write_dir / "lang_notes.md"
            published_file = final_dir / "lang_notes.md"
            staged_file.write_text(
                notes_path.read_text(encoding="utf-8"), encoding="utf-8"
            )
            lang_notes_path = str(published_file)
            print(f"[lqe_io] language eval notes → {published_file}")

    background_path = ""
    if prof and (prof.get("background") or "").strip():
        staged_file = write_dir / "background.md"
        published_file = final_dir / "background.md"
        staged_file.write_text(
            "# 项目背景\n\n" + prof["background"].strip() + "\n",
            encoding="utf-8",
        )
        background_path = str(published_file)
        print(f"[lqe_io] project background → {published_file}")

    basis = (
        getattr(args, "wordcount_basis", None)
        or (prof.get("wordcount_basis") if prof else None)
        or lang_cfg.get("wordcount_basis")
        or "target-words"
    )
    if lang_cfg.get("word_delim") == "none" and basis == "target-words":
        print(
            "[warn] target language has no word delimiter — 'target-words' basis will "
            "undercount severely; use source-chars",
            file=sys.stderr,
        )
    if basis == "source-chars":
        wordcount = sum(
            len(_RE_CJK.findall(segment.get("source_plain", segment["source"])))
            + len(
                re.findall(
                    r"[A-Za-z0-9]+",
                    segment.get("source_plain", segment["source"]),
                )
            )
            for segment in segments
        )
    else:
        wordcount = sum(
            len(segment.get("target_plain", segment["target"]).split())
            for segment in segments
        )

    checks_path = confirmed_rules_path = ""
    if prof:
        active_checks = _active_profile_asset_path(prof, "checks")
        active_confirmed = _active_profile_asset_path(prof, "confirmed_rules")
        checks = Path(active_checks) if active_checks else None
        checks_path = str(checks) if checks is not None and checks.exists() else ""
        confirmed = Path(active_confirmed) if active_confirmed else None
        common_confirmed = (
            confirmed.parent.parent / "common" / "confirmed_rules_common.md"
            if confirmed is not None
            else Path(prof["_dir"]).parent / "common" / "confirmed_rules_common.md"
        )
        parts = []
        if common_confirmed.exists():
            parts.append(
                f"<!-- ===== 共通确认规则（游戏级）: {common_confirmed.name} ===== -->\n"
                + common_confirmed.read_text(encoding="utf-8")
            )
        if confirmed is not None and confirmed.exists():
            parts.append(
                f"<!-- ===== 语言专有确认规则: {confirmed} ===== -->\n"
                + confirmed.read_text(encoding="utf-8")
            )
        if parts:
            staged_file = write_dir / "confirmed_rules.md"
            published_file = final_dir / "confirmed_rules.md"
            staged_file.write_text("\n\n".join(parts), encoding="utf-8")
            confirmed_rules_path = str(published_file)
            print(
                f"[lqe_io] confirmed rules: "
                f"{'共通+' if common_confirmed.exists() else ''}语言专有 → "
                f"{published_file}"
            )

    profile_policy = dict((prof or {}).get("scoring_policy", {}))
    for key in (
        "scorecard_profile",
        "severity_scale",
        "critical_gate",
        "repeat_dedup",
    ):
        if prof and key in prof and key not in profile_policy:
            profile_policy[key] = prof[key]
    if prof and "threshold" in prof and "threshold" not in profile_policy:
        profile_policy["threshold"] = prof["threshold"]
    scoring_policy = resolve_scoring_policy({}, profile_policy)

    normalized_profile = prof.get("_normalized_profile") if prof else None
    asset_inspection = prof.get("_asset_inspection") if prof else None
    capability_resolution = prof.get("_capability_resolution") if prof else None
    if normalized_profile is None:
        runtime_profile = {
            "name": "runtime/unprofiled",
            "language_pair": (
                f"{source_lang}-{target_lang}"
                if source_lang and target_lang
                else "und-und"
            ),
            "source_lang": source_lang or "und",
            "target_lang": target_lang or "und",
            "wordcount_basis": basis,
        }
        normalized_profile = normalize_profile(runtime_profile)
        asset_inspection = inspect_project_assets(
            normalized_profile,
            profile_dir=write_dir,
            allow_outside_root=True,
            strict_required=True,
        )
        capability_resolution = resolve_capabilities(
            normalized_profile,
            asset_statuses=asset_statuses(asset_inspection["snapshot"]),
            provider_registry=trusted_provider_registry(),
        )
    module_context_views, shadow_module_context_views = (
        _module_context_views_for_profile(prof)
        if prof
        else ({}, None)
    )
    project_asset_snapshot = (
        deepcopy(asset_inspection.get("snapshot"))
        if isinstance(asset_inspection, dict)
        else None
    )
    capability_resolution_path = ""
    project_asset_snapshot_path = ""
    if isinstance(capability_resolution, dict):
        staged_file = write_dir / "capability_resolution.json"
        published_file = final_dir / "capability_resolution.json"
        staged_file.write_text(
            json.dumps(capability_resolution, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        capability_resolution_path = str(published_file)
    if isinstance(project_asset_snapshot, dict):
        staged_file = write_dir / "project_asset_snapshot.json"
        published_file = final_dir / "project_asset_snapshot.json"
        staged_file.write_text(
            json.dumps(project_asset_snapshot, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        project_asset_snapshot_path = str(published_file)

    project_assets_path = ""
    copied_project_assets = {}
    if isinstance(asset_inspection, dict):
        staged_root = write_dir / "project_assets"
        copied = copy_project_assets(asset_inspection, staged_root)
        if copied:
            project_assets_path = str(final_dir / "project_assets")
            copied_project_assets = {
                asset_id: str(
                    final_dir / "project_assets" / asset_id / path.name
                )
                for asset_id, path in copied.items()
            }

    return {
        "job_runtime_contract_version": JOB_RUNTIME_CONTRACT_VERSION,
        "created_at": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "context_contract_version": 1,
        "input_guard_version": 1,
        "aipe_url": None,
        "check_scope": check_scope,
        "review_policy": review_policy,
        "project": prof.get("name", "") if prof else "",
        "language_pair": prof.get("language_pair", "") if prof else (
            f"{source_lang}-{target_lang}" if source_lang and target_lang else ""
        ),
        "source_lang": source_lang,
        "target_lang": target_lang,
        "lang_notes_path": lang_notes_path,
        "background_path": background_path,
        "checks_path": checks_path,
        "confirmed_rules_path": confirmed_rules_path,
        "threshold": scoring_policy["threshold"],
        "scoring_policy": scoring_policy,
        "sg_path": sg_path,
        "terms_path": terms_path,
        "terminology": [],
        "style_guide": "",
        "wordcount": wordcount,
        "wordcount_basis": basis,
        "iteration": 0,
        "profile_contract_version": (
            normalized_profile.get("profile_contract_version")
            if isinstance(normalized_profile, dict)
            else 1
        ),
        "profile_digest": (
            normalized_profile.get("source_profile_digest")
            if isinstance(normalized_profile, dict)
            else None
        ),
        "profile_overlay_digest": (
            file_sha256(Path(prof["_overlay_path"]))
            if prof and prof.get("_overlay_path")
            else None
        ),
        "context_pipeline": deepcopy(
            (normalized_profile or {}).get("context_pipeline", {"mode": "off"})
        ),
        "module_context_views": module_context_views,
        **(
            {"shadow_module_context_views": shadow_module_context_views}
            if shadow_module_context_views is not None
            else {}
        ),
        "normalized_capabilities": deepcopy(
            (normalized_profile or {}).get("capabilities", {})
        ),
        "capability_descriptors": deepcopy(
            (normalized_profile or {}).get("capability_descriptors", {})
        ),
        "capability_resolution": deepcopy(capability_resolution),
        "capability_resolution_digest": (
            capability_resolution.get("digest")
            if isinstance(capability_resolution, dict)
            else None
        ),
        "capability_resolution_path": capability_resolution_path,
        "project_asset_snapshot": project_asset_snapshot,
        "project_asset_snapshot_digest": (
            project_asset_snapshot.get("digest")
            if isinstance(project_asset_snapshot, dict)
            else None
        ),
        "project_asset_snapshot_path": project_asset_snapshot_path,
        "project_assets_path": project_assets_path,
        "project_asset_paths": copied_project_assets,
    }


def _staged_asset_replacements(
    staging_dir: Path,
    job_dir: Path,
    *,
    exclude: set[Path] | None = None,
) -> list[tuple[Path, Path]]:
    excluded = {path.resolve() for path in (exclude or set())}
    replacements = []
    for source in sorted(staging_dir.rglob("*")):
        if not source.is_file() or source.resolve() in excluded:
            continue
        destination = job_dir / source.relative_to(staging_dir)
        destination.parent.mkdir(parents=True, exist_ok=True)
        replacements.append((source, destination))
    return replacements


def _language_values_match(first: str, second: str) -> bool:
    return language_tags_match(first, second) or language_tags_match(second, first)


def _validate_sdlxliff_languages(result, args, prof: dict | None) -> tuple[str, str]:
    declarations = result.manifest.get("languages", [])
    if not declarations:
        raise SDLXLIFFImportError("SDLXLIFF input has no language declarations")
    normalized = []
    for index, declaration in enumerate(declarations):
        source = normalize_language_tag(declaration.get("source_language"))
        target = normalize_language_tag(declaration.get("target_language"))
        if not source or not target:
            raise SDLXLIFFImportError(
                f"language declaration {index} must include source-language and target-language"
            )
        normalized.append((source, target))
    expected_source, expected_target = normalized[0]
    for index, (source, target) in enumerate(normalized[1:], start=1):
        if source != expected_source or target != expected_target:
            raise SDLXLIFFImportError(
                "conflicting SDLXLIFF language declarations: "
                f"declaration 0={expected_source}->{expected_target}, "
                f"declaration {index}={source}->{target}"
            )

    profile_source = normalize_language_tag(prof.get("source_lang")) if prof else ""
    profile_target = normalize_language_tag(prof.get("target_lang")) if prof else ""
    cli_source = normalize_language_tag(getattr(args, "source_lang", None))
    cli_target = normalize_language_tag(getattr(args, "target_lang", None))
    for label, profile_value, cli_value in (
        ("source", profile_source, cli_source),
        ("target", profile_target, cli_target),
    ):
        if profile_value and cli_value and not _language_values_match(
            profile_value, cli_value
        ):
            raise SDLXLIFFImportError(
                f"profile and CLI {label} language conflict: "
                f"{profile_value!r} != {cli_value!r}"
            )

    for origin, source, target in (
        ("profile", profile_source, profile_target),
        ("CLI", cli_source, cli_target),
    ):
        if source and not language_tags_match(source, expected_source):
            raise SDLXLIFFImportError(
                f"{origin} source language {source!r} does not match "
                f"declared source language {expected_source!r}"
            )
        if target and not language_tags_match(target, expected_target):
            raise SDLXLIFFImportError(
                f"{origin} target language {target!r} does not match "
                f"declared target language {expected_target!r}"
            )
    return expected_source, expected_target


def _publish_sdlxliff_job(
    state_path: Path,
    *,
    manifest: dict,
    tm_candidates: dict,
    scope: dict,
    state: dict,
    staged_assets: dict[Path, Path] | None = None,
) -> None:
    job_dir = state_path.parent
    reserved_names = {
        "source_manifest.json",
        "tm_candidates.json",
        "scope.json",
    }
    if state_path.name.casefold() in reserved_names:
        raise ValueError(
            f"SDLXLIFF state path conflicts with reserved helper artifact: {state_path}"
        )
    paths = {
        "manifest": job_dir / "source_manifest.json",
        "candidates": job_dir / "tm_candidates.json",
        "scope": job_dir / "scope.json",
        "state": state_path,
    }
    assets = {
        Path(destination): Path(staged)
        for destination, staged in (staged_assets or {}).items()
    }
    destinations = [*paths.values(), *assets]
    canonical_destinations: dict[str, Path] = {}
    for destination in destinations:
        canonical = str(destination.resolve()).casefold()
        previous = canonical_destinations.get(canonical)
        if previous is not None:
            raise ValueError(
                "SDLXLIFF job artifacts resolve to the same path: "
                f"{previous}, {destination}"
            )
        canonical_destinations[canonical] = destination

    missing_staged = [path for path in assets.values() if not path.is_file()]
    if missing_staged:
        raise FileNotFoundError(
            "staged SDLXLIFF job asset is missing: "
            + ", ".join(str(path) for path in missing_staged)
        )
    existing = [path for path in destinations if os.path.lexists(path)]
    if existing:
        raise FileExistsError(
            "SDLXLIFF job artifact already exists; use a new job directory: "
            + ", ".join(str(path) for path in existing)
        )

    values = {
        "manifest": manifest,
        "candidates": tm_candidates,
        "scope": scope,
        "state": state,
    }
    serialized = {
        key: json.dumps(value, ensure_ascii=False, indent=2)
        for key, value in values.items()
    }
    job_dir.mkdir(parents=True, exist_ok=True)
    staged: dict[str, Path] = {}

    try:
        for key, payload in serialized.items():
            with tempfile.NamedTemporaryFile(
                mode="w",
                encoding="utf-8",
                dir=job_dir,
                prefix=f".{paths[key].name}.",
                suffix=".tmp",
                delete=False,
            ) as handle:
                staged[key] = Path(handle.name)
                handle.write(payload)
        replacements = [
            (source, destination)
            for destination, source in sorted(
                assets.items(), key=lambda item: str(item[0])
            )
        ]
        replacements.extend(
            (staged[key], paths[key])
            for key in ("manifest", "candidates", "scope", "state")
        )
        try:
            publish_replacement_transaction(replacements, overwrite=False)
        except FileExistsError as exc:
            raise FileExistsError(
                "SDLXLIFF job artifact appeared during publication"
            ) from exc
    finally:
        for path in staged.values():
            path.unlink(missing_ok=True)


def _read_sdlxliff_job(
    args,
    prof: dict | None,
    check_scope: dict,
    review_policy: dict,
) -> None:
    state_path = Path(args.out)
    job_dir = state_path.parent
    helper_paths = (
        job_dir / "source_manifest.json",
        job_dir / "tm_candidates.json",
        job_dir / "scope.json",
    )
    if state_path.name.casefold() in {
        helper_path.name.casefold() for helper_path in helper_paths
    } or any(
        state_path.resolve() == helper_path.resolve() for helper_path in helper_paths
    ):
        raise ValueError(
            f"SDLXLIFF --out path conflicts with reserved helper artifact: {state_path}"
        )
    formal_paths = (
        state_path,
        *helper_paths,
    )
    existing = [path for path in formal_paths if path.exists()]
    if existing:
        raise FileExistsError(
            "SDLXLIFF job artifact already exists; use a new job directory: "
            + ", ".join(str(path) for path in existing)
        )

    raw_options = prof.get("sdlxliff", {}) if prof else {}
    options = validate_sdlxliff_options(
        raw_options,
        cli_protect_exact_tm=getattr(args, "protect_exact_tm", False),
    )
    result = read_sdlxliff(Path(args.input), options=options)
    source_lang, target_lang = _validate_sdlxliff_languages(result, args, prof)

    for segment in result.segments:
        metadata = segment["metadata"]["sdlxliff"]
        metadata["content_type"] = segment.get("content_type")
        segment["context_note"] = metadata.get("comment") or None
        source_ref = segment.get("source_ref") or {}
        business_parts = [
            source_ref.get("relative_path"),
            source_ref.get("tu_id", source_ref.get("tu_index")),
            source_ref.get("sdl_segment_id", source_ref.get("segment_index")),
        ]
        identity = build_segment_identity(
            business_key="::".join(str(part) for part in business_parts),
            input_digest=source_digest(segment.get("source", "")),
            container=str(source_ref.get("relative_path") or "sdlxliff"),
            row_index=segment["id"],
        )
        segment.update(identity)
        segment["source_digest"] = source_digest(segment.get("source", ""))
        segment["input_status"] = "ready"
        segment["source_provenance"] = {
            "adapter": "sdlxliff@1",
            "source_ref": deepcopy(source_ref),
        }
        segment["iter"] = 0

    context_registry = _registry_for_profile(prof)
    resolution = prof.get("_capability_resolution") if prof else None
    shadow_registry = (
        _registry_for_profile(prof, include_shadow=True)
        if resolution
        and resolution.get("context_pipeline_mode") == "shadow"
        else None
    )
    for segment in result.segments:
        segment["context"] = _legacy_segment_context(
            segment,
            context_registry,
            source_provenance=segment["source_provenance"],
        )
        segment["context_provenance"] = deepcopy(
            segment["context"].get("provenance", {})
        )
        segment["context_status"] = segment["context"].get("status", "ready")
        segment["context_missing_required"] = segment["context"].get(
            "missing_required", []
        )
        segment["resolved_constraints"] = []
        if shadow_registry is not None:
            segment["shadow_context"] = _legacy_segment_context(
                segment,
                shadow_registry,
                source_provenance=segment["source_provenance"],
            )

    job_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(
        dir=job_dir, prefix=".read-assets."
    ) as asset_staging_dir:
        staging_dir = Path(asset_staging_dir)
        common = _prepare_read_assets(
            args,
            prof,
            check_scope,
            review_policy,
            result.segments,
            source_lang,
            target_lang,
            asset_dir=staging_dir,
            publish_dir=job_dir,
        )
        staged_project_asset_paths = {
            asset_id: str(
                staging_dir
                / "project_assets"
                / asset_id
                / Path(published_path).name
            )
            for asset_id, published_path in common.get(
                "project_asset_paths", {}
            ).items()
        }
        result.segments = _apply_runtime_project_context(
            result.segments,
            common,
            prof,
            context_registry,
            runtime_asset_paths=staged_project_asset_paths,
            phase="project_overrides",
        )
        (
            result.segments,
            context_overrides_audit,
            context_gap_report,
            context_overrides_guard,
        ) = _apply_explicit_context_overrides(
            args,
            result.segments,
            common,
            context_registry,
            shadow_registry=shadow_registry,
            staging_dir=staging_dir,
            job_dir=job_dir,
            staged_project_asset_paths=staged_project_asset_paths,
        )
        result.segments = _apply_runtime_project_context(
            result.segments,
            common,
            prof,
            context_registry,
            runtime_asset_paths=staged_project_asset_paths,
            phase="rules",
        )
        shadow_context_artifact = None
        if shadow_registry is not None:
            shadow_context_artifact = build_shadow_context_artifact(
                common,
                result.segments,
                shadow_registry,
            )
            staged_shadow = staging_dir / "shadow_context" / "context.json"
            staged_shadow.parent.mkdir(parents=True, exist_ok=True)
            staged_shadow.write_text(
                json.dumps(
                    shadow_context_artifact,
                    ensure_ascii=False,
                    indent=2,
                ),
                encoding="utf-8",
            )
        for segment in result.segments:
            segment["segment_revision_digest"] = _segment_revision(
                segment, common
            )
            segment["module_review_equivalence_keys"] = {
                module: module_review_equivalence_key(
                    segment, module, context_registry
                )
                for module in (
                    "terminology",
                    "precheck_review",
                    "accuracy",
                    "grammar",
                    "naturalness",
                    "suggestions",
                )
            }
        manifest_path = job_dir / "source_manifest.json"
        candidates_path = job_dir / "tm_candidates.json"
        candidates = {
            "candidate_ids": list(result.tm_candidates.get("candidate_ids", [])),
            "segments": [
                {
                    "id": item["segment_id"],
                    "evidence": item.get("evidence", {}),
                    "source_ref": item.get("source_ref", {}),
                }
                for item in result.tm_candidates.get("segments", [])
            ],
        }
        state = {
            "artifact_contract_version": 1,
            "input_format": "sdlxliff",
            "input_path": str(Path(args.input).resolve()),
            "input_paths": result.input_paths,
            "source_manifest_path": str(manifest_path),
            "tm_candidates_path": str(candidates_path),
            "source_col": "原文",
            "target_col": "译文",
            "headers": result.headers,
            "rows_raw": result.rows_raw,
            "text_type_markers": [],
            "resolved_context_descriptors": context_registry,
            "shadow_context_descriptors": shadow_registry,
            **(
                {
                    "shadow_context_path": str(
                        job_dir / "shadow_context" / "context.json"
                    ),
                    "shadow_context_digest": shadow_context_artifact[
                        "artifact_digest"
                    ],
                }
                if shadow_context_artifact is not None
                else {}
            ),
            **common,
            "segments": result.segments,
        }
        state["input_guard"] = input_guard_summary(result.segments)
        state["review_wordcount"] = state["wordcount"]
        if context_overrides_audit is not None:
            result.manifest["context_overrides"] = deepcopy(
                context_overrides_audit
            )
            result.manifest["context_gap_report"] = {
                "path": common["context_gap_report_path"],
                "digest": context_gap_report["report_digest"],
                "summary": deepcopy(context_gap_report["summary"]),
            }
        staged_assets = {
            destination: source
            for source, destination in _staged_asset_replacements(
                staging_dir, job_dir
            )
        }
        if context_overrides_guard is not None:
            override_path, expected_digest = context_overrides_guard
            if file_sha256(override_path) != expected_digest:
                raise ValueError(
                    f"context overrides changed while input was being read: {override_path}"
                )
        _publish_sdlxliff_job(
            state_path,
            manifest=result.manifest,
            tm_candidates=candidates,
            scope=check_scope,
            state=state,
            staged_assets=staged_assets,
        )
    print(
        f"[lqe_io] {len(result.segments)} segments → {args.out}  "
        f"wordcount={state['wordcount']}"
    )


def _cmd_read_locked(args):
    check_scope = build_check_scope(getattr(args, "no_terminology", False))
    review_policy = build_review_policy(
        getattr(args, "review_mode", "optimized")
    )
    out_path = Path(args.out)
    job_dir = out_path.parent
    scope_path = job_dir / "scope.json"
    if (
        out_path.name.casefold() == "scope.json"
        or out_path.resolve() == scope_path.resolve()
    ):
        print(
            f"[ERROR] --out path conflicts with reserved scope artifact: {scope_path}",
            file=sys.stderr,
        )
        sys.exit(2)
    if _paths_alias(out_path, Path(args.input)):
        print(
            f"[ERROR] --out path conflicts with --input: {args.input}",
            file=sys.stderr,
        )
        sys.exit(2)
    generated_asset_names = {
        "sg.txt",
        "terms.json",
        "lang_notes.md",
        "background.md",
        "confirmed_rules.md",
        "capability_resolution.json",
        "project_asset_snapshot.json",
        "tabular_source_manifest.json",
        "context_overrides.json",
        "context_gap_report.json",
    }
    if out_path.name.casefold() in generated_asset_names:
        print(
            f"[ERROR] --out path conflicts with generated asset: {out_path.name}",
            file=sys.stderr,
        )
        sys.exit(2)
    try:
        input_format = detect_input_format(
            Path(args.input), getattr(args, "input_format", "auto")
        )
    except ValueError as exc:
        print(f"[ERROR] {exc}", file=sys.stderr)
        sys.exit(1)
    if input_format == "sdlxliff" and out_path.name.casefold() in {
        "source_manifest.json",
        "tm_candidates.json",
        "scope.json",
    }:
        print(
            f"[ERROR] SDLXLIFF --out path conflicts with reserved helper artifact: {out_path}",
            file=sys.stderr,
        )
        sys.exit(2)
    tabular_only_options = {
        "--sheet": getattr(args, "sheet", None),
        "--key-col": getattr(args, "key_col", None),
        "--context-col": getattr(args, "context_cols", None),
        "--pivot-sheet": getattr(args, "pivot_sheet", None),
        "--pivot-key-col": getattr(args, "pivot_key_col", None),
        "--pivot-compare": getattr(args, "pivot_compare", None),
        "--pivot-authority": getattr(args, "pivot_authority", None),
        "--target-source-digest-col": getattr(
            args, "target_source_digest_col", None
        ),
    }
    if input_format == "sdlxliff" and any(
        value not in (None, [], "") for value in tabular_only_options.values()
    ):
        invalid = [
            flag for flag, value in tabular_only_options.items()
            if value not in (None, [], "")
        ]
        print(
            "[ERROR] tabular-only options are not valid for SDLXLIFF: "
            + ", ".join(invalid),
            file=sys.stderr,
        )
        sys.exit(2)
    prof = _load_project(args.project) if getattr(args, "project", None) else None
    if prof:
        if getattr(args, "profile_overlay", None):
            prof = _load_profile_overlay(args.profile_overlay, prof)
        _validate_project_profile(prof)
        active_style = _active_profile_asset_path(prof, "style_guide")
        active_terms = _active_profile_asset_path(prof, "terminology")
        if not args.style_guide and active_style:
            args.style_guide = active_style
        if active_terms:
            if not check_scope["terminology_enabled"]:
                print("[lqe_io] profile terminology overridden by --no-terminology")
            elif not args.terminology:
                args.terminology = active_terms
        print(f"[lqe_io] project: {prof.get('name', '?')} ({prof.get('language_pair', '?')})")

    protected_inputs = {
        "--input": Path(args.input),
        **_profile_reference_paths(prof),
    }
    if args.style_guide:
        protected_inputs["--style-guide"] = Path(args.style_guide)
    if args.terminology:
        protected_inputs["--terminology"] = Path(args.terminology)
    if getattr(args, "context_overrides", None):
        protected_inputs["--context-overrides"] = Path(args.context_overrides)
    planned_outputs = {
        "state": out_path,
        "scope": scope_path,
        "style guide copy": job_dir / "sg.txt",
        "terminology copy": job_dir / "terms.json",
        "language notes copy": job_dir / "lang_notes.md",
        "background copy": job_dir / "background.md",
        "confirmed rules copy": job_dir / "confirmed_rules.md",
        "capability resolution": job_dir / "capability_resolution.json",
        "project asset snapshot": job_dir / "project_asset_snapshot.json",
    }
    if getattr(args, "context_overrides", None):
        planned_outputs.update(
            {
                "context overrides copy": job_dir / "context_overrides.json",
                "context gap report": job_dir / "context_gap_report.json",
            }
        )
    if input_format == "tabular":
        planned_outputs["tabular source manifest"] = (
            job_dir / "tabular_source_manifest.json"
        )
    if input_format == "sdlxliff":
        planned_outputs.update(
            {
                "source manifest": job_dir / "source_manifest.json",
                "TM candidates": job_dir / "tm_candidates.json",
            }
        )
    try:
        validate_artifact_paths(
            planned_outputs,
            protected_inputs,
            context="read",
        )
    except ValueError as exc:
        print(f"[ERROR] {exc}", file=sys.stderr)
        sys.exit(2)

    if input_format == "sdlxliff":
        try:
            _read_sdlxliff_job(args, prof, check_scope, review_policy)
        except (OSError, ValueError) as exc:
            print(f"[ERROR] {exc}", file=sys.stderr)
            sys.exit(1)
        return

    if getattr(args, "protect_exact_tm", False):
        print(
            "[ERROR] --protect-exact-tm is only valid for SDLXLIFF input",
            file=sys.stderr,
        )
        sys.exit(2)
    if args.source_col is None or args.target_col is None:
        print(
            "[ERROR] tabular input requires --source-col and --target-col",
            file=sys.stderr,
        )
        sys.exit(2)

    no_header = getattr(args, "no_header", False)
    input_path = Path(args.input)
    input_sha256 = file_sha256(input_path)
    try:
        headers, data_rows, container, source_manifest = _read_tabular_source(
            input_path,
            sheet_name=getattr(args, "sheet", None),
            no_header=no_header,
        )
        if not data_rows:
            raise ValueError("no data rows found")
        widths = {len(row) for row in data_rows if any(_text(cell) for cell in row)}
        if len(widths) > 1:
            raise ValueError(f"tabular input has inconsistent row widths: {sorted(widths)}")
        si = _resolve_input_column(
            headers, args.source_col, no_header=no_header, label="--source-col"
        )
        ti = _resolve_input_column(
            headers, args.target_col, no_header=no_header, label="--target-col"
        )
        key_index = (
            _resolve_input_column(
                headers,
                args.key_col,
                no_header=no_header,
                label="--key-col",
            )
            if getattr(args, "key_col", None) is not None
            else None
        )
        digest_index = (
            _resolve_input_column(
                headers,
                args.target_source_digest_col,
                no_header=no_header,
                label="--target-source-digest-col",
            )
            if getattr(args, "target_source_digest_col", None) is not None
            else None
        )
        normalized_profile = prof.get("_normalized_profile") if prof else None
        text_type_marker_rules = deepcopy(
            (normalized_profile or {})
            .get("tabular", {})
            .get("text_type_marker_rules", [])
        )
        resolution = prof.get("_capability_resolution") if prof else None
        registry = _registry_for_profile(prof)
        shadow_registry = (
            _registry_for_profile(prof, include_shadow=True)
            if resolution
            and resolution.get("context_pipeline_mode") == "shadow"
            else None
        )
        context_cli_specs = _context_cli_specs(args)
        shadow_context_columns = (
            resolve_context_columns(
                headers,
                shadow_registry,
                cli_columns=context_cli_specs,
                profile=normalized_profile,
                no_header=no_header,
            )
            if shadow_registry is not None
            else None
        )
        formal_context_cli_specs = (
            _context_cli_specs_for_registry(
                context_cli_specs,
                source_registry=shadow_registry,
                target_registry=registry,
            )
            if shadow_registry is not None
            else context_cli_specs
        )
        resolved_context_columns = resolve_context_columns(
            headers,
            registry,
            cli_columns=formal_context_cli_specs,
            profile=normalized_profile,
            no_header=no_header,
        )
        pivot_guard_config = _prepare_pivot_guard(
            args,
            input_path=input_path,
            main_container=container,
            main_headers=headers,
            source_index=si,
            target_index=ti,
            key_index=key_index,
            resolved_context_columns=resolved_context_columns,
            registry=registry,
            no_header=no_header,
        )
    except (ContextContractError, OSError, ValueError, XLSImportError) as exc:
        print(f"[ERROR] {exc}", file=sys.stderr)
        sys.exit(1)

    # R3: 自动识别 max-length 列（UI 字段宽度上限），用于逐元素截断检查
    mi = None
    for idx, h in enumerate(headers):
        if h is not None and str(h).strip().lower() in _MAXLEN_KEYS:
            mi = idx
            break
    if mi is not None:
        print(f"[lqe_io] max-length column detected: '{headers[mi]}' (col {mi})")

    gi = None
    if getattr(args, "group_col", None):
        g = args.group_col
        gi = int(g) if str(g).isdigit() else (headers.index(g) if g in headers else None)
        if gi is None:
            print(f"[warn] group column '{g}' not found; grouping disabled", file=sys.stderr)
        else:
            print(f"[lqe_io] group column: '{headers[gi]}' (col {gi})")

    segments, rows_raw, text_type_markers = [], [], []
    text_type_context = None
    for i, row in enumerate(data_rows):
        if any(_text(c) for c in row):
            src = _text(_cell(row, si))
            tgt = _text(_cell(row, ti))
            context_state = extract_segment_context(
                row,
                resolved_context_columns,
                registry,
                profile=normalized_profile,
                source_provenance=_tabular_source_provenance(
                    source_manifest, row_index=i
                ),
            )
            shadow_context_state = (
                extract_segment_context(
                    row,
                    shadow_context_columns,
                    shadow_registry,
                    profile=normalized_profile,
                    source_provenance=_tabular_source_provenance(
                        source_manifest, row_index=i
                    ),
                )
                if shadow_registry is not None
                else None
            )
            row_content_type = _text(
                context_state.get("core", {}).get("content_type")
            )
            try:
                marker = _extract_text_type_marker(
                    src,
                    tgt,
                    row_content_type,
                    rules=text_type_marker_rules,
                )
            except ValueError as exc:
                print(f"[ERROR] {exc}", file=sys.stderr)
                sys.exit(1)
            if marker:
                text_type_context = marker
                text_type_markers.append({
                    "row_index": i,
                    "source": src,
                    "target": tgt,
                    "content_type": row_content_type or None,
                })
                continue
            seg_id = len(segments)
            identity = build_segment_identity(
                business_key=_cell(row, key_index),
                input_digest=input_sha256,
                container=container,
                row_index=i,
            )
            segment = {
                "id": seg_id,
                "row_index": i,
                **identity,
                "source": src,
                "target": tgt,
                "corrected": None,
                "content_type": row_content_type or None,
                "text_type_context": text_type_context,
                "context": context_state,
                "context_provenance": deepcopy(
                    context_state.get("provenance", {})
                ),
                "context_status": context_state.get("status", "ready"),
                "context_missing_required": context_state.get(
                    "missing_required", []
                ),
                "source_digest": source_digest(src),
                "input_status": "ready",
                "source_provenance": _tabular_source_provenance(
                    source_manifest, row_index=i
                ),
                "resolved_constraints": [],
                "max_len": _parse_maxlen(row[mi]) if mi is not None and mi < len(row) else None,
                "group": (str(row[gi]).strip() if gi is not None and gi < len(row) and row[gi] is not None and str(row[gi]).strip() else None),
                "iter": 0,
            }
            if shadow_context_state is not None:
                segment["shadow_context"] = shadow_context_state
            if digest_index is not None:
                apply_target_source_digest_guard(
                    segment, _cell(row, digest_index)
                )
            else:
                segment["input_warnings"] = [
                    {"code": "UNVERIFIED_TARGET_PROVENANCE"}
                ]
            segments.append(segment)
            rows_raw.append(list(row))

    if not segments:
        print("[ERROR] No data rows found.", file=sys.stderr)
        sys.exit(1)
    try:
        ensure_unique_business_keys(segments)
        pivot_guard = (
            _apply_pivot_guard(segments, rows_raw, pivot_guard_config)
            if pivot_guard_config is not None
            else None
        )
    except ValueError as exc:
        print(f"[ERROR] {exc}", file=sys.stderr)
        sys.exit(1)

    source_lang = _source_lang({"source_lang": getattr(args, "source_lang", None)}) \
        or _source_lang(prof if prof else {})
    lang = _target_lang({"target_lang": getattr(args, "target_lang", None)}) \
        or _target_lang(prof if prof else {})
    try:
        job_dir.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(
            dir=job_dir, prefix=".read-assets."
        ) as asset_staging_dir:
            staging_dir = Path(asset_staging_dir)
            common = _prepare_read_assets(
                args,
                prof,
                check_scope,
                review_policy,
                segments,
                source_lang,
                lang,
                asset_dir=staging_dir,
                publish_dir=job_dir,
            )
            staged_project_asset_paths = {
                asset_id: str(
                    staging_dir
                    / "project_assets"
                    / asset_id
                    / Path(published_path).name
                )
                for asset_id, published_path in common.get(
                    "project_asset_paths", {}
                ).items()
            }
            segments = _apply_runtime_project_context(
                segments,
                common,
                prof,
                registry,
                runtime_asset_paths=staged_project_asset_paths,
                phase="project_overrides",
            )
            (
                segments,
                context_overrides_audit,
                context_gap_report,
                context_overrides_guard,
            ) = _apply_explicit_context_overrides(
                args,
                segments,
                common,
                registry,
                shadow_registry=shadow_registry,
                staging_dir=staging_dir,
                job_dir=job_dir,
                staged_project_asset_paths=staged_project_asset_paths,
            )
            segments = _apply_runtime_project_context(
                segments,
                common,
                prof,
                registry,
                runtime_asset_paths=staged_project_asset_paths,
                phase="rules",
            )

            for segment in segments:
                segment["segment_revision_digest"] = _segment_revision(
                    segment, common
                )
                segment["module_review_equivalence_keys"] = {
                    module: module_review_equivalence_key(
                        segment, module, registry
                    )
                    for module in (
                        "terminology",
                        "precheck_review",
                        "accuracy",
                        "grammar",
                        "naturalness",
                        "suggestions",
                    )
                }
            total_wordcount = common["wordcount"]
            review_segments = [
                segment
                for segment in segments
                if segment.get("input_status") != "blocked"
            ]
            if common["wordcount_basis"] == "source-chars":
                review_wordcount = sum(
                    len(_RE_CJK.findall(segment.get("source", "")))
                    + len(re.findall(r"[A-Za-z0-9]+", segment.get("source", "")))
                    for segment in review_segments
                )
            else:
                review_wordcount = sum(
                    len(segment.get("target", "").split())
                    for segment in review_segments
                )
            common["wordcount"] = total_wordcount
            common["review_wordcount"] = review_wordcount
            common["review_wordcount_basis_digest"] = source_digest(
                json.dumps(
                    {
                        "wordcount": total_wordcount,
                        "review_wordcount": review_wordcount,
                        "blocked_ids": [
                            segment["id"]
                            for segment in segments
                            if segment.get("input_status") == "blocked"
                        ],
                    },
                    sort_keys=True,
                    separators=(",", ":"),
                )
            )
            shadow_context_artifact = None
            if shadow_registry is not None:
                shadow_context_artifact = build_shadow_context_artifact(
                    common,
                    segments,
                    shadow_registry,
                )
                staged_shadow = staging_dir / "shadow_context" / "context.json"
                staged_shadow.parent.mkdir(parents=True, exist_ok=True)
                staged_shadow.write_text(
                    json.dumps(
                        shadow_context_artifact,
                        ensure_ascii=False,
                        indent=2,
                    ),
                    encoding="utf-8",
                )
            tabular_manifest_path = job_dir / "tabular_source_manifest.json"
            staged_manifest = staging_dir / tabular_manifest_path.name
            source_manifest.update(
                {
                    "source_col": args.source_col,
                    "target_col": args.target_col,
                    "key_col": getattr(args, "key_col", None),
                    "group_col": getattr(args, "group_col", None),
                    "target_source_digest_col": getattr(
                        args, "target_source_digest_col", None
                    ),
                    "context_columns": resolved_context_columns,
                    "shadow_context_columns": shadow_context_columns,
                    "text_type_marker_rules": text_type_marker_rules,
                    "segments": len(segments),
                    **(
                        {"pivot_guard": pivot_guard}
                        if pivot_guard is not None
                        else {}
                    ),
                    **(
                        {
                            "context_overrides": deepcopy(
                                context_overrides_audit
                            ),
                            "context_gap_report": {
                                "path": common["context_gap_report_path"],
                                "digest": context_gap_report["report_digest"],
                                "summary": deepcopy(
                                    context_gap_report["summary"]
                                ),
                            },
                        }
                        if context_overrides_audit is not None
                        else {}
                    ),
                }
            )
            staged_manifest.write_text(
                json.dumps(source_manifest, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )

            state = {
                "artifact_contract_version": 1,
                "input_format": "tabular",
                "input_path": str(Path(args.input).resolve()),
                "input_sha256": input_sha256,
                "sheet_name": container,
                "tabular_source_manifest_path": str(tabular_manifest_path),
                "no_header": bool(no_header),
                "source_col": args.source_col,
                "target_col": args.target_col,
                "headers": headers,
                "rows_raw": rows_raw,
                "text_type_markers": text_type_markers,
                "text_type_marker_rules": text_type_marker_rules,
                "context_cols": resolved_context_columns,
                "resolved_context_descriptors": registry,
                "shadow_context_descriptors": shadow_registry,
                **(
                    {
                        "shadow_context_path": str(
                            job_dir / "shadow_context" / "context.json"
                        ),
                        "shadow_context_digest": shadow_context_artifact[
                            "artifact_digest"
                        ],
                    }
                    if shadow_context_artifact is not None
                    else {}
                ),
                **(
                    {"pivot_guard": pivot_guard}
                    if pivot_guard is not None
                    else {}
                ),
                **common,
                "input_guard": input_guard_summary(segments),
                "segments": segments,
            }
            staged_scope = staging_dir / "scope.json"
            staged_state = staging_dir / out_path.name
            staged_scope.write_text(
                json.dumps(check_scope, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            staged_state.write_text(
                json.dumps(state, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            asset_replacements = _staged_asset_replacements(
                staging_dir,
                job_dir,
                exclude={staged_scope, staged_state},
            )
            input_paths = [Path(args.input)]
            for configured in (args.style_guide, args.terminology):
                if configured:
                    input_paths.append(Path(configured))
            for _, destination in asset_replacements:
                for input_source in input_paths:
                    if _paths_alias(destination, input_source):
                        raise ValueError(
                            "generated job asset conflicts with input: "
                            f"{destination} == {input_source}"
                        )
            if file_sha256(input_path) != input_sha256:
                raise ValueError(
                    f"tabular input changed while it was being read: {input_path}"
                )
            if context_overrides_guard is not None:
                override_path, expected_digest = context_overrides_guard
                if file_sha256(override_path) != expected_digest:
                    raise ValueError(
                        "context overrides changed while input was being read: "
                        f"{override_path}"
                    )
            publish_replacement_transaction(
                [
                    *sorted(asset_replacements, key=lambda item: str(item[1])),
                    (staged_scope, scope_path),
                    (staged_state, out_path),
                ]
            )
    except (OSError, ValueError) as exc:
        print(f"[ERROR] {exc}", file=sys.stderr)
        sys.exit(1)
    print(
        f"[lqe_io] {len(segments)} segments → {args.out}  "
        f"wordcount={state['wordcount']}"
    )


def cmd_read(args):
    from lqe_split_contract import generation_lock

    state_path = Path(args.out)
    job_dir = state_path.parent
    read_lock_target = job_dir.parent / f"{job_dir.name}.lqe-read"
    try:
        with generation_lock(read_lock_target, exclusive=True):
            _cmd_read_locked(args)
    except (OSError, ValueError) as exc:
        raise SystemExit(f"[read] {exc}") from exc


def _reread_asset_path(old_job: Path, value: object, label: str) -> Path:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"historical state is missing {label}")
    path = Path(value)
    if not path.is_absolute():
        path = old_job / path
    if not path.is_file():
        raise ValueError(f"historical {label} is missing: {path}")
    return path


def _reread_segment_signature(segment: dict, input_format: str) -> tuple:
    if not isinstance(segment, dict):
        raise ValueError("historical state has a non-object segment")
    source = segment.get("source")
    target = segment.get("target")
    if not isinstance(source, str) or not isinstance(target, str):
        raise ValueError("historical segment source/target must be strings")
    if input_format == "tabular":
        row_index = segment.get("row_index", segment.get("id"))
        if type(row_index) is not int or row_index < 0:
            raise ValueError("historical tabular segment has no valid row_index")
        location = row_index
    else:
        source_ref = segment.get("source_ref")
        if not isinstance(source_ref, dict):
            raise ValueError("historical SDLXLIFF segment has no source_ref")
        location = (
            source_ref.get("relative_path"),
            source_ref.get("tu_id", source_ref.get("tu_index")),
            source_ref.get(
                "sdl_segment_id", source_ref.get("segment_index")
            ),
        )
        if any(value is None for value in location):
            raise ValueError("historical SDLXLIFF source_ref is incomplete")
    return location, source, target


def _assert_reread_coverage(
    old_segments: object,
    new_segments: object,
    input_format: str,
) -> None:
    if not isinstance(old_segments, list) or not isinstance(new_segments, list):
        raise ValueError("segment coverage must be represented by arrays")
    old_signatures = [
        _reread_segment_signature(segment, input_format)
        for segment in old_segments
    ]
    new_signatures = [
        _reread_segment_signature(segment, input_format)
        for segment in new_segments
    ]
    if old_signatures != new_signatures:
        raise ValueError(
            "original input no longer has the historical segment coverage"
        )


def _reread_context_specs(raw: object) -> list[str]:
    if raw in (None, {}):
        return []
    if not isinstance(raw, dict):
        raise ValueError("historical context column mapping is not reconstructable")
    specs = []
    for field_ref, mapping in sorted(raw.items()):
        if not isinstance(field_ref, str) or not isinstance(mapping, dict):
            raise ValueError(
                "historical context column mapping is not reconstructable"
            )
        column = mapping.get("column")
        if not isinstance(column, (str, int)) or isinstance(column, bool):
            raise ValueError(
                f"historical context column {field_ref!r} is not reconstructable"
            )
        specs.append(f"{field_ref}={column}")
    return specs


def _reread_tabular_options(
    old_job: Path,
    state: dict,
    input_path: Path,
) -> dict:
    expected_digest = state.get("input_sha256")
    if (
        not isinstance(expected_digest, str)
        or not re.fullmatch(r"[0-9a-fA-F]{64}", expected_digest)
    ):
        raise ValueError(
            "historical tabular state has no trustworthy input_sha256"
        )
    if file_sha256(input_path) != expected_digest.lower():
        raise ValueError("original input digest does not match historical state")

    no_header = state.get("no_header", False)
    if type(no_header) is not bool:
        raise ValueError("historical no_header option is invalid")
    source_col = state.get("source_col")
    target_col = state.get("target_col")
    for label, value in (("source_col", source_col), ("target_col", target_col)):
        if not isinstance(value, (str, int)) or isinstance(value, bool):
            raise ValueError(f"historical {label} is not reconstructable")

    sheet = state.get("sheet_name")
    if input_path.suffix.casefold() in {".csv", ".tsv"}:
        sheet = None
    elif not isinstance(sheet, str) or not sheet.strip():
        raise ValueError("historical sheet selection is not reconstructable")
    headers, rows, _, _ = _read_tabular_source(
        input_path,
        sheet_name=sheet,
        no_header=no_header,
    )
    source_index = _resolve_input_column(
        headers, source_col, no_header=no_header, label="--source-col"
    )
    target_index = _resolve_input_column(
        headers, target_col, no_header=no_header, label="--target-col"
    )
    old_segments = state.get("segments")
    if not isinstance(old_segments, list):
        raise ValueError("historical state has no segment array")
    for segment in old_segments:
        row_index, source, target = _reread_segment_signature(
            segment, "tabular"
        )
        if row_index >= len(rows):
            raise ValueError(
                f"historical segment row {row_index} is missing from original input"
            )
        row = rows[row_index]
        if _text(_cell(row, source_index)) != source or _text(
            _cell(row, target_index)
        ) != target:
            raise ValueError(
                f"historical segment row {row_index} no longer matches source/target"
            )

    manifest = None
    manifest_value = state.get("tabular_source_manifest_path")
    if manifest_value:
        manifest_path = _reread_asset_path(
            old_job, manifest_value, "tabular source manifest"
        )
        manifest = read_json(manifest_path)
        if not isinstance(manifest, dict):
            raise ValueError("historical tabular source manifest must be an object")
        if manifest.get("input_sha256") != expected_digest:
            raise ValueError(
                "historical tabular source manifest digest conflicts with state"
            )
        for field, expected in (
            ("source_col", source_col),
            ("target_col", target_col),
        ):
            if field in manifest and manifest[field] != expected:
                raise ValueError(
                    f"historical tabular source manifest {field} conflicts with state"
                )
    manifest = manifest or {}
    key_col = manifest.get("key_col")
    if key_col is None and any(
        segment.get("key_origin") == "input" for segment in old_segments
    ):
        raise ValueError("historical explicit business key column is unknown")
    group_col = manifest.get("group_col")
    if group_col is None and any(segment.get("group") for segment in old_segments):
        raise ValueError("historical group column is not reconstructable")

    context_columns = manifest.get(
        "context_columns", state.get("context_cols")
    )
    target_digest_col = manifest.get("target_source_digest_col")
    if state.get("input_guard_version") is not None and target_digest_col is None and any(
        not any(
            warning.get("code") == "UNVERIFIED_TARGET_PROVENANCE"
            for warning in (segment.get("input_warnings") or [])
            if isinstance(warning, dict)
        )
        for segment in old_segments
    ):
        raise ValueError(
            "historical target source-digest column is not reconstructable"
        )
    pivot = manifest.get("pivot_guard") or state.get("pivot_guard") or {}
    if pivot and not isinstance(pivot, dict):
        raise ValueError("historical pivot options are not reconstructable")
    if any(segment.get("input_status") == "blocked" for segment in old_segments):
        if target_digest_col is None and not pivot:
            raise ValueError("historical input guard options are not reconstructable")
    return {
        "sheet": sheet,
        "source_col": source_col,
        "target_col": target_col,
        "key_col": key_col,
        "context_cols": _reread_context_specs(context_columns),
        "target_source_digest_col": target_digest_col,
        "group_col": group_col,
        "pivot_sheet": pivot.get("pivot_sheet"),
        "pivot_key_col": (
            (pivot.get("key") or {}).get("pivot", {}).get("column")
            if pivot
            else None
        ),
        "pivot_compare": [
            f"{item.get('field')}={item.get('pivot', {}).get('column')}"
            for item in pivot.get("comparisons", [])
            if isinstance(item, dict)
        ],
        "pivot_authority": pivot.get("authority") if pivot else None,
        "no_header": no_header,
    }


def _reread_sdlxliff_options(
    old_job: Path,
    state: dict,
    input_path: Path,
) -> dict:
    manifest_path = _reread_asset_path(
        old_job, state.get("source_manifest_path"), "SDLXLIFF source manifest"
    )
    manifest = read_json(manifest_path)
    if not isinstance(manifest, dict):
        raise ValueError("historical SDLXLIFF source manifest must be an object")
    rules = manifest.get("rules") or {}
    if not isinstance(rules, dict):
        raise ValueError("historical SDLXLIFF rules are invalid")
    raw_options = {
        "tm_protection": manifest.get("tm_protection", "candidate-only"),
        "content_type_rules": rules.get("content_type", []),
        "exclude_rules": rules.get("exclusions", []),
    }
    options = validate_sdlxliff_options(raw_options)
    result = read_sdlxliff(input_path, options=options)
    expected_files = [
        (item.get("relative_path"), item.get("sha256"))
        for item in manifest.get("files", [])
        if isinstance(item, dict)
    ]
    actual_files = [
        (item.get("relative_path"), item.get("sha256"))
        for item in result.manifest.get("files", [])
    ]
    if not expected_files or expected_files != actual_files:
        raise ValueError("original SDLXLIFF input digest set does not match history")
    _assert_reread_coverage(
        state.get("segments"), result.segments, "sdlxliff"
    )
    has_rules = bool(raw_options["content_type_rules"] or raw_options["exclude_rules"])
    if has_rules and not str(state.get("project", "")).strip():
        raise ValueError(
            "historical SDLXLIFF rules require a reconstructable project profile"
        )
    return {
        "protect_exact_tm": (
            raw_options["tm_protection"]
            == "protect-exact-source-and-target"
        )
    }


def _remap_reread_state_paths(value: object, old_root: Path, new_root: Path):
    if isinstance(value, dict):
        return {
            key: _remap_reread_state_paths(item, old_root, new_root)
            for key, item in value.items()
        }
    if isinstance(value, list):
        return [
            _remap_reread_state_paths(item, old_root, new_root)
            for item in value
        ]
    if isinstance(value, str):
        old_prefix = str(old_root.resolve())
        if value == old_prefix or value.startswith(old_prefix + os.sep):
            return str(new_root.resolve()) + value[len(old_prefix):]
    return value


def cmd_reread(args):
    old_job_arg = Path(args.from_job)
    old_state_path = (
        old_job_arg if old_job_arg.is_file() else old_job_arg / "state.json"
    )
    old_job = old_state_path.parent.resolve()
    input_path = Path(args.input).resolve()
    new_job = Path(args.job).resolve()
    if not old_state_path.is_file():
        raise SystemExit(f"[reread] historical state is missing: {old_state_path}")
    if (
        new_job == old_job
        or new_job.is_relative_to(old_job)
        or _paths_alias(new_job, input_path)
        or (input_path.is_dir() and new_job.is_relative_to(input_path))
    ):
        raise SystemExit("[reread] old job, input, and target job must be distinct")

    original_state_bytes = old_state_path.read_bytes()
    try:
        state = json.loads(original_state_bytes)
        if not isinstance(state, dict):
            raise ValueError("historical state must be an object")
        if job_runtime_contract_version(state) != 1:
            raise ValueError("--from-job must be a historical runtime v1 job")
        old_policy = get_review_policy(state)
        review_mode = old_policy.get("mode")
        if old_policy != build_review_policy(review_mode):
            raise ValueError("historical review policy is not reconstructable")
        old_scope = get_check_scope(state)
        no_terminology = not old_scope.get("terminology_enabled", True)
        if old_scope != build_check_scope(no_terminology):
            raise ValueError("historical check scope is not reconstructable")
        source_lang = state.get("source_lang")
        target_lang = state.get("target_lang")
        if not all(
            isinstance(value, str) and value.strip()
            for value in (source_lang, target_lang)
        ):
            raise ValueError("historical language pair is not reconstructable")
        wordcount_basis = state.get("wordcount_basis")
        if wordcount_basis not in {"target-words", "source-chars"}:
            raise ValueError("historical wordcount basis is not reconstructable")
        project = state.get("project") or None
        if project is not None and not isinstance(project, str):
            raise ValueError("historical project ID is invalid")
        if args.profile_overlay and not project:
            raise ValueError("--profile-overlay requires a historical project ID")
        if not project and any(state.get(field) for field in ("sg_path", "terms_path")):
            raise ValueError(
                "historical ad-hoc style/terminology assets are not reconstructable"
            )
        input_format = state.get("input_format", "tabular")
        if input_format not in {"tabular", "sdlxliff"}:
            raise ValueError(f"unsupported historical input format: {input_format!r}")
    except (OSError, ValueError) as exc:
        raise SystemExit(f"[reread] {exc}") from exc

    plan = {
        "schema": "lqe.job-reread-migration-plan",
        "version": 1,
        "from_job": str(old_job),
        "input": str(input_path),
        "job": str(new_job),
        "validation_status": "pending",
        "inherited": {
            "project": project,
            "review_mode": review_mode,
            "review_policy": old_policy,
            "check_scope": old_scope,
            "source_lang": source_lang,
            "target_lang": target_lang,
            "wordcount_basis": wordcount_basis,
            "input_format": input_format,
            "source_col": state.get("source_col"),
            "target_col": state.get("target_col"),
            "sheet": state.get("sheet_name"),
            "no_header": state.get("no_header", False),
        },
        "profile_overlay": (
            str(Path(args.profile_overlay).resolve())
            if args.profile_overlay
            else None
        ),
        "context_overrides": (
            str(Path(args.context_overrides).resolve())
            if args.context_overrides
            else None
        ),
    }
    print(json.dumps(plan, ensure_ascii=False, sort_keys=True))

    try:
        if new_job.exists():
            raise ValueError(f"target job already exists: {new_job}")
        if not input_path.exists():
            raise ValueError(f"original input is missing: {input_path}")
        if detect_input_format(input_path, "auto") != input_format:
            raise ValueError("explicit input format does not match historical state")
        format_options = (
            _reread_tabular_options(old_job, state, input_path)
            if input_format == "tabular"
            else _reread_sdlxliff_options(old_job, state, input_path)
        )
    except (OSError, ValueError, SDLXLIFFImportError, XLSImportError) as exc:
        raise SystemExit(f"[reread] {exc}") from exc

    new_job.parent.mkdir(parents=True, exist_ok=True)
    staging_job = Path(
        tempfile.mkdtemp(prefix=f".{new_job.name}.reread.", dir=new_job.parent)
    )
    read_lock = new_job.parent / f".{staging_job.name}.lqe-read.lock"
    published = False
    try:
        read_args = argparse.Namespace(
            input=str(input_path),
            input_format=input_format,
            protect_exact_tm=format_options.get("protect_exact_tm", False),
            project=project,
            profile_overlay=args.profile_overlay,
            context_overrides=args.context_overrides,
            sheet=format_options.get("sheet"),
            source_col=format_options.get("source_col"),
            target_col=format_options.get("target_col"),
            key_col=format_options.get("key_col"),
            context_cols=format_options.get("context_cols", []),
            content_type_col=None,
            speaker_col=None,
            addressee_col=None,
            relationship_stage_col=None,
            scene_id_col=None,
            scene_tone_col=None,
            context_note_col=None,
            pivot_sheet=format_options.get("pivot_sheet"),
            pivot_key_col=format_options.get("pivot_key_col"),
            pivot_compare=format_options.get("pivot_compare", []),
            pivot_authority=format_options.get("pivot_authority"),
            target_source_digest_col=format_options.get(
                "target_source_digest_col"
            ),
            no_header=format_options.get("no_header", False),
            group_col=format_options.get("group_col"),
            terminology=None,
            no_terminology=no_terminology,
            style_guide=None,
            target_lang=target_lang,
            source_lang=source_lang,
            wordcount_basis=wordcount_basis,
            out=str(staging_job / "state.json"),
            review_mode=review_mode,
        )
        cmd_read(read_args)
        staged_state_path = staging_job / "state.json"
        staged_state = read_json(staged_state_path)
        require_current_job_runtime(staged_state, "reread")
        _assert_reread_coverage(
            state.get("segments"), staged_state.get("segments"), input_format
        )
        if get_review_policy(staged_state) != old_policy:
            raise ValueError("new job review policy differs from migration plan")
        if get_check_scope(staged_state) != old_scope:
            raise ValueError("new job check scope differs from migration plan")
        staged_state = _remap_reread_state_paths(
            staged_state, staging_job, new_job
        )
        write_json_atomic(staged_state_path, staged_state)
        if old_state_path.read_bytes() != original_state_bytes:
            raise ValueError("historical state changed during reread")
        if new_job.exists():
            raise FileExistsError(f"target job appeared during reread: {new_job}")
        os.rename(staging_job, new_job)
        published = True
    except (OSError, ValueError) as exc:
        raise SystemExit(f"[reread] {exc}") from exc
    finally:
        if not published:
            shutil.rmtree(staging_job, ignore_errors=True)
        read_lock.unlink(missing_ok=True)
    print(f"[reread] new runtime v2 job → {new_job}")

# ── lookup-terms ──────────────────────────────────────────────────────────────

def cmd_lookup_terms(args):
    state = read_json(args.state)
    terms = _load_terms(state)
    if not terms:
        print("[lqe_io] no terminology available.", file=sys.stderr)
        return

    term_map = _group_terms(terms)

    segs = state["segments"]
    if args.ids:
        id_set = set(int(x) for x in args.ids.split(","))
        segs = [s for s in segs if s["id"] in id_set]

    # 逐段匹配，避免跨段拼接产生误命中
    hits: dict[str, dict] = {}  # term_source → {senses, seg_ids}
    for seg in segs:
        src_text = seg["source"]
        for term_src, senses in term_map.items():
            if term_src in src_text:
                if term_src not in hits:
                    hits[term_src] = {"senses": senses, "seg_ids": []}
                hits[term_src]["seg_ids"].append(seg["id"])

    if not hits:
        print("[lookup-terms] no terminology matches found.")
        return

    print(f"[lookup-terms] {len(hits)} matches:\n")
    for src, info in sorted(hits.items(), key=lambda x: -len(x[0])):
        seg_ids = info["seg_ids"]
        id_str = f"  (segs: {seg_ids})" if len(seg_ids) <= 5 else f"  ({len(seg_ids)} segs)"
        tgt_str = " | ".join(s["target"] for s in info["senses"])
        print(f"  {src} → {tgt_str}{id_str}")


# ── apply-fixes ───────────────────────────────────────────────────────────────

def _protected_ids(args, *, allow_candidates: bool = False) -> set[int]:
    ids: set[int] = set()
    if getattr(args, "protected_ids", None):
        ids.update(int(x.strip()) for x in args.protected_ids.split(",") if x.strip())
    if getattr(args, "protected_file", None):
        data = read_json(args.protected_file)
        if isinstance(data, dict):
            if "protected_ids" in data:
                data = data["protected_ids"]
            elif "candidate_ids" in data:
                data = data["candidate_ids"] if allow_candidates else []
            else:
                data = data.get("segments") or []
        for item in data:
            if isinstance(item, int):
                ids.add(item)
            elif isinstance(item, dict):
                sid = item.get("id", item.get("seg_id", item.get("segment_id")))
                if sid is not None:
                    ids.add(int(sid))
    return ids


def _validated_tm_candidate_ids(
    state: dict, candidate_path: Path, payload: object
) -> set[int]:
    expected_path = state.get("tm_candidates_path")
    if not isinstance(expected_path, str) or not expected_path:
        raise ValueError("state has no tm_candidates_path")
    if not _paths_alias(candidate_path, Path(expected_path)):
        raise ValueError(
            "candidate file does not match state tm_candidates_path: "
            f"{candidate_path} != {expected_path}"
        )
    if not isinstance(payload, dict):
        raise ValueError("TM candidate payload must be an object")
    candidate_ids = payload.get("candidate_ids")
    candidates = payload.get("segments")
    if not isinstance(candidate_ids, list) or not all(
        type(segment_id) is int for segment_id in candidate_ids
    ):
        raise ValueError("TM candidate_ids must be an integer array")
    if len(candidate_ids) != len(set(candidate_ids)):
        raise ValueError("TM candidate_ids contains duplicates")
    if not isinstance(candidates, list):
        raise ValueError("TM candidate segments must be an array")

    candidate_by_id = {}
    for index, candidate in enumerate(candidates):
        if not isinstance(candidate, dict) or type(candidate.get("id")) is not int:
            raise ValueError(f"TM candidate segments[{index}] has an invalid id")
        segment_id = candidate["id"]
        if segment_id in candidate_by_id:
            raise ValueError(f"TM candidate segments contains duplicate id {segment_id}")
        candidate_by_id[segment_id] = candidate
    if set(candidate_ids) != set(candidate_by_id):
        raise ValueError(
            "TM candidate_ids must exactly match candidate segment evidence"
        )

    state_by_id = {segment.get("id"): segment for segment in state.get("segments", [])}
    for segment_id in candidate_ids:
        candidate = candidate_by_id[segment_id]
        state_segment = state_by_id.get(segment_id)
        if state_segment is None:
            raise ValueError(f"TM candidate id {segment_id} is not in state")
        if candidate.get("source_ref") != state_segment.get("source_ref"):
            raise ValueError(f"TM candidate id {segment_id} source_ref mismatch")
        evidence = candidate.get("evidence")
        metadata = (state_segment.get("metadata") or {}).get("sdlxliff") or {}
        if not is_exact_tm(evidence) or not is_exact_tm(metadata):
            raise ValueError(f"TM candidate id {segment_id} lacks exact-match evidence")
        for key in ("origin", "match_percent", "text_match"):
            if evidence.get(key) != metadata.get(key):
                raise ValueError(
                    f"TM candidate id {segment_id} evidence mismatch for {key}"
                )
    return set(candidate_ids)


def _state_protected_ids(state) -> set[int]:
    return {
        s["id"]
        for s in state.get("segments", [])
        if s.get("protected") or s.get("input_status") == "blocked"
    }


def _stage_json_replacement(path: Path, value: object) -> Path:
    payload = json.dumps(value, ensure_ascii=False, indent=2)
    staged = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            dir=path.parent,
            prefix=f".{path.name}.",
            suffix=".tmp",
            delete=False,
        ) as handle:
            staged = Path(handle.name)
            handle.write(payload)
        return staged
    except BaseException:
        if staged is not None and staged.exists():
            staged.unlink()
        raise


def _assert_json_unchanged(path: Path, expected: object, *, label: str) -> None:
    if read_json(path) != expected:
        raise ValueError(f"{label} changed during artifact publication: {path}")


def _stage_bound_result_replacements(
    errors_path: Path,
    errors_data: list[dict],
    state: dict,
    manifest: dict | None,
) -> tuple[list[Path], list[tuple[Path, Path]]]:
    staged_errors = _stage_json_replacement(errors_path, errors_data)
    staged = [staged_errors]
    replacements = [(staged_errors, errors_path)]
    if requires_bound_artifacts(state):
        if manifest is None:
            staged_errors.unlink(missing_ok=True)
            raise ValueError("bound result publication requires a generation")
        contract_path = result_contract_path(errors_path)
        staged_contract = _stage_json_replacement(
            contract_path,
            build_result_contract(manifest, errors_data),
        )
        staged.append(staged_contract)
        replacements.append((staged_contract, contract_path))
    return staged, replacements


def _publish_bound_result_update(
    errors_path: Path,
    errors_data: list[dict],
    state: dict,
    manifest: dict | None,
) -> None:
    staged: list[Path] = []
    try:
        staged, replacements = _stage_bound_result_replacements(
            errors_path,
            errors_data,
            state,
            manifest,
        )
        publish_replacement_transaction(replacements)
    finally:
        for path in staged:
            path.unlink(missing_ok=True)


def _publish_protection_transaction(
    state_path: Path,
    state: dict,
    output_path: Path,
    payload: dict,
) -> None:
    staged: list[Path] = []
    try:
        staged_output = _stage_json_replacement(output_path, payload)
        staged_state = _stage_json_replacement(state_path, state)
        staged.extend((staged_output, staged_state))
        publish_replacement_transaction(
            [
                (staged_output, output_path),
                (staged_state, state_path),
            ]
        )
    finally:
        for path in staged:
            path.unlink(missing_ok=True)


def _publish_write_transaction(
    state_path: Path,
    state: dict,
    errors_path: Path,
    errors_data: list,
    output_path: Path,
    staged_output: Path,
    *,
    publish_errors: bool,
    manifest: dict | None,
) -> None:
    staged: list[Path] = [staged_output]
    try:
        replacements = []
        if publish_errors:
            result_staged, result_replacements = _stage_bound_result_replacements(
                errors_path,
                errors_data,
                state,
                manifest,
            )
            staged.extend(result_staged)
            replacements.extend(result_replacements)
        staged_state = _stage_json_replacement(state_path, state)
        staged.append(staged_state)
        replacements.extend(
            [
                (staged_output, output_path),
                (staged_state, state_path),
            ]
        )
        publish_replacement_transaction(replacements)
    finally:
        for path in staged:
            path.unlink(missing_ok=True)


def _scrub_protected_entries(errors_data: list, protected_ids: set[int]) -> int:
    changed = 0
    for entry in errors_data:
        if entry.get("id") in protected_ids:
            if entry.get("errors") or entry.get("corrected") is not None:
                changed += len(entry.get("errors", [])) or 1
            entry["errors"] = []
            entry["corrected"] = None
    return changed


def _correction_candidates(errors_data: list[dict]) -> dict[int, str]:
    return {
        entry["id"]: entry["corrected"]
        for entry in errors_data
        if entry.get("corrected") is not None
        and not any(
            error.get("protected") for error in (entry.get("errors") or [])
        )
    }


def _cmd_protect_segments_locked(args, state_path: Path, state: dict):
    require_current_job_runtime(state, "protect-segments")
    protected_file = getattr(args, "protected_file", None)
    protected_payload = read_json(protected_file) if protected_file else None
    try:
        if isinstance(protected_payload, dict) and "candidate_ids" in protected_payload:
            ids = _validated_tm_candidate_ids(
                state, Path(protected_file), protected_payload
            )
            if getattr(args, "protected_ids", None):
                ids.update(
                    int(value.strip())
                    for value in args.protected_ids.split(",")
                    if value.strip()
                )
        else:
            ids = _protected_ids(args, allow_candidates=False)
    except (OSError, ValueError) as exc:
        raise SystemExit(f"[protect-segments] {exc}") from exc
    if not ids:
        print("[protect-segments] no ids supplied; state unchanged.")
        return

    out_path = Path(args.out) if args.out else state_path.parent / "tm_protected.json"
    try:
        validate_artifact_paths(
            {"protection decision": out_path},
            {
                "state": state_path,
                **state_reference_paths(state),
                **(
                    {"protected file": Path(protected_file)}
                    if protected_file
                    else {}
                ),
            },
            context="protect-segments",
        )
    except ValueError as exc:
        raise SystemExit(f"[protect-segments] {exc}") from exc

    seg_by_id = {s["id"]: s for s in state.get("segments", [])}
    valid = sorted(sid for sid in ids if sid in seg_by_id)
    unknown = sorted(sid for sid in ids if sid not in seg_by_id)
    for sid in valid:
        seg = seg_by_id[sid]
        seg["protected"] = True
        if seg.get("protected_reason") != "SOURCE_LOCKED":
            seg["protected_reason"] = args.reason
        working_target = current_target(seg)
        if working_target != seg.get("target", ""):
            seg["current_target"] = working_target
        seg["corrected"] = None

    payload = {
        "protected_ids": valid,
        "reason": args.reason,
        "source": "agent_decision",
    }
    if unknown:
        payload["unknown_ids"] = unknown
    try:
        _publish_protection_transaction(state_path, state, out_path, payload)
    except (OSError, ValueError) as exc:
        raise SystemExit(f"[protect-segments] {exc}") from exc
    print(f"[protect-segments] protected {len(valid)} segment(s) → {state_path}")
    print(f"[protect-segments] protected file → {out_path}")
    if unknown:
        print(f"[protect-segments] ignored unknown ids: {unknown[:20]}")


def cmd_protect_segments(args):
    from lqe_split_contract import generation_lock

    state_path = Path(args.state).resolve()
    try:
        with generation_lock(state_path.parent / "chunks", exclusive=True):
            state = read_json(state_path)
            _cmd_protect_segments_locked(args, state_path, state)
    except (OSError, ValueError) as exc:
        raise SystemExit(f"[protect-segments] {exc}") from exc


def _cmd_build_results_locked(args, state_path: Path, state: dict):
    require_current_job_runtime(state, "build-results")
    checks_path = Path(args.checks)
    out_path = Path(args.out)
    entries = normalize_check_entries(
        json.loads(checks_path.read_text(encoding="utf-8")),
        label=args.checks,
        review_policy=get_review_policy(state),
    )
    _validate_scope_or_exit(
        state,
        entries,
        issues_key="issues",
        label=Path(args.checks).name,
        command="build-results",
    )
    segments = deepcopy(state["segments"])
    state_ids = {segment["id"] for segment in segments}
    check_ids = {entry["id"] for entry in entries}
    missing = sorted(state_ids - check_ids)
    extra = sorted(check_ids - state_ids)
    if missing or extra:
        sys.exit(
            f"[build-results] check ids must match state segment ids: "
            f"missing={missing} extra={extra}"
        )
    results = build_results(
        segments,
        entries,
        review_policy=get_review_policy(state),
    )
    if requires_bound_artifacts(state):
        raise SystemExit(
            "[build-results] unbound checks cannot publish into a current job; "
            "use lqe_chunk.py merge"
        )
    try:
        validate_artifact_paths(
            {"results": out_path},
            {
                "state": state_path,
                "checks": checks_path,
                **state_reference_paths(state),
            },
            context="build-results",
        )
        write_json_atomic(out_path, results)
    except (OSError, ValueError) as exc:
        raise SystemExit(f"[build-results] {exc}") from exc


def cmd_build_results(args):
    from lqe_split_contract import generation_lock

    state_path = Path(args.state)
    try:
        with generation_lock(state_path.parent / "chunks", exclusive=True):
            state = read_json(state_path)
            _cmd_build_results_locked(args, state_path, state)
    except (OSError, ValueError) as exc:
        raise SystemExit(f"[build-results] {exc}") from exc


def _cmd_apply_fixes_locked(
    args,
    state_path: Path,
    state: dict,
    segments: list[dict],
    manifest: dict | None,
    revalidate,
):
    require_current_job_runtime(state, "apply-fixes")
    errors_data = read_json(args.errors)
    original_errors_data = deepcopy(errors_data)
    _validate_scope_or_exit(
        state,
        errors_data,
        issues_key="errors",
        label=Path(args.errors).name,
        command="apply-fixes",
    )
    protected_ids = _protected_ids(args) | _state_protected_ids(state)
    raw_attempted = _correction_candidates(errors_data)

    scrubbed = _scrub_protected_entries(errors_data, protected_ids)
    try:
        scoring_policy = resolve_scoring_policy(
            state,
            scoring_policy_overrides(args),
        )
    except (CheckFormatError, ValueError) as exc:
        raise SystemExit(f"[apply-fixes] {exc}") from exc
    scorecard_profile = load_scorecard_profile(
        scoring_policy["scorecard_profile"]
    )
    seg_ids = {segment["id"] for segment in segments}
    issues = _validate_errors(errors_data, seg_ids, scorecard_profile)
    for msg in issues:
        print(f"[validate] {msg}")
    try:
        verified = _verify_result_payload_with_segments(
            state,
            segments,
            manifest,
            errors_data,
            Path(args.errors),
            command="apply-fixes",
        )
        computation = score_errors(
            state,
            verified,
            scoring_policy,
            protected_ids=protected_ids,
        )
    except (CheckFormatError, ValueError) as exc:
        raise SystemExit(f"[apply-fixes] {exc}") from exc
    errors_data = computation["annotated_errors"]
    attempted = _correction_candidates(errors_data)

    corrections = {sid: text for sid, text in attempted.items() if sid not in protected_ids}
    segment_by_id = {segment["id"]: segment for segment in state["segments"]}
    protected_skipped = [
        {
            "id": sid,
            "reason": _protection_reason(segment_by_id[sid], protected_ids),
            "evidence": segment_by_id[sid].get("protection_evidence"),
            "attempted": text,
        }
        for sid, text in raw_attempted.items()
        if sid in protected_ids and sid in segment_by_id
    ]
    if not corrections and not protected_skipped:
        if scrubbed or computation["annotations_changed"]:
            revalidate()
            _assert_json_unchanged(
                Path(args.errors),
                original_errors_data,
                label="errors input",
            )
            _publish_bound_result_update(
                Path(args.errors),
                errors_data,
                state,
                manifest,
            )
        if scrubbed:
            print(
                f"[apply-fixes] scrubbed {scrubbed} protected-segment "
                f"issue(s) from {args.errors}"
            )
        print("[lqe_io] apply-fixes: no corrections found, state unchanged.")
        print(
            json.dumps(
                {"applied_count": 0, "lifecycle": "review_required"},
                ensure_ascii=False,
            )
        )
        return

    cur_iter = state.get("iteration", 0)
    history = state.get("error_history", [])
    iteration_targets = {
        segment["id"]: current_target(segment)
        for segment in state["segments"]
    }
    score_result = computation["output"]
    score = score_result["score"]
    supplied_score = getattr(args, "score", None)
    if supplied_score is not None and not math.isclose(
        float(supplied_score), score, abs_tol=0.005
    ):
        print(
            f"[apply-fixes] supplied score {float(supplied_score):g} differs "
            f"from recomputed score {score:g}; using recomputed score",
            file=sys.stderr,
        )
    cur_entry = {
        "iteration": cur_iter,
        "score": score,
        "status": score_result["status"],
        "errors": errors_data,
        "corrections_count": len(corrections),
        "protected_ids": sorted(protected_ids),
        "skipped_corrections": protected_skipped,
        "review_targets": {
            str(segment_id): target
            for segment_id, target in iteration_targets.items()
        },
    }
    history.append(cur_entry)
    state["error_history"] = history

    next_iter = cur_iter + (1 if corrections else 0)
    for seg in state["segments"]:
        if seg["id"] in protected_ids:
            seg["protected"] = True
            if not seg.get("protected_reason"):
                seg["protected_reason"] = "TM_100_MATCH"
            working_target = current_target(seg)
            if working_target != seg.get("target", ""):
                seg["current_target"] = working_target
            seg["corrected"] = None
            continue
        if seg["id"] in corrections:
            seg["current_target"] = corrections[seg["id"]]
            seg["corrected"] = corrections[seg["id"]]
            seg["iter"] = next_iter

    state["iteration"] = next_iter
    state["threshold"] = scoring_policy["threshold"]
    state["scoring_policy"] = scoring_policy
    state["pending_recheck"] = bool(corrections)

    archived = state_path.parent / f"errors_iter{cur_iter}.json"
    iter_out = state_path.parent / (_job_label(state_path) + f"_lqe_iter{cur_iter}.xlsx")
    try:
        validate_artifact_paths(
            {
                "iteration error archive": archived,
                "iteration report": iter_out,
            },
            {
                "state": state_path,
                "errors": Path(args.errors),
                **state_reference_paths(state),
            },
            context="apply-fixes",
        )
    except ValueError as exc:
        raise SystemExit(f"[apply-fixes] {exc}") from exc
    with tempfile.NamedTemporaryFile(
        prefix=f".{iter_out.stem}.",
        suffix=iter_out.suffix,
        dir=iter_out.parent,
        delete=False,
    ) as staging_file:
        staged_report = Path(staging_file.name)
    staged_paths = []
    try:
        _build_xlsx(
            state,
            [cur_entry],
            score,
            scoring_policy["threshold"],
            staged_report,
            scoring_policy["scorecard_profile"],
            announce=False,
            scoring_policy=scoring_policy,
            scoring_computation=computation,
            review_targets=iteration_targets,
        )
        replacements = []
        if scrubbed or computation["annotations_changed"]:
            result_staged, result_replacements = _stage_bound_result_replacements(
                Path(args.errors),
                errors_data,
                state,
                manifest,
            )
            staged_paths.extend(result_staged)
            replacements.extend(result_replacements)
        staged_archive = _stage_json_replacement(archived, errors_data)
        staged_state = _stage_json_replacement(state_path, state)
        staged_paths.extend((staged_archive, staged_state))
        replacements.extend(
            [
                (staged_archive, archived),
                (staged_report, iter_out),
                (staged_state, state_path),
            ]
        )
        revalidate()
        _assert_json_unchanged(
            Path(args.errors),
            original_errors_data,
            label="errors input",
        )
        publish_replacement_transaction(replacements)
    except (OSError, ValueError) as exc:
        raise SystemExit(f"[apply-fixes] {exc}") from exc
    finally:
        staged_report.unlink(missing_ok=True)
        for staged_path in staged_paths:
            staged_path.unlink(missing_ok=True)

    if scrubbed:
        print(
            f"[apply-fixes] scrubbed {scrubbed} protected-segment "
            f"issue(s) from {args.errors}"
        )
    print(f"[lqe_io] Applied {len(corrections)} corrections → iteration {next_iter}")
    print(f"[lqe_io] Errors archived → {archived}")
    print(f"[lqe_io] Output → {iter_out}")
    print(
        json.dumps(
            {
                "applied_count": len(corrections),
                "lifecycle": (
                    "pending_recheck" if corrections else "review_required"
                ),
            },
            ensure_ascii=False,
        )
    )


def cmd_apply_fixes(args):
    from lqe_chunk import verification_generation_lease

    state_path = Path(args.state)
    try:
        require_current_job_runtime(read_json(state_path), "apply-fixes")
        with verification_generation_lease(
            state_path,
            exclusive=True,
        ) as (state, segments, manifest, revalidate):
            _cmd_apply_fixes_locked(
                args,
                state_path,
                state,
                segments,
                manifest,
                revalidate,
            )
    except (OSError, ValueError) as exc:
        raise SystemExit(f"[apply-fixes] {exc}") from exc


# ── write ─────────────────────────────────────────────────────────────────────

_DARK_BLUE   = PatternFill("solid", fgColor="073763")
_LIGHT_BLUE  = PatternFill("solid", fgColor="CFE2F3")
_ORANGE      = PatternFill("solid", fgColor="FCE5CD")
_RED         = PatternFill("solid", fgColor="CC0000")
_GREEN       = PatternFill("solid", fgColor="006600")
_GREEN_LIGHT = PatternFill("solid", fgColor="D9EAD3")
_WHITE_FONT = Font(color="FFFFFF")
_BOLD       = Font(bold=True)
_CENTER     = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT_TOP   = Alignment(horizontal="left", vertical="top", wrap_text=True)
_LEFT_MID   = Alignment(horizontal="left", vertical="center")


def _validate_errors(errors_data: list, seg_ids: set, scorecard_profile: dict | None = None) -> list[str]:
    issues = []
    valid_categories = set(scorecard_category_order(scorecard_profile))
    for entry in errors_data:
        sid = entry.get("id")
        if sid not in seg_ids:
            issues.append(f"[seg {sid}] 未知 segment id")
            continue
        errs = entry.get("errors", [])
        for e in errs:
            raw_cat = e.get("category", "")
            cat = normalize_category_for_profile(raw_cat, scorecard_profile)
            sev = e.get("severity", "")
            if cat not in valid_categories:
                issues.append(f"[seg {sid}] 非法 category: '{raw_cat}'")
            if sev not in _VALID_SEVERITIES:
                issues.append(f"[seg {sid}] 非法 severity: '{sev}'")
            new_sev = apply_severity(cat, sev, scorecard_profile)
            if new_sev != sev:
                issues.append(f"[seg {sid}] {cat} severity {sev}→{new_sev} (auto-corrected)")
                e["severity"] = new_sev
    return issues


def _s(cell, fill=None, font=None, align=None):
    if fill:  cell.fill  = fill
    if font:  cell.font  = font
    if align: cell.alignment = align


def _set_excel_text(cell, value) -> None:
    cell.value = value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"


_GRAPHEME_PATTERN = regex.compile(r"\X")
_EMOJI_WIDTH_PATTERN = regex.compile(
    r"\p{Emoji_Presentation}|\p{Regional_Indicator}|\u20e3"
)
_EMOJI_BASE_PATTERN = regex.compile(r"\p{Emoji}")
_MAX_EXCEL_ROW_HEIGHT = 409.0


def _grapheme_units(grapheme: str) -> int:
    if grapheme == "\t":
        return 4
    if _EMOJI_WIDTH_PATTERN.search(grapheme) or (
        "\ufe0f" in grapheme and _EMOJI_BASE_PATTERN.search(grapheme)
    ):
        return 2
    units = 0
    for char in grapheme:
        if char in "\r\n":
            continue
        if unicodedata.combining(char) or unicodedata.category(char) in {"Mn", "Me", "Cf"}:
            continue
        char_units = 2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1
        units = max(units, char_units)
    return units


def _display_units(value) -> int:
    text = "" if value is None else str(value)
    return sum(
        _grapheme_units(grapheme)
        for grapheme in _GRAPHEME_PATTERN.findall(text)
    )


def _wrapped_line_count(value, column_width) -> int:
    text = "" if value is None else str(value)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    capacity = max(1, math.floor(float(column_width) * 0.88))
    total_lines = 0

    def place_characters(characters, used_units):
        additional_lines = 0
        for grapheme in _GRAPHEME_PATTERN.findall(characters):
            grapheme_units = _grapheme_units(grapheme)
            if not grapheme_units:
                continue
            if used_units and used_units + grapheme_units > capacity:
                additional_lines += 1
                used_units = 0
            used_units += grapheme_units
        return additional_lines, used_units

    for explicit_line in text.split("\n"):
        matches = list(re.finditer(r"\S+", explicit_line))
        if not matches:
            additional_lines, _ = place_characters(explicit_line, 0)
            total_lines += 1 + additional_lines
            continue

        wrapped_lines = 1
        additional_lines, used_units = place_characters(
            explicit_line[:matches[0].start()],
            0,
        )
        wrapped_lines += additional_lines
        previous_end = matches[0].start()
        for index, match in enumerate(matches):
            separator = (
                explicit_line[previous_end:match.start()]
                if index
                else ""
            )
            token_units = _display_units(match.group())
            separator_units = _display_units(separator)
            if used_units and used_units + separator_units + token_units > capacity:
                wrapped_lines += 1
                used_units = 0
                separator = ""
            additional_lines, used_units = place_characters(
                separator + match.group(),
                used_units,
            )
            wrapped_lines += additional_lines
            previous_end = match.end()
        additional_lines, _ = place_characters(
            explicit_line[previous_end:],
            used_units,
        )
        wrapped_lines += additional_lines
        total_lines += wrapped_lines

    return total_lines


def _wrapped_row_height(cells, minimum=15.75, *, context="wrapped row") -> float:
    maximum_lines = max(
        (_wrapped_line_count(value, column_width) for value, column_width in cells),
        default=1,
    )
    required_height = max(float(minimum), 2.0 + 16.5 * maximum_lines)
    if required_height > _MAX_EXCEL_ROW_HEIGHT:
        raise ValueError(
            f"{context}: {maximum_lines} wrapped lines require {required_height:g} pt; "
            f"Excel/WPS row height is limited to {_MAX_EXCEL_ROW_HEIGHT:g} pt"
        )
    return required_height


def _fit_wrapped_row(cells, minimum=15.75, *, context="wrapped row") -> tuple[float, float]:
    maximum_lines = max(
        (_wrapped_line_count(value, column_width) for value, column_width in cells),
        default=1,
    )
    for font_size, line_height in ((11.0, 16.5), (10.0, 15.0), (9.0, 13.5)):
        required_height = max(float(minimum), 2.0 + line_height * maximum_lines)
        if required_height <= _MAX_EXCEL_ROW_HEIGHT:
            return required_height, font_size
    raise ValueError(
        f"{context}: {maximum_lines} wrapped lines still require {required_height:g} pt "
        f"at {font_size:g} pt; Excel/WPS row height is limited to "
        f"{_MAX_EXCEL_ROW_HEIGHT:g} pt"
    )


def _fit_or_span_wrapped_row(
    cells,
    minimum=15.75,
    *,
    context="wrapped row",
) -> tuple[float, float, int]:
    try:
        height, font_size = _fit_wrapped_row(
            cells,
            minimum,
            context=context,
        )
        return height, font_size, 1
    except ValueError:
        maximum_lines = max(
            (_wrapped_line_count(value, column_width) for value, column_width in cells),
            default=1,
        )
        total_height = max(float(minimum), 2.0 + 13.5 * maximum_lines)
        row_span = math.ceil(total_height / _MAX_EXCEL_ROW_HEIGHT)
        return total_height / row_span, 9.0, row_span


def _set_row_font_size(worksheet, row: int, columns: int, font_size: float) -> None:
    if font_size >= 11.0:
        return
    for column in range(1, columns + 1):
        cell = worksheet.cell(row=row, column=column)
        font = copy(cell.font)
        font.sz = font_size
        cell.font = font


def _segment_filename(state: dict, segment: dict, source_row=None) -> str:
    if state.get("input_format") == "sdlxliff":
        metadata = segment.get("metadata") or {}
        sdl_metadata = metadata.get("sdlxliff") or {}
        file_original = _text(sdl_metadata.get("file_original"))
        if file_original:
            return file_original
        source_ref = segment.get("source_ref") or {}
        relative_path = _text(source_ref.get("relative_path"))
        if relative_path:
            return relative_path
        return Path(state.get("input_path") or "").name
    fallback = Path(state.get("input_path") or "").stem
    headers = state.get("headers") or []
    try:
        source_path_index = headers.index("来源相对路径")
    except ValueError:
        return fallback
    if source_row is None:
        segments = state.get("segments") or []
        rows = state.get("rows_raw") or []
        if len(rows) == len(segments):
            segment_id = segment.get("id")
            for index, candidate in enumerate(segments):
                if candidate.get("id") == segment_id:
                    source_row = rows[index]
                    break
    if not source_row or source_path_index >= len(source_row):
        return fallback
    value = source_row[source_path_index]
    return _text(value) or fallback


def _report_source_table(state: dict) -> tuple[list[str], list[list[object]]]:
    segments = state.get("segments") or []
    if state.get("input_format") == "sdlxliff":
        headers = ["来源文件", "TU ID", "SDL Segment ID", "原文", "译文"]
        rows = []
        for segment in segments:
            source_ref = segment.get("source_ref") or {}
            rows.append(
                [
                    _segment_filename(state, segment),
                    source_ref.get("tu_id") or "",
                    source_ref.get("sdl_segment_id") or "",
                    segment.get("source", ""),
                    segment.get("target", ""),
                ]
            )
        return headers, rows

    rows_raw = state.get("rows_raw") or []
    if len(rows_raw) != len(segments):
        raise ValueError(
            "tabular report rows_raw/segments length mismatch: "
            f"rows_raw={len(rows_raw)} segments={len(segments)}"
        )
    return list(state.get("headers") or []), [list(row) for row in rows_raw]


def _protection_reason(segment: dict, protected_ids: set[int]) -> str:
    if segment.get("id") not in protected_ids:
        return ""
    return _text(segment.get("protected_reason")) or "TM_100_MATCH"


def _protection_evidence(segment: dict, protected_ids: set[int]) -> str:
    reason = _protection_reason(segment, protected_ids)
    if not reason:
        return ""
    evidence = segment.get("protection_evidence")
    if evidence is None:
        return reason
    return json.dumps(
        {"reason": reason, "evidence": evidence},
        ensure_ascii=False,
        sort_keys=True,
    )


def _build_xlsx(
    state,
    history,
    score,
    threshold,
    out_path,
    scorecard_profile_id="legacy",
    *,
    announce=True,
    scoring_policy=None,
    scoring_computation=None,
    report_contract_results=None,
    reference_suggestions=None,
    review_targets=None,
):
    validate_error_history_term_contract(
        history,
        segments=state.get("segments", []),
        label="report history",
    )
    reference_suggestions = reference_suggestions or {}
    if scoring_policy is not None:
        scorecard_profile_id = scoring_policy["scorecard_profile"]
        threshold = scoring_policy["threshold"]
    scorecard_profile = load_scorecard_profile(scorecard_profile_id)
    severity_points = scorecard_severity_points(
        scorecard_profile,
        (scoring_policy or {}).get("severity_scale", "lisa"),
    )
    categories = scorecard_category_order(scorecard_profile)
    check_scope = get_check_scope(state)
    review_policy = get_review_policy(state)
    terminology_status = (
        "Enabled"
        if check_scope["terminology_enabled"]
        else "Disabled by runtime request"
    )
    enabled_modules = ", ".join(check_scope["enabled_modules"])
    scope_summary = (
        f"Terminology check: {terminology_status}; "
        f"Enabled modules: {enabled_modules}"
    )
    segments = state["segments"]
    seg_map = {s["id"]: s for s in segments}
    current_review_targets = {
        segment["id"]: (
            review_targets[segment["id"]]
            if review_targets is not None
            and segment["id"] in review_targets
            else current_target(segment)
        )
        for segment in segments
    }
    cat_counts: dict[str, dict[str, int]] = {
        cat: {"Neutral": 0, "Minor": 0, "Major": 0, "Critical": 0}
        for cat in categories
    }
    rep_counts: dict[str, dict[str, int]] = {
        cat: {"Neutral": 0, "Minor": 0, "Major": 0, "Critical": 0}
        for cat in categories
    }
    detail_rows: list[dict] = []
    all_protected_ids = set()
    for entry in history:
        all_protected_ids.update(entry.get("protected_ids", []))
        all_protected_ids.update(
            result["id"]
            for result in entry.get("errors", [])
            if any(issue.get("protected") for issue in result.get("errors", []))
        )
    for seg in segments:
        if seg.get("protected"):
            all_protected_ids.add(seg["id"])
    max_iter = max((entry["iteration"] for entry in history), default=0)
    latest_entry = history[-1] if history else None

    def _entry_review_target(entry, segment):
        snapshots = entry.get("review_targets")
        if isinstance(snapshots, dict):
            value = snapshots.get(
                str(segment["id"]),
                snapshots.get(segment["id"]),
            )
            if isinstance(value, str):
                return value
        has_term_spans = any(
            result.get("id") == segment["id"]
            and any("term_spans" in issue for issue in result.get("errors", []))
            for result in entry.get("errors", [])
        )
        if has_term_spans:
            raise ValueError(
                f"iteration {entry.get('iteration')} segment {segment['id']} "
                "is missing its review_targets snapshot"
            )
        if entry is latest_entry:
            return current_review_targets[segment["id"]]
        return segment.get("target", "")

    for entry in history:
        fixed = entry["iteration"] < max_iter
        for e_seg in entry["errors"]:
            seg = seg_map.get(e_seg["id"])
            if not seg:
                continue
            if seg["id"] in all_protected_ids:
                continue
            corrected = e_seg.get("corrected")
            review_entry = {**e_seg, "corrected": corrected}
            for e in e_seg.get("errors", []):
                cat = normalize_category_for_profile(e.get("category", "Other"), scorecard_profile)
                sev = apply_severity(cat, e.get("severity", "Minor"), scorecard_profile)
                if entry is latest_entry:
                    if e.get("repeated"):
                        if cat in rep_counts:
                            rep_counts[cat][sev] = rep_counts[cat].get(sev, 0) + 1
                    elif cat in cat_counts:
                        cat_counts[cat][sev] = cat_counts[cat].get(sev, 0) + 1
                review_status, edit_status, check_source = _issue_review_columns(
                    e,
                    seg["id"],
                )
                detail_rows.append({
                    "filename": _segment_filename(state, seg),
                    "seg_id":   seg["id"],
                    "source":   seg["source"],
                    "original": _entry_review_target(entry, seg),
                    "issue":    e,
                    "entry":    review_entry,
                    "corrected": corrected,
                    "parent":   scorecard_category_parent(cat, scorecard_profile),
                    "category": cat,
                    "severity": sev,
                    "iteration": f"Iter {entry['iteration']}",
                    "comment":  ("[Repeated] " if e.get("repeated") else "") + e.get("comment", ""),
                    "fixed":    fixed,
                    "processing": _issue_processing_label(
                        e,
                        review_entry,
                        protected=False,
                    ),
                    "review_status": review_status,
                    "edit_status": edit_status,
                    "check_source": check_source,
                })

    if scoring_computation is not None:
        for category in categories:
            cat_counts[category] = {
                severity: int(
                    scoring_computation.get("category_counts", {})
                    .get(category, {})
                    .get(severity, 0)
                )
                for severity in ("Neutral", "Minor", "Major", "Critical")
            }
            rep_counts[category] = {
                severity: int(
                    scoring_computation.get("repeated_counts", {})
                    .get(category, {})
                    .get(severity, 0)
                )
                for severity in ("Neutral", "Minor", "Major", "Critical")
            }

    total_counts = {"Neutral": 0, "Minor": 0, "Major": 0, "Critical": 0}
    total_rep    = {"Neutral": 0, "Minor": 0, "Major": 0, "Critical": 0}
    for c in cat_counts.values():
        for sev, n in c.items():
            total_counts[sev] += n
    for c in rep_counts.values():
        for sev, n in c.items():
            total_rep[sev] += n
    total_raw = sum(
        raw_points(counts, scorecard_profile, severity_points)
        for counts in cat_counts.values()
    )
    total_weighted = (
        scoring_computation["total_weighted"]
        if scoring_computation is not None
        else sum(
            weighted_points(cat, counts, scorecard_profile, severity_points)
            for cat, counts in cat_counts.items()
        )
    )

    current_entries = (
        {e["id"]: e for e in history[-1].get("errors", [])}
        if history
        else {}
    )

    def _review_entry(seg):
        current_entry = current_entries.get(seg["id"])
        if current_entry is None:
            return {"errors": [], "corrected": None}
        return current_entry

    def _review_summary(errs, field):
        values = []
        for issue in errs:
            value = _text(issue.get(field))
            if value:
                values.append(value)
        return "\n".join(values)

    def _problem_summary(errs):
        comments = [_text(issue.get("comment")) for issue in errs]
        comments = [comment for comment in comments if comment]
        if len(comments) <= 1:
            return comments[0] if comments else ""
        return "\n".join(
            f"{index}. {comment}"
            for index, comment in enumerate(comments, start=1)
        )

    def _matching_term_spans(errs, field, display_text):
        matches = []
        for issue in errs:
            term_spans = issue.get("term_spans")
            if term_spans is None and issue.get("category") != "Terminology":
                continue
            if not isinstance(term_spans, dict):
                raise ValueError("Terminology issue is missing term_spans")
            spans = term_spans.get(field)
            if not isinstance(spans, list):
                raise ValueError(f"term_spans.{field} must be a list")
            for span in spans:
                if not isinstance(span, dict):
                    raise ValueError(
                        f"term_spans.{field} entries must be objects"
                    )
                start = span.get("start")
                end = span.get("end")
                span_text = span.get("text")
                if not (
                    type(start) is int
                    and type(end) is int
                    and isinstance(span_text, str)
                    and 0 <= start < end <= len(display_text)
                    and display_text[start:end] == span_text
                ):
                    raise ValueError(
                        f"term_spans.{field} does not match the report baseline"
                    )
                matches.append(span)
        return matches

    def _review_suggestion(seg, entry, errs, is_protected):
        if is_protected:
            return "", "已保护"
        if (
            not review_policy["minor_edits_allowed"]
            and errs
            and not any(
            issue.get("severity") in {"Major", "Critical"}
            for issue in errs
            )
        ):
            return "", "未生成建议，需人工处理"
        reference = reference_suggestions.get(seg["id"])
        if reference is not None:
            return reference, "建议待确认"
        corrected = entry.get("corrected")
        if corrected is not None:
            unresolved = any(
                issue.get("needs_confirmation") is True
                or issue.get("edit") is None
                for issue in errs
            )
            return (
                corrected,
                "部分修正，仍需确认" if unresolved else "可直接采用",
            )
        if errs:
            return "", "未生成建议，需人工处理"
        return "", ""

    review_headers = [
        "Segment ID",
        "原文",
        "原译",
        "AI/建议译文",
        "建议状态",
        "错误类别",
        "严重度",
        "问题说明",
        "审校结论",
        "审校终稿或备注",
    ]
    review_widths = [10, 38, 38, 38, 20, 20, 10, 42, 14, 38]
    status_fills = {
        "可直接采用": _GREEN_LIGHT,
        "建议待确认": PatternFill("solid", fgColor="FFF2CC"),
        "部分修正，仍需确认": _ORANGE,
        "未生成建议，需人工处理": PatternFill("solid", fgColor="F4CCCC"),
        "已保护": PatternFill("solid", fgColor="D9D9D9"),
    }

    wb  = openpyxl.Workbook()
    ws  = wb.active
    ws.title = "LQA Scorecard"
    latest_status = history[-1].get("status") if history else None
    critical_gate_fail = bool(
        (scoring_policy or {}).get("critical_gate")
        and total_counts.get("Critical", 0)
    )
    status = latest_status or (
        "REVIEW_NOT_RUN"
        if score is None
        else ("FAIL" if critical_gate_fail or score < threshold else "PASS")
    )
    score_display = "N/A" if score is None else f"{score:.2f}"

    intro = wb.create_sheet("说明·导读", 0)
    intro_wrap = Alignment(wrap_text=True, vertical="top")
    intro_rows = [
        ("LQE 质检报告 · 新人导读", "", "", ""),
        (
            f"本次结果：{status} ｜ 得分：{score_display} ｜ "
            f"合格线：{threshold}",
            "",
            "",
            "",
        ),
        ("Check scope", scope_summary, "", ""),
        ("", "", "", ""),
        ("三步读报告", "", "", ""),
        ("顺序", "位置", "看什么", "要做什么"),
        (
            "1",
            "说明·导读",
            "报告结构、列含义、状态和审校规则",
            "第一次使用时先读一遍；之后可直接进入正文",
        ),
        (
            "2",
            "LQA Scorecard",
            "PASS/FAIL、分数、错误分布和逐错误明细",
            "先判断整体风险，再按严重度处理下方问题",
        ),
        (
            "3",
            "LQE Results",
            "一段一行的汇总审校区；同段多个问题集中显示",
            "填写审校结论，并在需要时写入终稿或备注",
        ),
        ("", "", "", ""),
        ("LQA Scorecard 怎么读", "", "", ""),
        ("项目", "含义", "怎么看", "注意"),
        (
            "Status",
            "本次质检的最终判定",
            "PASS 表示达到合格线；FAIL 表示未达到或触发关键错误门槛",
            "先看判定，再看分数和错误分布",
        ),
        (
            "Final score",
            "按当前计分规则计算的最终分数",
            "与 Threshold 比较",
            "分数不代替逐条审校",
        ),
        (
            "Threshold",
            "项目合格线",
            "Final score 低于该值通常为 FAIL",
            "Critical 门槛可能直接导致 FAIL",
        ),
        (
            "Error summary",
            "按错误类别和严重度汇总数量与罚分",
            "优先处理 Critical、Major，再处理 Minor",
            "Raw penalty 为原始罚分；Weighted penalty 为加权罚分",
        ),
        (
            "总数 / 重复数",
            "例如 8 / 3 表示共 8 个，其中 3 个按重复错误计",
            "用于理解同类问题是否反复出现",
            "重复数已包含在总数内",
        ),
        ("", "", "", ""),
        ("审校区 10 列说明", "", "", ""),
        ("列", "含义", "审校动作", "补充"),
        (
            "Segment ID",
            "文本段的唯一编号",
            "用编号筛选、搜索和沟通问题",
            "Scorecard 中同段多错误可能出现多行",
        ),
        (
            "原文",
            "待翻译的源语言文本",
            "核对语义、语境、角色和功能",
            "术语问题影响的源词显示为红色字体",
        ),
        (
            "原译",
            "本轮实际送审译文；首轮为输入原译，后续轮为上一轮已应用译文",
            "确认问题是否真实存在",
            "红字表示术语问题词；红色删除线表示删除/替换，重叠时保留删除线",
        ),
        (
            "AI/建议译文",
            "可供审校参考的完整建议译文；可能为空",
            "结合原文、问题说明和项目规则判断是否采用",
            "红色字体表示新增或替换内容；不等于已确认终稿",
        ),
        (
            "建议状态",
            "建议译文的可用程度和确认要求",
            "按下方状态说明决定人工介入程度",
            "它不是错误严重度",
        ),
        (
            "错误类别",
            "问题所属的 LQE 类别",
            "判断问题性质并检查是否还有同类问题",
            "Scorecard 以“父类别 · 子类别”合并显示",
        ),
        (
            "严重度",
            "问题对理解、功能、合规或体验的影响等级",
            "按 Critical → Major → Minor → Neutral 排优先级",
            "它不是 AI 建议置信度",
        ),
        (
            "问题说明",
            "为什么判为问题，以及具体问题位置",
            "据此复核；不同意时在备注中写明理由",
            "同段多问题在 Results 中按编号汇总",
        ),
        (
            "审校结论",
            "审校人员对建议的最终处理选择",
            "从下拉框选择：接受、修改后接受、拒绝、待确认",
            "建议不要留空后直接交付",
        ),
        (
            "审校终稿或备注",
            "最终采用的译文，或拒绝/待确认的说明",
            "修改后接受时填写完整终稿；拒绝或待确认时写明原因",
            "保持占位符、标签和换行结构正确",
        ),
        ("", "", "", ""),
        ("建议状态说明", "", "", ""),
        ("状态", "表示什么", "审校动作", "是否可直接交付"),
        (
            "可直接采用",
            "建议完整，当前没有待确认项",
            "快速核对上下文、术语和格式",
            "仍需审校人员最终确认",
        ),
        (
            "建议待确认",
            "提供了参考译文，但属于宽松生成或存在判断空间",
            "逐句核对后决定接受、修改或拒绝",
            "否",
        ),
        (
            "部分修正，仍需确认",
            "已修正可确定部分，但仍有未解决内容",
            "补全剩余修改，并填写完整终稿",
            "否",
        ),
        (
            "未生成建议，需人工处理",
            "已确认有问题，但没有可靠完整建议",
            "人工改译并填写终稿",
            "否",
        ),
        (
            "已保护",
            "该段受锁定、TM 或其他保护规则约束",
            "不要修改；如认为保护规则有误，单独反馈",
            "按原文档规则处理",
        ),
        ("", "", "", ""),
        ("审校结论说明", "", "", ""),
        ("结论", "何时选择", "需要填写什么", "结果"),
        (
            "接受",
            "建议译文无需修改即可采用",
            "通常无需补写终稿；必要时可备注",
            "采用建议译文",
        ),
        (
            "修改后接受",
            "建议方向正确，但需要人工调整",
            "在“审校终稿或备注”填写完整终稿",
            "采用人工终稿",
        ),
        (
            "拒绝",
            "建议不成立、不合适或原译无需修改",
            "写明理由；如仍需修改，可同时填写替代终稿",
            "不采用当前建议",
        ),
        (
            "待确认",
            "缺少上下文、规则或客户决定，暂时无法定稿",
            "写明待确认点和所需信息",
            "暂不交付",
        ),
        ("", "", "", ""),
        ("阅读与交付提示", "", "", ""),
        (
            "差异标记",
            "原文/原译中的术语问题词为红色；原译差异为红色删除线；建议译文差异为红色字体",
            "差异只帮助定位修改，不代表建议一定正确",
            "",
        ),
        (
            "隐藏数据",
            "LQE Results 隐藏无问题段、同段额外审计行和技术列",
            "需要追溯时可取消隐藏；Scorecard 本身不隐藏行列",
            "",
        ),
        (
            "建议为空",
            "不代表没有问题，而是没有可靠的完整建议",
            "根据问题说明人工改译，并填写终稿",
            "",
        ),
        (
            "交付前",
            "检查所有非保护问题是否已有审校结论",
            "重点复核占位符、标签、数字、术语和换行",
            "",
        ),
    ]
    for row in intro_rows:
        intro.append(row)
        for cell in intro[intro.max_row]:
            cell.alignment = intro_wrap

    intro.merge_cells("A1:D1")
    intro.merge_cells("A2:D2")
    intro.merge_cells("B3:D3")
    section_rows = [5, 11, 19, 32, 40, 47]
    header_rows = [6, 12, 20, 33, 41]
    for row in section_rows:
        intro.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        for cell in intro[row]:
            cell.fill = _DARK_BLUE
            cell.font = _WHITE_FONT
        intro.cell(row, 1).font = Font(color="FFFFFF", bold=True, size=11)
        intro.row_dimensions[row].height = 24
    for row in header_rows:
        for cell in intro[row]:
            cell.fill = _LIGHT_BLUE
            cell.font = _BOLD
            cell.alignment = _CENTER
        intro.row_dimensions[row].height = 24

    intro["A1"].fill = _DARK_BLUE
    intro["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    intro["A1"].alignment = _CENTER
    intro.row_dimensions[1].height = 32
    intro["A2"].fill = _RED if status == "FAIL" else _GREEN
    intro["A2"].font = Font(bold=True, color="FFFFFF")
    intro["A2"].alignment = _CENTER
    intro.row_dimensions[2].height = 24
    intro["A3"].fill = _LIGHT_BLUE
    intro["A3"].font = _BOLD
    intro["B3"].fill = _LIGHT_BLUE
    intro.row_dimensions[3].height = 30

    for row in range(1, intro.max_row + 1):
        if row in section_rows or row in header_rows or row in (1, 2, 3):
            continue
        values = [intro.cell(row, column).value for column in range(1, 5)]
        if not any(value not in (None, "") for value in values):
            intro.row_dimensions[row].height = 8
            continue
        try:
            height, font_size = _fit_wrapped_row(
                [
                    (value, width)
                    for value, width in zip(values, [18, 48, 36, 42])
                ],
                minimum=24,
                context=f"说明·导读 row {row}",
            )
        except ValueError:
            height, font_size = 60, 9
        intro.row_dimensions[row].height = min(height, 72)
        _set_row_font_size(intro, row, 4, font_size)

    for row, fill in [
        (34, _GREEN_LIGHT),
        (35, PatternFill("solid", fgColor="FFF2CC")),
        (36, _ORANGE),
        (37, PatternFill("solid", fgColor="F4CCCC")),
        (38, PatternFill("solid", fgColor="D9D9D9")),
    ]:
        intro.cell(row, 1).fill = fill
        intro.cell(row, 1).font = _BOLD

    for column, width in zip("ABCD", [18, 48, 36, 42]):
        intro.column_dimensions[column].width = width
    intro.freeze_panes = "A6"
    intro.sheet_view.showGridLines = False
    intro.sheet_view.selection[0].activeCell = "A1"
    intro.sheet_view.selection[0].sqref = "A1"
    wb.active = intro

    def _db_row(row, height=14.25):
        ws.row_dimensions[row].height = height
        for col in range(1, 11):
            _s(ws.cell(row=row, column=col), fill=_DARK_BLUE, font=_WHITE_FONT)

    for r in [1, 2, 3]:
        ws.row_dimensions[r].height = 23.25
        _db_row(r, height=23.25)
    ws.merge_cells("A1:J3")
    c = ws["A1"]
    c.value = "LQA Scorecard"
    c.font  = Font(color="FFFFFF", size=16, bold=True)
    c.alignment = _CENTER

    source_lang = state.get("source_lang") or "-"
    target_lang = state.get("target_lang") or "-"
    blocked_segment_count = sum(
        segment.get("input_status") == "blocked" for segment in segments
    )
    reviewable_segment_count = len(segments) - blocked_segment_count
    unverified_target_count = sum(
        any(
            (
                warning.get("code")
                if isinstance(warning, dict)
                else warning
            ) == "UNVERIFIED_TARGET_PROVENANCE"
            for warning in segment.get("input_warnings", [])
        )
        for segment in segments
    )
    info = [
        ("File", Path(state["input_path"]).name, "Wordcount", state.get("wordcount", 0)),
        ("Source language", source_lang, "Target language", target_lang),
        ("Total iterations", len(history), "Threshold", threshold),
        ("Date", date.today().isoformat(), "Terminology check", terminology_status),
    ]
    for ri, (l1, v1, l2, v2) in enumerate(info, start=4):
        _db_row(ri, height=15.0)
        ws.cell(row=ri, column=1, value=l1)
        ws.merge_cells(f"B{ri}:D{ri}")
        c = ws.cell(row=ri, column=2)
        _set_excel_text(c, v1)
        c.fill = _DARK_BLUE; c.font = Font(color="FFFFFF")
        ws.cell(row=ri, column=5, value=l2)
        ws.merge_cells(f"F{ri}:H{ri}")
        c = ws.cell(row=ri, column=6)
        _set_excel_text(c, v2)
        c.fill = _DARK_BLUE; c.font = Font(color="FFFFFF")
    for column, value in ((9, "Check scope"), (10, scope_summary)):
        c = ws.cell(row=4, column=column)
        _set_excel_text(c, value)
        c.fill = _DARK_BLUE
        c.font = _WHITE_FONT
        c.alignment = _LEFT_TOP
    scorecard_context_summary = [
        (
            5,
            "Segment coverage",
            f"Total: {len(segments)} | Reviewable: {reviewable_segment_count} | "
            f"Blocked: {blocked_segment_count}",
        ),
        (
            6,
            "Input warning",
            (
                "UNVERIFIED_TARGET_PROVENANCE: "
                f"{unverified_target_count} segment(s)"
                if unverified_target_count
                else "None"
            ),
        ),
    ]
    for row, label, value in scorecard_context_summary:
        ws.row_dimensions[row].height = 30
        for column, cell_value in ((9, label), (10, value)):
            c = ws.cell(row=row, column=column)
            _set_excel_text(c, cell_value)
            c.fill = _DARK_BLUE
            c.font = _WHITE_FONT
            c.alignment = _LEFT_TOP
    if unverified_target_count:
        ws.cell(row=6, column=10).fill = _ORANGE
        ws.cell(row=6, column=10).font = Font(bold=True)

    ws.row_dimensions[8].height = 6

    _db_row(9)
    ws.merge_cells("A9:J9")
    c = ws["A9"]
    c.value = "LQA results"
    c.alignment = _LEFT_MID

    for row, height in [(10, 34.0), (11, 34.0), (12, 34.0)]:
        ws.row_dimensions[row].height = height

    for row, label, val, val_fill, val_font in [
        (10, "Status",      status,          _RED if status == "FAIL" else _GREEN,
                                             Font(color="FFFFFF", bold=True)),
        (11, "Final score", "N/A" if score is None else round(score, 4), _ORANGE, Font(bold=True)),
        (12, "Threshold",   threshold,       _ORANGE, Font(bold=True)),
    ]:
        c = ws.cell(row=row, column=1, value=label)
        _s(c, fill=_LIGHT_BLUE, align=_CENTER)
        c = ws.cell(row=row, column=2, value=val)
        _s(c, fill=val_fill, font=val_font, align=_CENTER)

    ws.merge_cells("C10:J12")
    c = ws["C10"]
    c.value = "Overall feedback"
    _s(c, fill=_LIGHT_BLUE, align=_CENTER)

    cur_row = 13
    if len(history) > 1:
        ws.row_dimensions[cur_row].height = 6
        cur_row += 1
        _db_row(cur_row)
        ws.merge_cells(f"A{cur_row}:J{cur_row}")
        c = ws.cell(row=cur_row, column=1, value="Iteration log")
        c.alignment = _LEFT_MID
        cur_row += 1
        for col, hdr in [(1,"Iteration"),(2,"Score"),(3,"Errors found"),(4,"Corrections applied")]:
            c = ws.cell(row=cur_row, column=col, value=hdr)
            _s(c, fill=_LIGHT_BLUE, align=_CENTER, font=_BOLD)
        ws.row_dimensions[cur_row].height = 14.25
        cur_row += 1
        for entry in history:
            s  = entry.get("score")
            ec = sum(len(e.get("errors", [])) for e in entry["errors"])
            for col, val in [
                (1, f"Iter {entry['iteration']}"),
                (2, round(s, 2) if s is not None else ""),
                (3, ec),
                (4, entry.get("corrections_count", 0)),
            ]:
                c = ws.cell(row=cur_row, column=col, value=val)
                _s(c, fill=_ORANGE, align=_CENTER)
            ws.row_dimensions[cur_row].height = 14.25
            cur_row += 1

    ws.row_dimensions[cur_row].height = 6
    cur_row += 1

    _db_row(cur_row)
    ws.merge_cells(f"A{cur_row}:J{cur_row}")
    c = ws.cell(row=cur_row, column=1, value="Error summary")
    c.alignment = _LEFT_MID
    cur_row += 1

    summary_headers = [
        "Error category",
        "Weight",
        "Neutral\n(total / repeated)",
        "Minor\n(total / repeated)",
        "Major\n(total / repeated)",
        "Critical\n(total / repeated)",
        "Raw penalty",
        "Weighted penalty",
    ]
    ws.row_dimensions[cur_row].height = 30
    for col, val in enumerate(summary_headers, start=1):
        c = ws.cell(row=cur_row, column=col, value=val)
        _s(c, fill=_LIGHT_BLUE, font=_BOLD, align=_CENTER)
    cur_row += 1

    def _severity_pair(counts, repeated, severity):
        return f"{counts.get(severity, 0)} / {repeated.get(severity, 0)}"

    ws.row_dimensions[cur_row].height = 14.25
    total_summary = [
        "TOTAL",
        None,
        _severity_pair(total_counts, total_rep, "Neutral"),
        _severity_pair(total_counts, total_rep, "Minor"),
        _severity_pair(total_counts, total_rep, "Major"),
        _severity_pair(total_counts, total_rep, "Critical"),
        total_raw,
        round(total_weighted, 2),
    ]
    for col, val in enumerate(total_summary, start=1):
        c = ws.cell(row=cur_row, column=col, value=val)
        _s(c, fill=_ORANGE, align=_CENTER)
    cur_row += 1

    for cat in categories:
        counts = cat_counts[cat]
        r = raw_points(counts, scorecard_profile, severity_points)
        w = weighted_points(cat, counts, scorecard_profile, severity_points)
        ws.row_dimensions[cur_row].height = 30 if len(cat) > 12 else 20
        rep = rep_counts[cat]
        category_summary = [
            cat,
            scorecard_category_weight(cat, scorecard_profile),
            _severity_pair(counts, rep, "Neutral"),
            _severity_pair(counts, rep, "Minor"),
            _severity_pair(counts, rep, "Major"),
            _severity_pair(counts, rep, "Critical"),
            r,
            round(w, 2),
        ]
        for col, val in enumerate(category_summary, start=1):
            c = ws.cell(row=cur_row, column=col, value=val)
            _s(c, align=_CENTER)
        cur_row += 1

    ws.row_dimensions[cur_row].height = 6
    cur_row += 1

    scorecard_widths = [18, *review_widths[1:]]
    for column, width in enumerate(scorecard_widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width

    scorecard_detail_header = cur_row
    ws.row_dimensions[cur_row].height = 14.25
    for col, hdr in enumerate(review_headers, start=1):
        c = ws.cell(row=cur_row, column=col, value=hdr)
        _s(c, fill=_DARK_BLUE, font=_WHITE_FONT, align=_CENTER)
    cur_row += 1

    scorecard_reviewer_rows = []
    for dr in detail_rows:
        seg = seg_map[dr["seg_id"]]
        is_protected = dr["seg_id"] in all_protected_ids
        entry = dr["entry"]
        errs = [] if is_protected else entry.get("errors", [])
        suggestion, suggestion_status = _review_suggestion(
            seg,
            entry,
            errs,
            is_protected,
        )
        suggestion_for_diff = (
            suggestion
            if suggestion_status in {
                "可直接采用",
                "建议待确认",
                "部分修正，仍需确认",
            }
            else None
        )
        rich_source, rich_original, rich_suggestion = build_review_rich_texts(
            dr["source"],
            dr["original"],
            suggestion_for_diff,
            source_spans=_matching_term_spans(
                [dr["issue"]],
                "source",
                dr["source"],
            ),
            target_spans=_matching_term_spans(
                [dr["issue"]],
                "target",
                dr["original"],
            ),
        )
        category = " · ".join(
            value
            for value in (dr["parent"], dr["category"])
            if value
        )
        detail_values = [
            dr["seg_id"],
            rich_source,
            rich_original,
            rich_suggestion,
            suggestion_status,
            category,
            dr["severity"],
            dr["comment"],
            "",
            "",
        ]
        for col, val in enumerate(detail_values, start=1):
            c = ws.cell(row=cur_row, column=col)
            _set_excel_text(c, val)
            _s(
                c,
                align=(
                    _CENTER
                    if col in (1, 5, 7, 9)
                    else _LEFT_TOP
                ),
            )
        if suggestion_status:
            ws.cell(row=cur_row, column=5).fill = status_fills[
                suggestion_status
            ]
        try:
            row_height, row_font_size = _fit_wrapped_row(
                [
                    (value, scorecard_widths[index])
                    for index, value in enumerate(detail_values)
                ],
                minimum=30,
                context=f"LQA Scorecard row {cur_row}",
            )
        except ValueError:
            row_height, row_font_size = 120, 9
        ws.row_dimensions[cur_row].height = min(row_height, 120)
        _set_row_font_size(ws, cur_row, len(detail_values), row_font_size)
        scorecard_reviewer_rows.append(cur_row)
        cur_row += 1

    scorecard_validation = DataValidation(
        type="list",
        formula1='"接受,修改后接受,拒绝,待确认"',
        allow_blank=True,
    )
    ws.add_data_validation(scorecard_validation)
    if scorecard_reviewer_rows:
        scorecard_validation.add(
            f"I{min(scorecard_reviewer_rows)}:"
            f"I{max(scorecard_reviewer_rows)}"
        )
        ws.auto_filter.ref = (
            f"A{scorecard_detail_header}:J{max(scorecard_reviewer_rows)}"
        )

    ws2 = wb.create_sheet("LQE Results")
    ws2.freeze_panes = "A2"
    ws2.sheet_view.showGridLines = False

    def _fmt_errors(errs):
        return "\n".join(issue_detail(error) for error in errs)

    _WRAP_TOP = Alignment(wrap_text=True, vertical="top")
    report_headers, report_rows = _report_source_table(state)

    visible_headers = review_headers
    used_headers = set(visible_headers)
    reserved_technical_headers = {
        "处理方式",
        *AUDIT_HEADER_BASES.values(),
        "术语原文（结构化）",
        "术语库译文（结构化）",
        "错误详情",
        "Protected",
        "Protection Evidence",
        "Segment Key",
        "Input Status",
        "Content Type",
        "Context Status",
        "Context JSON",
        "Context Provenance JSON",
        "Context Digest",
        "Resolved Constraints JSON",
        "Capability Resolution Digest",
        "Project Asset Snapshot Digest",
        "LQE_Iter",
    }

    def _source_header(base: str) -> str:
        candidate = base
        suffix = 2
        if candidate in used_headers or candidate in reserved_technical_headers:
            candidate = f"{base}（原始数据）"
        while candidate in used_headers or candidate in reserved_technical_headers:
            candidate = f"{base}（原始数据 {suffix}）"
            suffix += 1
        used_headers.add(candidate)
        return candidate

    technical_source_headers = [
        _source_header(str(header))
        for header in report_headers
    ]
    processing_header = "处理方式"
    segment_audit_header = AUDIT_HEADER_BASES["segment_id"]
    issue_number_header = AUDIT_HEADER_BASES["issue_number"]
    review_status_header = AUDIT_HEADER_BASES["review_status"]
    edit_status_header = AUDIT_HEADER_BASES["edit_status"]
    check_source_header = AUDIT_HEADER_BASES["check_source"]

    technical_headers = technical_source_headers + [
        processing_header,
        segment_audit_header,
        issue_number_header,
        review_status_header,
        edit_status_header,
        check_source_header,
        "术语原文（结构化）",
        "术语库译文（结构化）",
        "错误详情",
        "Protected",
        "Protection Evidence",
        "Segment Key",
        "Input Status",
        "Content Type",
        "Context Status",
        "Context JSON",
        "Context Provenance JSON",
        "Context Digest",
        "Resolved Constraints JSON",
        "Capability Resolution Digest",
        "Project Asset Snapshot Digest",
        "LQE_Iter",
    ]
    ws2_headers = visible_headers + technical_headers
    visible_widths = review_widths
    for column, width in enumerate(visible_widths, start=1):
        ws2.column_dimensions[get_column_letter(column)].width = width
    technical_start = len(visible_headers) + 1
    for column in range(technical_start, len(ws2_headers) + 1):
        dimension = ws2.column_dimensions[get_column_letter(column)]
        dimension.width = 16
        dimension.hidden = True

    for ci, header in enumerate(ws2_headers, start=1):
        cell = ws2.cell(row=1, column=ci)
        _set_excel_text(cell, header)
        _s(cell, fill=_DARK_BLUE, font=_WHITE_FONT, align=_CENTER)
    ws2.row_dimensions[1].height = 24
    ws2.auto_filter.ref = f"A1:J1"

    ri = 2
    reviewer_rows = []
    for segment_index, seg in enumerate(segments):
        raw_row = report_rows[segment_index]
        is_protected = seg["id"] in all_protected_ids
        entry = _review_entry(seg)
        errs = [] if is_protected else entry.get("errors", [])
        suggestion, suggestion_status = _review_suggestion(
            seg,
            entry,
            errs,
            is_protected,
        )
        source_text = seg.get("source", "")
        original_text = current_review_targets[seg["id"]]
        suggestion_for_diff = (
            suggestion
            if suggestion_status in {
                "可直接采用",
                "建议待确认",
                "部分修正，仍需确认",
            }
            else None
        )
        rich_source, rich_original, rich_suggestion = build_review_rich_texts(
            source_text,
            original_text,
            suggestion_for_diff,
            source_spans=_matching_term_spans(errs, "source", source_text),
            target_spans=_matching_term_spans(errs, "target", original_text),
        )
        visible_data = [
            seg["id"],
            rich_source,
            rich_original,
            rich_suggestion,
            suggestion_status,
            _review_summary(errs, "category"),
            _review_summary(errs, "severity"),
            _problem_summary(errs),
            "",
            "",
        ]
        first_row = ri
        reviewer_rows.append(first_row)
        row_issues = errs or [None]
        for issue_index, issue in enumerate(row_issues, start=1):
            review_status, edit_status, check_source = _issue_review_columns(
                issue,
                seg["id"],
            )
            processing = _issue_processing_label(
                issue,
                entry,
                protected=is_protected,
            )
            term_fields = terminology_issue_fields(issue)
            technical_data = (
                list(raw_row)
                + [
                    processing,
                    seg["id"],
                    issue_index if issue is not None else "",
                    review_status,
                    edit_status,
                    check_source,
                    term_fields["term_source"] if term_fields else "",
                    (
                        "\n".join(term_fields["expected_targets"])
                        if term_fields
                        else ""
                    ),
                    _fmt_errors([issue]) if issue is not None else "",
                    "Yes" if is_protected else "No",
                    _protection_evidence(seg, all_protected_ids),
                    *context_audit_values(state, seg),
                    seg.get("iter", 0),
                ]
            )
            row_data = (
                visible_data if issue_index == 1 else [""] * len(visible_headers)
            ) + technical_data
            for ci, value in enumerate(row_data, start=1):
                cell = ws2.cell(row=ri, column=ci)
                _set_excel_text(cell, value)
                cell.alignment = _WRAP_TOP
                if ci <= len(visible_headers) and suggestion_status:
                    cell.fill = status_fills[suggestion_status]
            if issue_index == 1:
                try:
                    row_height, row_font_size = _fit_wrapped_row(
                        [
                            (value, visible_widths[index])
                            for index, value in enumerate(visible_data)
                        ],
                        minimum=30,
                        context=f"LQE Results row {ri}",
                    )
                except ValueError:
                    row_height, row_font_size = 120, 9
                row_height = min(row_height, 120)
                row_span = 1
                ws2.row_dimensions[ri].height = row_height
                _set_row_font_size(
                    ws2,
                    ri,
                    len(ws2_headers),
                    row_font_size,
                )
                if not (errs or is_protected or suggestion):
                    ws2.row_dimensions[ri].hidden = True
                    ws2.row_dimensions[ri].outlineLevel = 1
            else:
                row_span = 1
                ws2.row_dimensions[ri].hidden = True
                ws2.row_dimensions[ri].outlineLevel = 1
            ri += row_span

    decision_validation = DataValidation(
        type="list",
        formula1='"接受,修改后接受,拒绝,待确认"',
        allow_blank=True,
    )
    decision_validation.error = "请选择下拉列表中的审校结论。"
    decision_validation.errorTitle = "无效审校结论"
    decision_validation.prompt = "请选择接受、修改后接受、拒绝或待确认。"
    decision_validation.promptTitle = "审校结论"
    ws2.add_data_validation(decision_validation)
    if reviewer_rows:
        decision_validation.add(
            f"I{min(reviewer_rows)}:I{max(reviewer_rows)}"
        )
        ws2.auto_filter.ref = f"A1:J{max(reviewer_rows)}"

    ws.sheet_view.tabSelected = False
    ws2.sheet_view.tabSelected = False
    ws2.sheet_view.selection[0].activeCell = "A1"
    ws2.sheet_view.selection[0].sqref = "A1"
    intro.sheet_view.tabSelected = True
    wb.active = intro

    if report_contract_results is not None:
        attach_report_contract(wb, state, report_contract_results)
    wb.save(str(out_path))
    if announce:
        print(f"[lqe_io] Output → {out_path}")


def _cmd_write_locked(
    args,
    state_path: Path,
    state: dict,
    segments: list[dict],
    manifest: dict | None,
    revalidate,
):
    require_current_job_runtime(state, "write")
    errors_path = Path(args.errors)
    final_errors_data = read_json(errors_path)
    original_errors_data = deepcopy(final_errors_data)
    _validate_scope_or_exit(
        state,
        final_errors_data,
        issues_key="errors",
        label=errors_path.name,
        command="write",
    )
    protected_ids = _state_protected_ids(state)
    scrubbed = _scrub_protected_entries(final_errors_data, protected_ids)
    try:
        scoring_policy = resolve_scoring_policy(
            state,
            scoring_policy_overrides(args),
        )
        scorecard_profile = load_scorecard_profile(
            scoring_policy["scorecard_profile"]
        )
        validation_messages = _validate_errors(
            final_errors_data,
            {segment["id"] for segment in state.get("segments", [])},
            scorecard_profile,
        )
        for message in validation_messages:
            print(f"[validate] {message}")
        final_errors_data = _verify_result_payload_with_segments(
            state,
            segments,
            manifest,
            final_errors_data,
            errors_path,
            command="write",
        )
        computation = score_errors(
            state,
            final_errors_data,
            scoring_policy,
            protected_ids=protected_ids,
        )
    except (CheckFormatError, ValueError) as exc:
        raise SystemExit(f"[write] {exc}") from exc
    final_errors_data = computation["annotated_errors"]
    score_result = computation["output"]
    score = score_result["score"]
    suggestion_path = state_path.parent / ARTIFACT_NAME
    original_suggestion_data = (
        read_json(suggestion_path) if suggestion_path.is_file() else None
    )
    try:
        reference_suggestions = load_reference_suggestions(
            state_path.parent,
            segments,
            manifest,
            final_errors_data,
            review_policy=get_review_policy(state),
        )
    except (CheckFormatError, OSError, ValueError) as exc:
        raise SystemExit(f"[write] {exc}") from exc
    supplied_score = float(args.score)
    if not math.isclose(supplied_score, score, abs_tol=0.005):
        print(
            f"[write] supplied score {supplied_score:g} differs from "
            f"recomputed score {score:g}; using recomputed score",
            file=sys.stderr,
        )
    state["threshold"] = scoring_policy["threshold"]
    state["scoring_policy"] = scoring_policy
    state["pending_recheck"] = False

    final_entry = {
        "iteration": state.get("iteration", 0),
        "score": score,
        "status": score_result["status"],
        "errors": final_errors_data,
        "corrections_count": 0,
        "protected_ids": sorted(protected_ids),
        "review_targets": {
            str(segment["id"]): current_target(segment)
            for segment in state["segments"]
        },
    }

    history = state.get("error_history", [])
    final_iter = state.get("iteration", 0)
    skipped_corrections = [
        skipped
        for entry in history
        if entry.get("iteration") == final_iter
        for skipped in entry.get("skipped_corrections", [])
    ]
    if skipped_corrections:
        final_entry["skipped_corrections"] = skipped_corrections
    history = [
        entry for entry in history if entry.get("iteration") != final_iter
    ]
    history.append(final_entry)
    state["error_history"] = history

    out_path = state_path.parent / (_job_label(state_path) + "_lqe.xlsx")
    try:
        validate_artifact_paths(
            {"LQE report": out_path},
            {
                "state": state_path,
                "errors": errors_path,
                **(
                    {"reference suggestions": suggestion_path}
                    if suggestion_path.is_file()
                    else {}
                ),
                **state_reference_paths(state),
            },
            context="write",
        )
    except ValueError as exc:
        raise SystemExit(f"[write] {exc}") from exc
    with tempfile.NamedTemporaryFile(
        prefix=f".{out_path.stem}.",
        suffix=out_path.suffix,
        dir=out_path.parent,
        delete=False,
    ) as staging_file:
        staging_path = Path(staging_file.name)
    try:
        try:
            _build_xlsx(
                state,
                history,
                score,
                scoring_policy["threshold"],
                staging_path,
                scoring_policy["scorecard_profile"],
                announce=False,
                scoring_policy=scoring_policy,
                scoring_computation=computation,
                report_contract_results=final_errors_data,
                reference_suggestions=reference_suggestions,
            )
        except ValueError as exc:
            raise SystemExit(f"[write] {exc}") from exc
        revalidate()
        _assert_json_unchanged(
            errors_path,
            original_errors_data,
            label="errors input",
        )
        if original_suggestion_data is not None:
            _assert_json_unchanged(
                suggestion_path,
                original_suggestion_data,
                label="reference suggestions input",
            )
        _publish_write_transaction(
            state_path,
            state,
            errors_path,
            final_errors_data,
            out_path,
            staging_path,
            publish_errors=final_errors_data != original_errors_data,
            manifest=manifest,
        )
    finally:
        staging_path.unlink(missing_ok=True)

    if scrubbed:
        print(f"[write] scrubbed {scrubbed} protected-segment issue(s) from {errors_path}")
    print(f"[lqe_io] Output → {out_path}")


def cmd_write(args):
    from lqe_chunk import verification_generation_lease

    state_path = Path(args.state)
    try:
        require_current_job_runtime(read_json(state_path), "write")
        with verification_generation_lease(
            state_path,
            exclusive=True,
        ) as (state, segments, manifest, revalidate):
            _cmd_write_locked(
                args,
                state_path,
                state,
                segments,
                manifest,
                revalidate,
            )
    except ValueError as exc:
        raise SystemExit(f"[write] {exc}") from exc


# ── pre-check（实现在 lqe_checks.py）─────────────────────────────────────────

def cmd_pre_check(args):
    from lqe_checks import run_pre_check
    from lqe_split_contract import generation_lock

    state_path = Path(args.state)
    try:
        with generation_lock(state_path.parent / "chunks", exclusive=True):
            require_current_job_runtime(read_json(state_path), "pre-check")
            run_pre_check(
                state_path,
                Path(args.out) if args.out else None,
            )
    except (OSError, ValueError) as exc:
        raise SystemExit(f"[pre-check] {exc}") from exc


# ── export ───────────────────────────────────────────────────────────────────

def _export_sdlxliff_xlsx(
    state_path: Path,
    state: dict,
    result_entries: dict,
    *,
    out_path: Path | None = None,
) -> Path:
    headers, rows = _report_source_table(state)
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Corrected"
    worksheet.append(headers)
    for index, segment in enumerate(state.get("segments") or []):
        source_row = rows[index]
        entry = result_entries[segment["id"]]
        protected = bool(segment.get("protected")) or (
            _processing_label(entry) == "已保护，不修改"
        )
        corrected = entry.get("corrected")
        output_row = list(source_row)
        if not protected and corrected is not None:
            output_row[4] = corrected
        row_number = worksheet.max_row + 1
        for column, value in enumerate(output_row, start=1):
            _set_excel_text(
                worksheet.cell(row=row_number, column=column),
                value,
            )
    out_path = out_path or state_path.parent / (
        _job_label(state_path) + "_corrected.xlsx"
    )
    workbook.save(out_path)
    workbook.close()
    return out_path


def _verification_segments(state_path: Path, state: dict) -> list[dict]:
    from lqe_chunk import load_verification_segments

    try:
        return load_verification_segments(
            state_path,
            state=state,
        )
    except (OSError, ValueError) as exc:
        raise CheckFormatError(str(exc)) from exc


def _verify_result_payload_with_segments(
    state: dict,
    segments: list[dict],
    manifest: dict | None,
    errors_data: list,
    errors_path: Path,
    *,
    command: str,
) -> list[dict]:
    _validate_scope_or_exit(
        state,
        errors_data,
        issues_key="errors",
        label=errors_path.name,
        command=command,
    )
    try:
        if requires_bound_artifacts(state):
            contract_path = result_contract_path(errors_path)
            if manifest is None or not contract_path.is_file():
                raise CheckFormatError(
                    f"{errors_path.name}: bound result contract is required"
                )
            validate_result_contract(
                read_json(contract_path),
                manifest,
                errors_data,
                label=errors_path.name,
            )
        bound = requires_bound_artifacts(state)
        return verify_results(
            segments,
            errors_data,
            str(errors_path),
            allow_internal_provenance=bound,
            require_internal_provenance=bound,
            review_policy=get_review_policy(state),
        )
    except (CheckFormatError, OSError, ValueError) as exc:
        raise SystemExit(f"[{command}] {exc}") from exc


def _verify_result_payload(
    state_path: Path,
    state: dict,
    errors_data: list,
    errors_path: Path,
    *,
    command: str,
) -> tuple[list[dict], list[dict]]:
    from lqe_chunk import verification_generation_lease

    try:
        with verification_generation_lease(
            state_path,
            exclusive=False,
        ) as (live_state, segments, manifest, _):
            if live_state != state:
                raise CheckFormatError("state changed while loading results")
            verified = _verify_result_payload_with_segments(
                live_state,
                segments,
                manifest,
                errors_data,
                errors_path,
                command=command,
            )
            return segments, verified
    except (OSError, ValueError) as exc:
        raise SystemExit(f"[{command}] {exc}") from exc


def _state_no_header(state: dict) -> bool:
    if "no_header" in state:
        value = state["no_header"]
        if not isinstance(value, bool):
            raise ValueError("state.no_header must be a boolean")
        return value

    source_column = state.get("source_col")
    return isinstance(source_column, int) or (
        isinstance(source_column, str) and source_column.isdigit()
    )


def _state_column_index(state: dict, field: str) -> int:
    value = state.get(field)
    if _state_no_header(state) or isinstance(value, int):
        try:
            return int(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"cannot locate {field} {value!r}") from exc
    headers = state.get("headers") or []
    if value not in headers:
        raise ValueError(f"cannot locate {field} {value!r}")
    return headers.index(value)


def _validate_legacy_tabular_rows(
    state: dict,
    segments: list[dict],
    row_at,
) -> None:
    source_index = _state_column_index(state, "source_col")
    target_index = _state_column_index(state, "target_col")
    for segment in segments:
        row = row_at(segment)
        if row is None:
            raise ValueError(
                f"source row for segment {segment['id']} is missing"
            )
        source = _text(_cell(row, source_index))
        target = _text(_cell(row, target_index))
        if source != segment.get("source", "") or target != segment.get(
            "target", ""
        ):
            raise ValueError(
                f"source row for segment {segment['id']} changed after read"
            )


def _validate_export_source_digest(state: dict, source_path: Path) -> str:
    digest = file_sha256(source_path)
    expected = state.get("input_sha256")
    if expected is not None:
        if not isinstance(expected, str) or not re.fullmatch(
            r"[0-9a-f]{64}", expected
        ):
            raise ValueError("state.input_sha256 is invalid")
        if digest != expected:
            raise ValueError(f"source input changed after read: {source_path}")
    return digest


def _validate_sdl_source_snapshot(state: dict) -> dict[Path, str]:
    manifest_path = state.get("source_manifest_path")
    if not isinstance(manifest_path, str) or not manifest_path.strip():
        raise ValueError("SDLXLIFF state has no source_manifest_path")
    manifest = read_json(manifest_path)
    manifest_files = manifest.get("files") if isinstance(manifest, dict) else None
    if not isinstance(manifest_files, list):
        raise ValueError("SDLXLIFF source manifest files must be an array")

    input_root_value = state.get("input_path")
    if not isinstance(input_root_value, str) or not input_root_value.strip():
        raise ValueError("SDLXLIFF state has no input_path")
    input_root = Path(input_root_value)
    if input_root.is_file():
        live_paths = [input_root]
        relative_paths = [input_root.name]
    elif input_root.is_dir():
        live_paths = sorted(
            (
                path
                for path in input_root.rglob("*")
                if path.is_file() and path.suffix.casefold() == ".sdlxliff"
            ),
            key=lambda path: path.relative_to(input_root).as_posix(),
        )
        relative_paths = [
            path.relative_to(input_root).as_posix() for path in live_paths
        ]
    else:
        raise ValueError(f"SDLXLIFF input path is missing: {input_root}")

    recorded_paths = state.get("input_paths")
    if not isinstance(recorded_paths, list) or not all(
        isinstance(value, str) and value.strip() for value in recorded_paths
    ):
        raise ValueError("SDLXLIFF state.input_paths must be a string array")
    if [str(path.resolve()) for path in live_paths] != [
        str(Path(value).resolve()) for value in recorded_paths
    ]:
        raise ValueError("SDLXLIFF source file set changed after read")

    manifest_by_path = {}
    for index, item in enumerate(manifest_files):
        if not isinstance(item, dict):
            raise ValueError(
                f"SDLXLIFF source manifest files[{index}] must be an object"
            )
        relative_path = item.get("relative_path")
        digest = item.get("sha256")
        if not isinstance(relative_path, str) or not relative_path:
            raise ValueError(
                f"SDLXLIFF source manifest files[{index}].relative_path is invalid"
            )
        if not isinstance(digest, str) or not re.fullmatch(r"[0-9a-f]{64}", digest):
            raise ValueError(
                f"SDLXLIFF source manifest files[{index}].sha256 is invalid"
            )
        if relative_path in manifest_by_path:
            raise ValueError(
                f"SDLXLIFF source manifest has duplicate path: {relative_path}"
            )
        manifest_by_path[relative_path] = digest
    if set(manifest_by_path) != set(relative_paths):
        raise ValueError("SDLXLIFF source manifest file set is inconsistent")

    snapshot = {}
    for path, relative_path in zip(live_paths, relative_paths):
        digest = file_sha256(path)
        if digest != manifest_by_path[relative_path]:
            raise ValueError(f"SDLXLIFF source changed after read: {path}")
        snapshot[path.resolve()] = digest
    return snapshot


def _recheck_sdl_source_snapshot(
    state: dict,
    snapshot: dict[Path, str],
) -> None:
    try:
        current = _validate_sdl_source_snapshot(state)
    except ValueError as exc:
        raise ValueError(
            f"SDLXLIFF source changed during export: {exc}"
        ) from exc
    if current != snapshot:
        raise ValueError("SDLXLIFF source changed during export")


def _validate_export_paths(
    state_path: Path,
    state: dict,
    out_path: Path,
    errors_path: Path | None,
) -> None:
    protected_inputs = {
        "state": state_path,
        **state_reference_paths(state),
    }
    if errors_path is not None:
        protected_inputs["errors"] = errors_path
        if requires_bound_artifacts(state):
            protected_inputs["errors contract"] = result_contract_path(
                errors_path
            )
    validate_artifact_paths(
        {"corrected export": out_path},
        protected_inputs,
        context="export",
    )


def _cmd_export_locked(
    args,
    state_path: Path,
    state: dict,
    segments: list[dict],
    manifest: dict | None,
    revalidate,
):
    require_current_job_runtime(state, "export")
    errors_path = Path(args.errors) if getattr(args, "errors", None) else None
    overlay_entries = None
    original_overlay = None
    if errors_path is not None:
        raw_entries = read_json(errors_path)
        original_overlay = deepcopy(raw_entries)
        _validate_scope_or_exit(
            state,
            raw_entries,
            issues_key="errors",
            label=errors_path.name,
            command="export",
        )
        _scrub_protected_entries(raw_entries, _state_protected_ids(state))
        overlay_entries = _verify_result_payload_with_segments(
            state,
            segments,
            manifest,
            raw_entries,
            errors_path,
            command="export",
        )

    def revalidate_inputs():
        revalidate()
        if errors_path is not None:
            _assert_json_unchanged(
                errors_path,
                original_overlay,
                label="errors input",
            )
    seg_map = {s["id"]: s for s in segments}
    result_entries = {
        segment["id"]: {
            "errors": [],
            "corrected": (
                current_target(segment)
                if current_target(segment) != segment.get("target", "")
                else None
            ),
        }
        for segment in segments
    }

    if overlay_entries is not None:
        for e in overlay_entries:
            seg = seg_map.get(e["id"])
            if seg is None:
                continue
            baseline = result_entries[e["id"]]["corrected"]
            result_entries[e["id"]] = {
                **e,
                "corrected": (
                    e.get("corrected")
                    if e.get("corrected") is not None
                    else baseline
                ),
            }

    counts = {
        "建议修改": 0,
        "需要人工确认": 0,
        "保持原译": 0,
        "已保护": 0,
    }

    def export_kind(segment):
        if segment.get("protected") or segment.get("input_status") == "blocked":
            return "已保护"
        label = _processing_label(result_entries[segment["id"]])
        if label == "已保护，不修改":
            return "已保护"
        if label in ("建议修改", "需要人工确认"):
            return label
        return "保持原译"

    def print_summary(out_path):
        print(
            f"[export] 建议修改 {counts['建议修改']} / "
            f"需要人工确认 {counts['需要人工确认']} / "
            f"保持原译 {counts['保持原译']} / "
            f"已保护 {counts['已保护']} → {out_path}"
        )

    if state.get("input_format") == "sdlxliff":
        for segment in segments:
            counts[export_kind(segment)] += 1
        out_path = state_path.parent / (
            _job_label(state_path) + "_corrected.xlsx"
        )
        staged = None
        try:
            _validate_export_paths(
                state_path, state, out_path, errors_path
            )
            source_snapshot = _validate_sdl_source_snapshot(state)
            with tempfile.NamedTemporaryFile(
                dir=out_path.parent,
                prefix=f".{out_path.stem}.",
                suffix=out_path.suffix,
                delete=False,
            ) as handle:
                staged = Path(handle.name)
            _export_sdlxliff_xlsx(
                state_path,
                state,
                result_entries,
                out_path=staged,
            )
            _recheck_sdl_source_snapshot(state, source_snapshot)
            revalidate_inputs()
            publish_replacement_transaction([(staged, out_path)])
        except (OSError, ValueError) as exc:
            raise SystemExit(f"[export] {exc}") from exc
        finally:
            if staged is not None:
                staged.unlink(missing_ok=True)
        print_summary(out_path)
        return

    try:
        no_header = _state_no_header(state)
        ti = _state_column_index(state, "target_col")
    except ValueError as exc:
        print(f"[export] {exc}", file=sys.stderr)
        sys.exit(1)

    src_path = Path(state["input_path"])
    if src_path.suffix.lower() in (".csv", ".tsv"):
        delim = "\t" if src_path.suffix.lower() == ".tsv" else ","
        out_path = state_path.parent / (
            _job_label(state_path) + "_corrected" + src_path.suffix.lower()
        )
        enc = "utf-8-sig" if src_path.suffix.lower() == ".csv" else "utf-8"
        staged = None
        try:
            _validate_export_paths(
                state_path, state, out_path, errors_path
            )
            source_digest = _validate_export_source_digest(state, src_path)
            raw_rows = list(
                csv.reader(
                    io.StringIO(src_path.read_bytes().decode("utf-8-sig")),
                    delimiter=delim,
                )
            )
            offset = 0 if no_header else 1
            if state.get("input_sha256") is None:
                _validate_legacy_tabular_rows(
                    state,
                    segments,
                    lambda segment: (
                        raw_rows[
                            offset
                            + int(
                                segment.get(
                                    "row_index", segment.get("id", 0)
                                )
                            )
                        ]
                        if 0
                        <= offset
                        + int(
                            segment.get("row_index", segment.get("id", 0))
                        )
                        < len(raw_rows)
                        else None
                    ),
                )
            for seg in segments:
                row_idx = offset + int(
                    seg.get("row_index", seg.get("id", 0))
                )
                if row_idx < 0 or row_idx >= len(raw_rows):
                    raise ValueError(
                        f"source row for segment {seg['id']} is missing"
                    )
                row = raw_rows[row_idx]
                kind = export_kind(seg)
                corrected = result_entries[seg["id"]].get("corrected")
                if (
                    kind != "已保护"
                    and corrected is not None
                    and ti < len(row)
                ):
                    row[ti] = corrected
                counts[kind] += 1
            with tempfile.NamedTemporaryFile(
                mode="w",
                newline="",
                encoding=enc,
                dir=out_path.parent,
                prefix=f".{out_path.name}.",
                suffix=".tmp",
                delete=False,
            ) as handle:
                staged = Path(handle.name)
                csv.writer(handle, delimiter=delim).writerows(raw_rows)
            if file_sha256(src_path) != source_digest:
                raise ValueError(
                    f"source input changed during export: {src_path}"
                )
            revalidate_inputs()
            publish_replacement_transaction([(staged, out_path)])
        except (OSError, ValueError) as exc:
            raise SystemExit(f"[export] {exc}") from exc
        finally:
            if staged is not None:
                staged.unlink(missing_ok=True)
        print_summary(out_path)
        return

    out_path = state_path.parent / (_job_label(state_path) + "_corrected.xlsx")
    staged = None
    workbook = None
    try:
        _validate_export_paths(state_path, state, out_path, errors_path)
        source_digest = _validate_export_source_digest(state, src_path)
        workbook = (
            workbook_for_corrected_export(src_path)
            if src_path.suffix.casefold() == ".xls"
            else openpyxl.load_workbook(str(src_path))
        )
        sheet_name = state.get("sheet_name")
        worksheet = (
            workbook[sheet_name]
            if sheet_name in workbook.sheetnames
            else workbook.active
        )
        start_row = 1 if no_header else 2
        if state.get("input_sha256") is None:
            _validate_legacy_tabular_rows(
                state,
                segments,
                lambda segment: (
                    [
                        worksheet.cell(
                            row=start_row
                            + int(
                                segment.get(
                                    "row_index", segment.get("id", 0)
                                )
                            ),
                            column=column,
                        ).value
                        for column in range(1, worksheet.max_column + 1)
                    ]
                    if start_row
                    <= start_row
                    + int(
                        segment.get("row_index", segment.get("id", 0))
                    )
                    <= worksheet.max_row
                    else None
                ),
            )
        for seg in segments:
            row_num = start_row + int(
                seg.get("row_index", seg.get("id", 0))
            )
            if row_num < start_row or row_num > worksheet.max_row:
                raise ValueError(
                    f"source row for segment {seg['id']} is missing"
                )
            kind = export_kind(seg)
            corrected = result_entries[seg["id"]].get("corrected")
            if kind != "已保护" and corrected is not None:
                _set_excel_text(
                    worksheet.cell(row=row_num, column=ti + 1), corrected
                )
            counts[kind] += 1
        with tempfile.NamedTemporaryFile(
            dir=out_path.parent,
            prefix=f".{out_path.stem}.",
            suffix=out_path.suffix,
            delete=False,
        ) as handle:
            staged = Path(handle.name)
        workbook.save(str(staged))
        workbook.close()
        workbook = None
        if file_sha256(src_path) != source_digest:
            raise ValueError(f"source input changed during export: {src_path}")
        revalidate_inputs()
        publish_replacement_transaction([(staged, out_path)])
    except (OSError, ValueError) as exc:
        raise SystemExit(f"[export] {exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()
        if staged is not None:
            staged.unlink(missing_ok=True)
    print_summary(out_path)


def cmd_export(args):
    from lqe_chunk import verification_generation_lease
    from lqe_split_contract import generation_lock, state_fingerprint

    state_path = Path(args.state)
    errors_path = Path(args.errors) if getattr(args, "errors", None) else None
    try:
        require_current_job_runtime(read_json(state_path), "export")
        if errors_path is not None:
            with verification_generation_lease(
                state_path,
                exclusive=False,
            ) as (state, segments, manifest, revalidate):
                _cmd_export_locked(
                    args,
                    state_path,
                    state,
                    segments,
                    manifest,
                    revalidate,
                )
            return

        chunks_dir = state_path.parent / "chunks"
        with generation_lock(chunks_dir, exclusive=False):
            state = read_json(state_path)
            expected_state_fingerprint = state_fingerprint(state)

            def revalidate_state():
                if state_fingerprint(read_json(state_path)) != expected_state_fingerprint:
                    raise ValueError("state changed during export")

            _cmd_export_locked(
                args,
                state_path,
                state,
                deepcopy(state["segments"]),
                None,
                revalidate_state,
            )
    except (OSError, ValueError) as exc:
        raise SystemExit(f"[export] {exc}") from exc


# ── ingest-corpus (stub) ──────────────────────────────────────────────────────

def cmd_ingest_corpus(args):
    # TODO: 接口格式待确认（JSON 直传 vs 文件上传）
    print("[lqe_io] ingest-corpus: AIPE RAG ingest interface TBD, skipping.")


# ── main ──────────────────────────────────────────────────────────────────────

def _add_scoring_policy_args(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--threshold", type=float, default=None)
    parser.add_argument(
        "--scorecard-profile",
        default=None,
        dest="scorecard_profile",
        help="评分卡 profile；省略时继承 state.scoring_policy",
    )
    parser.add_argument(
        "--severity-scale",
        choices=["lisa", "mqm"],
        default=None,
        dest="severity_scale",
        help="严重度乘数档；省略时继承 state.scoring_policy",
    )
    critical = parser.add_mutually_exclusive_group()
    critical.add_argument(
        "--critical-gate", action="store_true", dest="critical_gate"
    )
    critical.add_argument(
        "--no-critical-gate", action="store_false", dest="critical_gate"
    )
    repeat = parser.add_mutually_exclusive_group()
    repeat.add_argument(
        "--repeat-dedup", action="store_true", dest="repeat_dedup"
    )
    repeat.add_argument(
        "--no-repeat-dedup", action="store_false", dest="repeat_dedup"
    )
    parser.set_defaults(critical_gate=None, repeat_dedup=None)


def main():
    p = argparse.ArgumentParser()
    sub = p.add_subparsers(dest="cmd", required=True)

    r = sub.add_parser("read")
    r.add_argument("--input", required=True)
    r.add_argument(
        "--input-format",
        choices=["auto", "tabular", "sdlxliff"],
        default="auto",
        dest="input_format",
    )
    r.add_argument(
        "--protect-exact-tm",
        action="store_true",
        dest="protect_exact_tm",
        help="保护同时满足 origin=TM、100%% 和 SourceAndTarget 的 SDLXLIFF 段",
    )
    r.add_argument("--project", default=None, help="项目档案：projects/<名>/profile.json 或目录/文件路径；提供 SG/术语/词数基准/checks/confirmed_rules 默认值，显式参数优先")
    r.add_argument("--profile-overlay", default=None, dest="profile_overlay",
                   help="内部 profile overlay JSON；不能修改项目身份或语言")
    r.add_argument(
        "--context-overrides",
        default=None,
        dest="context_overrides",
        help=(
            "经人工或授权来源核实的 job 级上下文 sidecar；只允许正式 "
            "foundation/enforce capability 字段，整批校验后原子发布"
        ),
    )
    r.add_argument("--sheet", default=None,
                   help="表格输入的主工作表；CSV/TSV 不适用")
    r.add_argument("--source-col", default=None, dest="source_col", help="列名或列索引（0-based，配合 --no-header）；表格输入必填")
    r.add_argument("--target-col", default=None, dest="target_col", help="列名或列索引（0-based，配合 --no-header）；表格输入必填")
    r.add_argument("--key-col", default=None, dest="key_col",
                   help="稳定业务 key 列；无值时按输入摘要、容器和行号生成")
    r.add_argument("--context-col", action="append", default=[], dest="context_cols",
                   help="通用上下文列 FIELD=COLUMN，可重复")
    for flag, destination in (
        ("--content-type-col", "content_type_col"),
        ("--speaker-col", "speaker_col"),
        ("--addressee-col", "addressee_col"),
        ("--relationship-stage-col", "relationship_stage_col"),
        ("--scene-id-col", "scene_id_col"),
        ("--scene-tone-col", "scene_tone_col"),
        ("--context-note-col", "context_note_col"),
    ):
        r.add_argument(flag, default=None, dest=destination)
    r.add_argument("--pivot-sheet", default=None, dest="pivot_sheet")
    r.add_argument("--pivot-key-col", default=None, dest="pivot_key_col")
    r.add_argument("--pivot-compare", action="append", default=[], dest="pivot_compare")
    r.add_argument(
        "--pivot-authority",
        choices=["authoritative", "diagnostic"],
        default=None,
        dest="pivot_authority",
    )
    r.add_argument("--target-source-digest-col", default=None,
                   dest="target_source_digest_col")
    r.add_argument("--no-header", action="store_true", dest="no_header", help="文件无表头行，source-col/target-col 为整数索引")
    r.add_argument("--group-col", default=None, dest="group_col", help="成组文本（对联/题目）的组标识列名或索引；同组段落 Step 2 合并评估")
    terminology = r.add_mutually_exclusive_group()
    terminology.add_argument("--terminology", default=None, help="术语表文件路径（.csv/.tsv/.json/.xlsx）")
    terminology.add_argument("--no-terminology", action="store_true", dest="no_terminology",
                             help="禁用术语、专名和术语审计检查")
    r.add_argument("--style-guide", default=None, dest="style_guide", help="风格指南文件路径（.txt/.md/.docx/.xlsx）")
    r.add_argument("--target-lang", default=None, dest="target_lang",
                   help="目标语言代码（en/th/zh 等）；挂载 target_languages/<code>/ 目标语言属性。项目 profile 必须显式写 target_lang")
    r.add_argument("--source-lang", default=None, dest="source_lang",
                   help="源语言代码（zh/en 等）；项目 profile 必须显式写 source_lang，显式参数优先")
    r.add_argument("--wordcount-basis", default=None, choices=["target-words", "source-chars"],
                   dest="wordcount_basis",
                   help="词数基准：target-words=译文空格分词（EN 等）；source-chars=源文 CJK 字符数+拉丁词数（泰语等无空格译文用）")
    r.add_argument("--out", default="lqe_state.json")
    r.add_argument(
        "--review-mode",
        choices=["optimized", "full"],
        default="optimized",
        dest="review_mode",
        help=(
            "审校输出模式：optimized=降本规则；full=完整建议行为。"
            "Agent 应在初始化前向用户确认。"
        ),
    )

    rr = sub.add_parser("reread")
    rr.add_argument("--from-job", required=True, dest="from_job")
    rr.add_argument("--input", required=True)
    rr.add_argument("--job", required=True)
    rr.add_argument("--profile-overlay", default=None, dest="profile_overlay")
    rr.add_argument("--context-overrides", default=None, dest="context_overrides")

    af = sub.add_parser("apply-fixes")
    af.add_argument("--state",     required=True)
    af.add_argument("--errors",    required=True)
    af.add_argument("--score",     default=None, help="本轮分数（来自 lqe_calc.py 输出）")
    _add_scoring_policy_args(af)
    af.add_argument("--protected-ids", default=None, help="逗号分隔的 TM 100%% match segment ids")
    af.add_argument("--protected-file", default=None, help="TM 100%% match protected ids JSON 文件")

    ps = sub.add_parser("protect-segments")
    ps.add_argument("--state", required=True)
    ps.add_argument("--protected-ids", default=None, help="逗号分隔的已确认段 id")
    ps.add_argument("--protected-file", default=None, help="已确认段 id JSON 文件")
    ps.add_argument("--reason", default="TM_100_MATCH")
    ps.add_argument("--out", default=None, help="输出 protected ids JSON，默认 {job}/tm_protected.json")

    br = sub.add_parser("build-results")
    br.add_argument("--state", required=True)
    br.add_argument("--checks", required=True)
    br.add_argument("--out", required=True)

    w = sub.add_parser("write")
    w.add_argument("--state",     required=True)
    w.add_argument("--errors",    required=True)
    w.add_argument("--score",     required=True)
    _add_scoring_policy_args(w)

    pc = sub.add_parser("pre-check")
    pc.add_argument("--state", required=True)
    pc.add_argument("--out", default=None, help="输出路径（默认 {job_dir}/errors_precheck.json）")

    lt = sub.add_parser("lookup-terms")
    lt.add_argument("--state", required=True)
    lt.add_argument("--ids", default=None, help="逗号分隔的 seg id，不传则扫描全部段落")

    ex = sub.add_parser("export")
    ex.add_argument("--state", required=True)
    ex.add_argument("--errors", default=None, help="可选 errors.json：state.corrected 为空时用其 corrected 填充建议修正（单轮 FAIL 导出用）")

    ic = sub.add_parser("ingest-corpus")
    ic.add_argument("--state",    required=True)
    ic.add_argument("--aipe-url", required=True, dest="aipe_url")

    args = p.parse_args()
    {
        "read":           cmd_read,
        "reread":         cmd_reread,
        "pre-check":      cmd_pre_check,
        "protect-segments": cmd_protect_segments,
        "build-results":  cmd_build_results,
        "apply-fixes":    cmd_apply_fixes,
        "write":          cmd_write,
        "export":         cmd_export,
        "lookup-terms":   cmd_lookup_terms,
        "ingest-corpus":  cmd_ingest_corpus,
    }[args.cmd](args)


if __name__ == "__main__":
    main()
