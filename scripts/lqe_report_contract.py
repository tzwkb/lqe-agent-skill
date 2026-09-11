"""Bind generated LQE workbooks to the state and verified results they render."""

from __future__ import annotations

from datetime import date, datetime, time
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText

from lqe_engine import requires_bound_artifacts
from lqe_provenance import AUDIT_HEADER_BASES, issue_detail, issue_review_columns
from lqe_split_contract import canonical_digest


SHEET_NAME = "_LQE_CONTRACT"
SCHEMA = "lqe.report-contract"
VERSION = 7
LEGACY_VERSION = 5
NATIVE_WRITER = "lqe_io._build_xlsx"
NATIVE_OUTPUT_MODE = "native_internal_writer"
REFERENCE_TEMPLATE_MODE = "reference_only"

CONTEXT_AUDIT_HEADERS = (
    "Segment Key",
    "Input Status",
    "Content Type",
    "Context Status",
    "Context JSON",
    "Context Provenance JSON",
    "Context Digest",
    "Resolved Constraints JSON",
    "Capability Resolution Digest",
    "Project Asset Snapshot Digest",
)


def _audit_column(headers: list[object], base: str) -> int:
    candidates = [
        index
        for index, header in enumerate(headers, start=1)
        if header == base
        or (
            isinstance(header, str)
            and header.startswith(f"{base}（审计 ")
        )
    ]
    if not candidates:
        raise ValueError(f"LQE Results is missing required audit column: {base}")
    return candidates[-1]


def context_audit_values(state: dict, segment: dict) -> tuple[object, ...]:
    context = segment.get("context") or {}
    if not isinstance(context, dict):
        raise ValueError("report segment context must be an object")
    context_provenance = segment.get("context_provenance") or {}
    if not isinstance(context_provenance, dict):
        raise ValueError("report segment context_provenance must be an object")
    resolved_constraints = segment.get("resolved_constraints") or []
    if not isinstance(resolved_constraints, list):
        raise ValueError("report segment resolved_constraints must be an array")
    core = context.get("core") or {}
    if not isinstance(core, dict):
        raise ValueError("report segment context.core must be an object")
    input_status = segment.get("input_status", "ready")
    content_type = segment.get("content_type") or core.get("content_type") or ""
    return (
        segment.get("segment_key", "") or "",
        input_status,
        content_type,
        "blocked" if input_status == "blocked" else context.get("status", "ready"),
        json.dumps(context, ensure_ascii=False, sort_keys=True),
        json.dumps(context_provenance, ensure_ascii=False, sort_keys=True),
        segment.get("segment_revision_digest", "") or "",
        json.dumps(resolved_constraints, ensure_ascii=False, sort_keys=True),
        state.get("capability_resolution_digest", "") or "",
        state.get("project_asset_snapshot_digest", "") or "",
    )


def _normalized_audit_value(value: object) -> object:
    return "" if value is None else value


def _validate_current_results_shape(workbook, state: dict, results: list[dict]) -> None:
    current_runtime = state.get("job_runtime_contract_version") == 2
    if not requires_bound_artifacts(state):
        return
    if "LQE Results" not in workbook.sheetnames:
        raise ValueError("child report is missing LQE Results")
    sheet = workbook["LQE Results"]
    headers = [cell.value for cell in sheet[1]]
    if not headers or headers[-1] != "LQE_Iter":
        raise ValueError(
            "LQE Results must place LQE_Iter in the last column"
        )
    context_columns = {}
    if current_runtime:
        missing_context_headers = [
            header for header in CONTEXT_AUDIT_HEADERS if header not in headers
        ]
        if missing_context_headers:
            raise ValueError(
                "LQE Results is missing required context audit columns: "
                + ", ".join(missing_context_headers)
            )
        context_columns = {
            header: _audit_column(headers, header)
            for header in CONTEXT_AUDIT_HEADERS
        }
    columns = {
        key: _audit_column(headers, base)
        for key, base in AUDIT_HEADER_BASES.items()
    }
    if "LQA Scorecard" not in workbook.sheetnames:
        raise ValueError("current report is missing LQA Scorecard")
    try:
        detail_column = headers.index("错误详情") + 1
    except ValueError as exc:
        raise ValueError(
            "LQE Results is missing required audit column: 错误详情"
        ) from exc

    state_ids = [segment.get("id") for segment in state.get("segments", [])]
    result_ids = [entry.get("id") for entry in results if isinstance(entry, dict)]
    if (
        any(type(value) is not int for value in state_ids + result_ids)
        or len(result_ids) != len(results)
        or len(set(state_ids)) != len(state_ids)
        or len(set(result_ids)) != len(result_ids)
        or set(result_ids) != set(state_ids)
    ):
        raise ValueError("report results do not cover the current state exactly")
    by_id = {entry["id"]: entry for entry in results}

    expected = []
    segments_by_id = {segment["id"]: segment for segment in state.get("segments", [])}
    for segment_id in state_ids:
        errors = by_id[segment_id].get("errors")
        if not isinstance(errors, list):
            raise ValueError("report result errors must be arrays")
        for issue_number, issue in enumerate(errors or [None], start=1):
            if issue is not None and not isinstance(
                issue.get("review_provenance"), dict
            ):
                raise ValueError(
                    "current report errors require explicit review_provenance"
                )
            review_status, edit_status, source = issue_review_columns(
                issue,
                segment_id,
            )
            expected_row = (
                segment_id,
                issue_number if issue is not None else None,
                review_status,
                edit_status,
                source,
                issue_detail(issue),
            )
            if current_runtime:
                expected_row += context_audit_values(
                    state,
                    segments_by_id[segment_id],
                )
            expected.append(expected_row)

    actual = []
    for row in range(2, sheet.max_row + 1):
        segment_id = sheet.cell(row, columns["segment_id"]).value
        if segment_id in (None, ""):
            continue
        issue_number = sheet.cell(row, columns["issue_number"]).value
        if issue_number == "":
            issue_number = None
        actual_row = (
            segment_id,
            issue_number,
            sheet.cell(row, columns["review_status"]).value,
            sheet.cell(row, columns["edit_status"]).value,
            sheet.cell(row, columns["check_source"]).value,
            sheet.cell(row, detail_column).value or "",
        )
        if current_runtime:
            actual_row += tuple(
                _normalized_audit_value(
                    sheet.cell(row, context_columns[header]).value
                )
                for header in CONTEXT_AUDIT_HEADERS
            )
        actual.append(actual_row)
    if actual != expected:
        raise ValueError(
            "LQE Results rows do not match current segment/error provenance"
        )


def _cell_content(cell) -> dict:
    value = cell.value
    if value in (None, ""):
        return {"data_type": "blank", "value": None}
    if isinstance(value, CellRichText):
        return {"data_type": "text", "value": str(value)}
    elif isinstance(value, (date, datetime, time)):
        return {"data_type": "date", "value": value.isoformat()}
    elif isinstance(value, bool):
        return {"data_type": "boolean", "value": value}
    elif isinstance(value, (int, float)):
        normalized = (
            int(value)
            if isinstance(value, float) and value.is_integer()
            else value
        )
        return {"data_type": "number", "value": normalized}
    elif isinstance(value, str):
        return {
            "data_type": "formula" if cell.data_type == "f" else "text",
            "value": value,
        }
    return {"data_type": "text", "value": str(value)}


def _visible_sheet_digest(workbook, sheet_name: str) -> str:
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"child report is missing {sheet_name}")
    sheet = workbook[sheet_name]
    payload = {
        "max_row": sheet.max_row,
        "max_column": sheet.max_column,
        "merged_ranges": sorted(str(cell_range) for cell_range in sheet.merged_cells.ranges),
        "cells": [
            [_cell_content(cell) for cell in row]
            for row in sheet.iter_rows(
                min_row=1,
                max_row=sheet.max_row,
                min_col=1,
                max_col=sheet.max_column,
            )
        ],
    }
    return canonical_digest(payload)


def _visible_results_digest(workbook) -> str:
    return _visible_sheet_digest(workbook, "LQE Results")


def _visible_scorecard_digest(workbook) -> str:
    return _visible_sheet_digest(workbook, "LQA Scorecard")


def build_report_contract(workbook, state: dict, results: list[dict]) -> dict:
    _validate_current_results_shape(workbook, state, results)
    current_runtime = state.get("job_runtime_contract_version") == 2
    payload = {
        "schema": SCHEMA,
        "version": VERSION if current_runtime else LEGACY_VERSION,
        "state_digest": canonical_digest(state),
        "results_digest": canonical_digest(results),
        "visible_results_digest": _visible_results_digest(workbook),
        "visible_scorecard_digest": (
            _visible_scorecard_digest(workbook)
            if "LQA Scorecard" in workbook.sheetnames
            else None
        ),
        "delivery": {
            "writer": NATIVE_WRITER,
            "output_mode": NATIVE_OUTPUT_MODE,
            "reference_template_mode": REFERENCE_TEMPLATE_MODE,
            "reference_template_applied": False,
            "job_id": state.get("job_id") or "unbound",
            "job_created_at": state.get("created_at"),
            "scorecard_profile": (
                state.get("scoring_policy", {}).get("scorecard_profile")
                if isinstance(state.get("scoring_policy"), dict)
                else None
            ),
        },
    }
    if current_runtime:
        payload.update({
            "job_runtime_contract_version": 2,
            "context_contract_version": state.get("context_contract_version"),
            "context_audit_headers": list(CONTEXT_AUDIT_HEADERS),
        })
    payload["contract_digest"] = canonical_digest(payload)
    return payload


def attach_report_contract(workbook, state: dict, results: list[dict]) -> None:
    encoded = json.dumps(
        build_report_contract(workbook, state, results),
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    if SHEET_NAME in workbook.sheetnames:
        del workbook[SHEET_NAME]
    sheet = workbook.create_sheet(SHEET_NAME)
    sheet["A1"] = encoded
    sheet.sheet_state = "veryHidden"


def validate_report_contract(workbook, state: dict, results: list[dict]) -> None:
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"child report is missing {SHEET_NAME}")
    raw = workbook[SHEET_NAME]["A1"].value
    if not isinstance(raw, str):
        raise ValueError("child report contract is invalid")
    try:
        actual = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise ValueError("child report contract is invalid") from exc
    expected = build_report_contract(workbook, state, results)
    if actual != expected:
        fields = sorted(
            key
            for key in set(actual) | set(expected)
            if actual.get(key) != expected.get(key)
        )
        raise ValueError(
            "child report is stale for current state/errors "
            f"(mismatch: {', '.join(fields)})"
        )


def verify_native_report(output_path, state: dict, results: list[dict]) -> dict:
    """Fail closed unless *output_path* is a current native LQE report.

    This is deliberately stricter than ``validate_report_contract``: aggregation
    and historical readers may accept a legacy contract, while final delivery
    must prove that the native writer created the workbook.
    """
    output_path = Path(output_path)
    if not output_path.is_file():
        raise ValueError(f"native LQE report is missing: {output_path}")
    if state.get("job_runtime_contract_version") != 2:
        raise ValueError("native report verification requires runtime contract v2")
    workbook = load_workbook(
        str(output_path),
        rich_text=True,
        data_only=False,
    )
    try:
        validate_report_contract(workbook, state, results)
        raw = workbook[SHEET_NAME]["A1"].value
        contract = json.loads(raw)
        if contract.get("version") != VERSION:
            raise ValueError(
                f"native report contract version must be {VERSION}"
            )
        delivery = contract.get("delivery")
        if not isinstance(delivery, dict):
            raise ValueError("native report is missing delivery provenance")
        expected = {
            "writer": NATIVE_WRITER,
            "output_mode": NATIVE_OUTPUT_MODE,
            "reference_template_mode": REFERENCE_TEMPLATE_MODE,
            "reference_template_applied": False,
            "job_id": state.get("job_id") or "unbound",
            "job_created_at": state.get("created_at"),
            "scorecard_profile": (
                state.get("scoring_policy", {}).get("scorecard_profile")
                if isinstance(state.get("scoring_policy"), dict)
                else None
            ),
        }
        if delivery != expected:
            raise ValueError(
                "native report delivery provenance does not match current job"
            )
        return {
            "report": str(output_path),
            "contract_digest": contract["contract_digest"],
            "writer": delivery["writer"],
            "job_id": delivery["job_id"],
        }
    finally:
        workbook.close()
