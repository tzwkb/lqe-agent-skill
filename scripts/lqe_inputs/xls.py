"""Read legacy BIFF ``.xls`` workbooks without rewriting the source file."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any


class XLSImportError(ValueError):
    pass


@dataclass(frozen=True)
class XLSImportResult:
    sheet_name: str
    sheet_names: list[str]
    headers: list[Any]
    data_rows: list[list[Any]]
    manifest: dict


def _xlrd_module():
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover - depends on host runtime
        raise XLSImportError(
            "legacy .xls input requires xlrd>=2.0; no job artifacts were published"
        ) from exc
    version = getattr(xlrd, "__version__", "0")
    try:
        major = int(str(version).split(".", 1)[0])
    except ValueError as exc:  # pragma: no cover - defensive
        raise XLSImportError(f"cannot determine xlrd version: {version!r}") from exc
    if major < 2:
        raise XLSImportError(f"legacy .xls input requires xlrd>=2.0, found {version}")
    return xlrd


def _cell_value(book, cell):
    xlrd = _xlrd_module()
    value = cell.value
    if cell.ctype == xlrd.XL_CELL_DATE:
        value = xlrd.xldate.xldate_as_datetime(value, book.datemode)
    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
        value = bool(value)
    elif cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
        return None
    elif cell.ctype == xlrd.XL_CELL_ERROR:
        return f"#XLERR:{int(value)}"
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _sheet_rows(book, sheet) -> list[list[Any]]:
    rows = []
    for row_index in range(sheet.nrows):
        rows.append(
            [
                _cell_value(book, sheet.cell(row_index, column_index))
                for column_index in range(sheet.ncols)
            ]
        )
    return rows


def read_xls(
    path: Path,
    *,
    sheet_name: str | None = None,
    no_header: bool = False,
) -> XLSImportResult:
    path = Path(path)
    if path.suffix.casefold() != ".xls":
        raise XLSImportError(f"legacy XLS adapter only accepts .xls: {path}")
    xlrd = _xlrd_module()
    try:
        book = xlrd.open_workbook(str(path), on_demand=True)
    except Exception as exc:
        raise XLSImportError(f"cannot read legacy .xls workbook {path}: {exc}") from exc
    try:
        names = list(book.sheet_names())
        if not names:
            raise XLSImportError(f"legacy .xls workbook has no sheets: {path}")
        selected = sheet_name or names[0]
        if selected not in names:
            raise XLSImportError(
                f"sheet {selected!r} not found; available sheets: {names}"
            )
        rows = _sheet_rows(book, book.sheet_by_name(selected))
        width = max((len(row) for row in rows), default=0)
        if no_header:
            defaults = [
                "Key", "Source", "Target", "Status", "Comment", "Scope",
                "File", "Reviewer Note",
            ]
            headers = [
                defaults[index] if index < len(defaults) else f"col{index}"
                for index in range(width)
            ]
            data_rows = rows
        else:
            headers = rows[0] if rows else []
            data_rows = rows[1:] if rows else []
        return XLSImportResult(
            sheet_name=selected,
            sheet_names=names,
            headers=headers,
            data_rows=data_rows,
            manifest={
                "schema": "lqe.tabular-source-manifest",
                "version": 1,
                "adapter": "xls.xlrd@1",
                "format": "xls",
                "sheet": selected,
                "sheet_names": names,
                "rows": len(rows),
                "columns": width,
                "coordinate_system": "zero_based_data_row_and_column",
                "limitations": [
                    "cell_values_and_coordinates_only",
                    "formulas_styles_comments_drawings_not_round_tripped",
                    "corrected_export_is_xlsx",
                ],
            },
        )
    finally:
        book.release_resources()


def workbook_for_corrected_export(path: Path):
    """Return an openpyxl workbook containing the value grid of every XLS sheet."""
    path = Path(path)
    xlrd = _xlrd_module()
    try:
        book = xlrd.open_workbook(str(path), on_demand=True)
    except Exception as exc:
        raise XLSImportError(f"cannot read legacy .xls workbook {path}: {exc}") from exc
    try:
        import openpyxl

        output = openpyxl.Workbook()
        output.remove(output.active)
        for name in book.sheet_names():
            source = book.sheet_by_name(name)
            target = output.create_sheet(title=name[:31] or "Sheet")
            for row_index, row in enumerate(_sheet_rows(book, source), start=1):
                for column_index, value in enumerate(row, start=1):
                    target.cell(row=row_index, column=column_index, value=value)
        if not output.worksheets:
            output.create_sheet("Sheet1")
        return output
    finally:
        book.release_resources()
